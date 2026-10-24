"""
theme.py - one look for the whole PO -> Packing List project
=============================================================
* PALETTE / FONT  : single source of truth shared by the Streamlit preview
                    and the Excel files (same yellow, same grey, same font).
* inject_theme()  : global Streamlit CSS (called once by the router).
                    Cards, rounded inputs, accent buttons, and EVERY table
                    cell (HTML previews, sidebar tables) centred.
* hero()          : page header banner used by the router.
* polish_workbook : Excel finishing - print setup + view + look.
                    Every cell is centred (horizontal + vertical), one font,
                    clean gridline-free view, comfortable row heights and
                    column widths that never truncate. It never touches cell
                    values, formulas, fills, borders, merges, rotations or
                    wrap flags, so every buyer's packing-list layout / rules
                    stay exactly as they are.
* install_xlsx_polish() : wraps st.download_button once per process so every
                    .xlsx the buyer modules hand to the user is polished.
* install_dataframe_safe() : display-only helpers for st.dataframe
                    (Arrow-safe columns, centred text).
"""
from __future__ import annotations

import io
from copy import copy

import streamlit as st

# ------------------------------------------------------------
# Palette - the SAME hex values are used in Excel fills and in CSS
# ------------------------------------------------------------
FONT_XLSX = "Arial"                 # font used by the Excel packing lists
FONT_UI = "Arial, 'Segoe UI', Roboto, Helvetica, sans-serif"

YELLOW = "FFFF00"                   # manual-entry cells (CUT QTY, CTN MEAS ...)
GREY = "D9D9D9"                     # ORDER QTY / header grey
DARK = "1F2937"                     # table header
OK_BG, OK_FG = "DCFCE7", "166534"
WARN_BG, WARN_FG = "FEF3C7", "92400E"
BAD_BG, BAD_FG = "FEE2E2", "991B1B"
ACCENT = "2563EB"
ACCENT_DARK = "1D4ED8"
BORDER = "CBD5E1"
SOFT = "F1F5F9"


def _css() -> str:
    return f"""
<style>
/* ================= smooth: no grey-out / flicker on rerun ================= */
[data-stale="true"], .stale-element {{
    opacity: 1 !important;
    transition: none !important;
}}
[data-testid="stAppViewContainer"] .main .block-container {{
    padding-top: 1.6rem;
    max-width: 100%;
}}

/* ================= typography ================= */
html, body, [data-testid="stAppViewContainer"], [data-testid="stSidebar"] {{
    font-family: {FONT_UI};
}}
h1, h2, h3, h4 {{ letter-spacing: -0.01em; }}
h2 {{ font-weight: 700 !important; }}
h3, h4, h5 {{ font-weight: 650 !important; }}
[data-testid="stCaptionContainer"] {{ opacity: .8; }}
hr {{ margin: 1.1rem 0 !important; opacity: .35; }}

/* ================= hero banner (router) ================= */
.po-hero {{
    display: flex; align-items: center; justify-content: center; gap: 16px;
    text-align: center;
    padding: 22px 26px; margin: 0 0 18px 0;
    border-radius: 18px;
    background: linear-gradient(120deg, #{DARK} 0%, #1e3a8a 55%, #{ACCENT} 100%);
    box-shadow: 0 10px 28px rgba(30, 58, 138, .22);
}}
.po-hero .po-hero-icon {{
    font-size: 40px; line-height: 1;
    background: rgba(255,255,255,.14); border-radius: 14px; padding: 10px 14px;
}}
.po-hero .po-hero-title {{
    margin: 0; color: #fff; font-size: 30px; font-weight: 800; letter-spacing: -0.02em;
}}
.po-hero .po-hero-sub {{
    margin: 4px 0 0 0; color: rgba(255,255,255,.82); font-size: 14px; font-weight: 400;
}}

/* ================= file uploaders = drop cards ================= */
[data-testid="stFileUploader"] > section,
[data-testid="stFileUploaderDropzone"] {{
    border: 2px dashed rgba(37, 99, 235, .45) !important;
    border-radius: 16px !important;
    background: rgba(37, 99, 235, .04) !important;
    transition: background .15s ease, border-color .15s ease;
}}
[data-testid="stFileUploader"] > section:hover,
[data-testid="stFileUploaderDropzone"]:hover {{
    background: rgba(37, 99, 235, .09) !important;
    border-color: #{ACCENT} !important;
}}
[data-testid="stFileUploader"] label p {{ font-weight: 700; }}
[data-testid="stFileUploaderFile"] {{ border-radius: 10px; }}

/* ================= buttons / inputs ================= */
.stButton > button, .stDownloadButton > button {{
    border-radius: 10px;
    font-weight: 600;
    padding: .45rem 1rem;
    border: 1px solid rgba(100, 116, 139, .35);
    transition: background .12s ease, border-color .12s ease,
                box-shadow .12s ease, transform .12s ease;
}}
.stButton > button:hover, .stDownloadButton > button:hover {{
    border-color: #{ACCENT};
    box-shadow: 0 4px 12px rgba(37, 99, 235, .18);
    transform: translateY(-1px);
}}
.stButton > button:active, .stDownloadButton > button:active {{ transform: translateY(0); }}
.stDownloadButton > button, .stButton > button[kind="primary"] {{
    background: #{ACCENT}; color: #fff; border-color: #{ACCENT};
}}
.stDownloadButton > button:hover, .stButton > button[kind="primary"]:hover {{
    background: #{ACCENT_DARK}; border-color: #{ACCENT_DARK}; color: #fff;
}}
.stNumberInput input, .stTextInput input, .stTextArea textarea {{
    border-radius: 8px !important;
    text-align: center !important;
}}
.stNumberInput input:focus, .stTextInput input:focus, .stTextArea textarea:focus {{
    border-color: #{ACCENT} !important;
    box-shadow: 0 0 0 3px rgba(37, 99, 235, .16) !important;
}}
[data-baseweb="select"] > div {{ border-radius: 8px !important; }}
[data-baseweb="select"] [class*="singleValue"], [data-baseweb="select"] input {{
    text-align: center;
}}

/* ================= alerts, expanders, tabs, metrics ================= */
[data-testid="stAlert"] {{
    border-radius: 12px !important;
    border-left-width: 5px !important;
    text-align: center;
}}
[data-testid="stAlert"] > div {{ justify-content: center; }}
[data-testid="stExpander"] {{
    border-radius: 12px !important;
    border: 1px solid rgba(100, 116, 139, .28) !important;
    overflow: hidden;
}}
[data-testid="stExpander"] summary {{ font-weight: 600; }}
button[data-baseweb="tab"] {{ font-weight: 600; }}
button[data-baseweb="tab"][aria-selected="true"] {{ color: #{ACCENT}; }}
[data-baseweb="tab-highlight"] {{ background-color: #{ACCENT} !important; height: 3px !important; }}
[data-testid="stMetric"] {{
    text-align: center;
    border: 1px solid rgba(100, 116, 139, .28);
    border-radius: 14px; padding: 12px 8px;
    background: rgba(100, 116, 139, .05);
}}
[data-testid="stMetric"] > div {{ align-items: center; justify-content: center; }}
[data-testid="stMetricLabel"] {{ justify-content: center; }}
[data-testid="stDataFrame"], [data-testid="stDataEditor"] {{
    border-radius: 12px; overflow: hidden;
    border: 1px solid rgba(100, 116, 139, .28);
}}

/* ================= sidebar ================= */
[data-testid="stSidebar"] {{ border-right: 1px solid rgba(100, 116, 139, .25); }}
[data-testid="stSidebar"] h3 {{
    font-size: 15px; text-transform: uppercase; letter-spacing: .06em; opacity: .85;
}}

/* ================= HTML preview tables: ONE look, EVERY cell centred ================= */
.plist-wrap, table.plist, table.plist td, table.plist th,
table.order-tbl, table.order-tbl td, table.order-tbl th,
.sidebar-qty, .sidebar-qty td, .sidebar-qty th {{
    font-family: {FONT_UI} !important;
}}
table.plist td, table.plist th,
table.order-tbl td, table.order-tbl th,
table.order-tbl .rowlbl, .desc-cell,
.sidebar-qty td, .sidebar-qty th, .sidebar-qty td:first-child {{
    text-align: center !important;
    vertical-align: middle !important;
}}
:where(table.plist td, table.plist th, table.order-tbl td) {{ color: #000; }}
:where(table.plist td, table.order-tbl td) {{ background: #fff; }}
table.plist, table.order-tbl, .sidebar-qty {{ border-collapse: collapse; }}
table.plist td.yellow, td.yellow {{ background: #{YELLOW} !important; }}
table.plist td.gray,   td.gray   {{ background: #{GREY} !important; font-weight: 700; }}
table.plist th {{ background: #{GREY}; color: #000; font-weight: 700; }}
table.plist tr:hover td:not([class]) {{ background: #{SOFT}; }}
.plist-wrap, .order-scroll {{
    border-radius: 10px; overflow: auto;
    box-shadow: 0 2px 10px rgba(15, 23, 42, .08);
}}
.sidebar-qty th {{ background: #{DARK}; color: #fff; }}
.sidebar-qty td:first-child {{ background: #{SOFT}; }}
.sidebar-qty tr.remain-positive td.remain-cell-sb {{ background:#{WARN_BG}; color:#{WARN_FG}; }}
.sidebar-qty tr.remain-zero     td.remain-cell-sb {{ background:#{OK_BG};   color:#{OK_FG};   }}
.sidebar-qty tr.remain-negative td.remain-cell-sb {{ background:#{BAD_BG};  color:#{BAD_FG};  }}
</style>
"""


def inject_theme() -> None:
    st.markdown(_css(), unsafe_allow_html=True)


def hero(title: str, subtitle: str = "", icon: str = "📦") -> None:
    """Page header banner (replaces a plain st.title)."""
    sub = f'<p class="po-hero-sub">{subtitle}</p>' if subtitle else ""
    st.markdown(
        f'<div class="po-hero"><div class="po-hero-icon">{icon}</div>'
        f'<div><h1 class="po-hero-title">{title}</h1>{sub}</div></div>',
        unsafe_allow_html=True,
    )


# ------------------------------------------------------------
# Excel: look + print polish
# ------------------------------------------------------------
_MIN_ROW_HEIGHT = 16.0         # only raises rows that were explicitly set smaller
_MAX_AUTO_WIDTH = 45.0


def _center_and_font(ws) -> None:
    """Centre every cell (keeping wrap / shrink / rotation) and use one font."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment

    last_row, last_col = ws.max_row, ws.max_column
    if not last_row or not last_col:
        return
    for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=last_col):
        for c in row:
            if isinstance(c, MergedCell):
                continue
            a = c.alignment
            c.alignment = Alignment(
                horizontal="center", vertical="center",
                wrap_text=a.wrap_text, shrink_to_fit=a.shrink_to_fit,
                text_rotation=a.text_rotation, indent=0,
            )
            f = c.font
            if f is not None and f.name != FONT_XLSX:
                f2 = copy(f)
                f2.name = FONT_XLSX
                f2.scheme = None
                c.font = f2


def _row_heights(ws) -> None:
    for r, dim in list(ws.row_dimensions.items()):
        if dim.height is not None and 0 < dim.height < _MIN_ROW_HEIGHT:
            dim.height = _MIN_ROW_HEIGHT


def _explicit_width_cols(ws) -> set:
    """Column indexes that already have a width defined (incl. grouped <col> ranges)."""
    from openpyxl.utils import column_index_from_string
    covered = set()
    for key, dim in ws.column_dimensions.items():
        if dim.width is None:
            continue
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        covered.update(range(lo, hi + 1))
    return covered


def _widen_truncated_columns(ws) -> None:
    """Give a width only to columns that have none yet and whose text would be cut."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter

    covered = _explicit_width_cols(ws)
    merged = {(r, c) for rng in ws.merged_cells.ranges
              for r in range(rng.min_row, rng.max_row + 1)
              for c in range(rng.min_col, rng.max_col + 1)}
    for ci in range(1, (ws.max_column or 0) + 1):
        if ci in covered:
            continue
        need = 0
        for ri in range(1, (ws.max_row or 0) + 1):
            if (ri, ci) in merged:
                continue
            c = ws.cell(row=ri, column=ci)
            if isinstance(c, MergedCell) or c.value is None:
                continue
            if c.alignment.wrap_text or c.alignment.text_rotation:
                continue
            txt = str(c.value)
            if txt.startswith("="):
                continue
            longest = max((len(p) for p in txt.split("\n")), default=0)
            bold = 1.1 if (c.font and c.font.b) else 1.0
            need = max(need, longest * bold)
        width = min(need * 1.15 + 2, _MAX_AUTO_WIDTH)
        if width > 9:
            ws.column_dimensions[get_column_letter(ci)].width = round(width, 1)


def _fit_overflowing_text(ws) -> None:
    """Text wider than its cell used to spill over the empty cells to its right
    (it was left-aligned). Centring it inside the narrow cell would clip it, so:
      1. unmerged text -> 'Center Across Selection' over the same empty cells
         it used to spill into (no merging, nothing moves);
      2. otherwise (merged, or a neighbour is in the way) -> shrink-to-fit."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment
    from openpyxl.utils import column_index_from_string

    default_w = ws.sheet_format.defaultColWidth or 8.43
    widths = {}
    for key, dim in ws.column_dimensions.items():
        if dim.width is None:
            continue
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        for ci in range(lo, hi + 1):
            widths[ci] = dim.width

    def col_w(ci):
        return widths.get(ci, default_w)

    merged_cells = set()
    span = {}                                   # top-left -> total merged width
    for rng in ws.merged_cells.ranges:
        span[(rng.min_row, rng.min_col)] = sum(
            col_w(ci) for ci in range(rng.min_col, rng.max_col + 1))
        for r in range(rng.min_row, rng.max_row + 1):
            for ci in range(rng.min_col, rng.max_col + 1):
                merged_cells.add((r, ci))

    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for c in row:
            if isinstance(c, MergedCell) or not isinstance(c.value, str):
                continue
            a = c.alignment
            if a.wrap_text or a.text_rotation or c.value.startswith("="):
                continue
            size = (c.font.sz or 11) if c.font else 11
            bold = 1.1 if (c.font and c.font.b) else 1.0
            need = max((len(p) for p in c.value.split("\n")), default=0) * (size / 11.0) * bold
            merged_here = (c.row, c.column) in span
            avail = span.get((c.row, c.column), col_w(c.column))
            if need <= avail:
                continue

            if not merged_here:
                chain, total, ci = [], avail, c.column + 1
                while total < need and ci <= max_col:
                    nxt = ws.cell(row=c.row, column=ci)
                    if (c.row, ci) in merged_cells or nxt.value is not None:
                        break
                    chain.append(nxt)
                    total += col_w(ci)
                    ci += 1
                if chain and total >= need:
                    for cell in [c] + chain:
                        cell.alignment = Alignment(
                            horizontal="centerContinuous", vertical="center",
                            wrap_text=False, text_rotation=0, indent=0)
                    continue

            c.alignment = Alignment(
                horizontal="center", vertical="center", shrink_to_fit=True,
                wrap_text=False, text_rotation=0, indent=0)


def polish_workbook(wb) -> None:
    from openpyxl.worksheet.properties import PageSetupProperties

    for ws in wb.worksheets:
        # ---- print setup (unchanged behaviour) ----
        ps = ws.page_setup
        if not ps.orientation:
            ps.orientation = "landscape"
        if not ps.paperSize:
            ps.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = (
            ws.sheet_properties.pageSetUpPr or PageSetupProperties()
        )
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if ps.fitToWidth is None:
            ps.fitToWidth = 1
        if ps.fitToHeight is None:
            ps.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        m = ws.page_margins
        m.left = m.right = 0.4
        m.top = m.bottom = 0.5

        # ---- look (no values / formulas / fills / borders / merges touched) ----
        _center_and_font(ws)
        _row_heights(ws)
        _widen_truncated_columns(ws)
        _fit_overflowing_text(ws)
        ws.sheet_view.showGridLines = False
        if ws.sheet_properties.tabColor is None:
            ws.sheet_properties.tabColor = ACCENT


def _style_only_patch(data: bytes) -> bytes:
    """Workbooks with images / drawings / charts: openpyxl can't round-trip them
    losslessly, so ONLY xl/styles.xml (centre + font) and the sheet views
    (gridlines off) are edited - every other part is copied byte-for-byte."""
    import zipfile
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    q = lambda t: f"{{{ns}}}{t}"  # noqa: E731

    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                root = etree.fromstring(raw)
                for xf in root.iter(q("xf")):
                    if xf.getparent().tag != q("cellXfs"):
                        continue
                    xf.set("applyAlignment", "1")
                    al = xf.find(q("alignment"))
                    if al is None:
                        al = etree.SubElement(xf, q("alignment"))
                        # <alignment> must come before <protection>
                        prot = xf.find(q("protection"))
                        if prot is not None:
                            xf.remove(al)
                            prot.addprevious(al)
                    al.set("horizontal", "center")
                    al.set("vertical", "center")
                    al.attrib.pop("indent", None)
                for fnt in root.iter(q("font")):
                    if fnt.getparent().tag != q("fonts"):
                        continue
                    nm = fnt.find(q("name"))
                    if nm is not None:
                        nm.set("val", FONT_XLSX)
                    for tag in ("scheme",):
                        el = fnt.find(q(tag))
                        if el is not None:
                            fnt.remove(el)
                raw = etree.tostring(root, xml_declaration=True,
                                     encoding="UTF-8", standalone=True)
            elif (item.filename.startswith("xl/worksheets/")
                  and item.filename.endswith(".xml")):
                root = etree.fromstring(raw)
                for sv in root.iter(q("sheetView")):
                    sv.set("showGridLines", "0")
                raw = etree.tostring(root, xml_declaration=True,
                                     encoding="UTF-8", standalone=True)
            zout.writestr(item, raw)
    return out.getvalue()


def polish_xlsx_bytes(data: bytes) -> bytes:
    import zipfile
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        has_media = any(n.startswith(("xl/media/", "xl/drawings/", "xl/charts/"))
                        for n in z.namelist())
    if has_media:
        try:
            return _style_only_patch(data)
        except Exception:
            return data                               # never break a workbook

    wb = load_workbook(io.BytesIO(data))        # formulas are kept as formulas
    polish_workbook(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def install_xlsx_polish() -> None:
    """Idempotent: wrap st.download_button so every .xlsx gets polished."""
    if getattr(st.download_button, "_po_polished", False):
        return
    original = st.download_button

    def wrapped(label, data=None, file_name=None, mime=None, *a, **kw):
        try:
            is_xlsx = (mime == _XLSX_MIME) or str(file_name or "").lower().endswith(".xlsx")
            if is_xlsx and data is not None:
                raw = data.getvalue() if hasattr(data, "getvalue") else (
                    data if isinstance(data, (bytes, bytearray)) else None)
                if raw:
                    data = polish_xlsx_bytes(bytes(raw))
        except Exception:
            pass                                   # never block a download
        return original(label, data, file_name, mime, *a, **kw)

    wrapped._po_polished = True
    st.download_button = wrapped


# ------------------------------------------------------------
# Tables: avoid the failed-Arrow-conversion + retry on every rerun
# (columns that mix '' / 'TOTAL' strings with ints) and centre the text.
# Display only - st.dataframe is read-only, so nothing the user edits changes.
# ------------------------------------------------------------
def _arrow_safe(df):
    try:
        import pandas as pd
        if not isinstance(df, pd.DataFrame):
            return df
        obj_cols = [c for c in df.columns if df[c].dtype == object]
        mixed = [c for c in obj_cols
                 if len({type(v) for v in df[c].tolist() if v is not None}) > 1]
        if not mixed:
            return df
        df = df.copy()
        for c in mixed:
            df[c] = df[c].map(lambda v: "" if v is None else str(v))
        return df
    except Exception:
        return df


_CENTER = {"text-align": "center"}


def _centered(data):
    """Return a Styler with centred text (existing colours / formats are kept;
    the centre rule is added last so it wins over any earlier left-align)."""
    try:
        import pandas as pd
        from pandas.io.formats.style import Styler
        if isinstance(data, Styler):
            return data.set_properties(**_CENTER)
        if isinstance(data, pd.DataFrame) and data.columns.is_unique and data.index.is_unique:
            return data.style.set_properties(**_CENTER)
    except Exception:
        pass
    return data


def install_dataframe_safe() -> None:
    if getattr(st.dataframe, "_po_safe", False):
        return
    original = st.dataframe

    def wrapped(data=None, *a, **kw):
        safe = _arrow_safe(data)
        styled = _centered(safe)
        if styled is not safe:
            try:
                return original(styled, *a, **kw)
            except Exception:
                pass                              # fall back to the plain table
        return original(safe, *a, **kw)

    wrapped._po_safe = True
    st.dataframe = wrapped
