#!/usr/bin/env python3
"""
MUJI PO (発注一覧) -> Packing list generator  (Creative Collections Ltd.)

Usage:  python muji_packing_list.py <PO.xlsx> <output.xlsx> [--dest Vietnam,Malaysia] [--style BE43WA7S]

Rules derived from the two reference packing lists (Vietnam solid, Malaysia assort+solid):
  * One carton = 24 pcs.
  * PO rows that carry an assort JAN (cols R-Y) are ASSORT cartons: one carton per assort
    (W = number of assorts = number of cartons); pcs per size in the carton = V (balls in assort) x Q (pcs/ball).
  * PO rows with no assort JAN carry the SOLID qty (AB). Solid cartons are single-size, 24 pcs each.
    A remainder (<24) becomes one partial carton of that size.
  * PO total (AC) per size = solid (AB) + assort pcs (X).  Used as ORDER QTY in the summary.
  * Weights: Net = Gross - 1.21 (solid) / Gross - 1.5 (assort).  CBM = ctn x 55 x 37 x 32 / 1,000,000.
  * Carton type A = solid, B = assort.  Carton numbering restarts at 1 in every section (as in the references).
"""
import argparse, re, sys
from collections import defaultdict, OrderedDict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L

SIZES = ['XS', 'S', 'M', 'L', 'XL', 'XXL']
SCOL = {s: 8 + i for i, s in enumerate(SIZES)}            # H..M
CARTON_PCS = 24
CARTON_TXT = '幅(（W) 55 cm與行(D)37 cm  高之(H) 32 cm (Type-A)'
SOLID_TARE, ASSORT_TARE = 1.21, 1.5
COLOR_EN = {'068': 'Navy', '069': 'Dark Navy'}

# ---- weights known from the references (BE43WA7S only) ---------------------------------
SOLID_GROSS_KNOWN = {'BE43WA7S': {'S': 11.2, 'M': 11.57, 'L': 12.2, 'XL': 12.71}}
ASSORT_GROSS_KNOWN = {('BE43WA7S', (('S', 8), ('M', 8), ('L', 4), ('XL', 4))): 11.5}


def unit_net(style):
    """per-piece net weight (kg) by size, from reference solid cartons; XS/XXL extrapolated."""
    k = SOLID_GROSS_KNOWN.get(style)
    if not k:
        return None
    u = {s: (g - SOLID_TARE) / CARTON_PCS for s, g in k.items()}
    u['XS'] = u['S'] - (u['M'] - u['S'])
    u['XXL'] = u['XL'] + (u['XL'] - u['L'])
    return u


def est_assort_gross(style, mix):
    u = unit_net(style)
    if not u:
        return None
    k = (10.0) / sum(u[s] * n for s, n in (('S', 8), ('M', 8), ('L', 4), ('XL', 4)))  # calibrate to Malaysia 11.5
    return round(sum(u[s] * n for s, n in mix) * k + ASSORT_TARE, 2)


# ---- read PO ----------------------------------------------------------------------------
def code(v):  # '025:S' -> ('025','S')
    a, _, b = str(v).partition(':')
    return a.strip(), b.strip()


def read_po(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r[1] or not r[4]:
            continue
        cc, _ = code(r[13])
        rows.append(dict(dest=r[1], style=r[4], color_code=cc, size=code(r[14])[1], jan=str(r[15]),
                         ball=r[16] or 1, ajan=str(r[17]) if r[17] else None,
                         v=r[21], w=r[22], x=r[23], y=r[24] or 1, solid=r[27], total=r[28]))
    return rows


def build_groups(rows):
    """-> {(dest, style, color): dict(solid{size:(jan,qty)}, assort{ajan:{size:pcs}}, acount{ajan:n}, order{size:qty})}"""
    G = OrderedDict()
    issues = []
    for r in rows:
        g = G.setdefault((r['dest'], r['style'], r['color_code']),
                         dict(solid={}, assort=OrderedDict(), acount={}, order=defaultdict(int), assort_pcs=defaultdict(int)))
        if r['ajan']:
            pcs = r['v'] * r['ball'] * r['y']
            g['assort'].setdefault(r['ajan'], {})[r['size']] = pcs
            g['acount'].setdefault(r['ajan'], set()).add(r['w'])
            g['assort_pcs'][r['size']] += r['x']
            if r['x'] != r['v'] * r['w']:
                issues.append(f"{r['dest']} {r['style']} {r['size']}: assort total X={r['x']} != V*W")
        else:
            g['solid'][r['size']] = (r['jan'], r['solid'] or 0)
            g['order'][r['size']] += r['total'] or 0
    for k, g in G.items():
        for s in set(g['solid']) | set(g['assort_pcs']):
            solid = g['solid'].get(s, (None, 0))[1]
            if solid + g['assort_pcs'].get(s, 0) != g['order'].get(s, 0):
                issues.append(f"{k[0]} {k[1]} {s}: solid {solid} + assort {g['assort_pcs'].get(s,0)} != PO total {g['order'].get(s,0)}")
        for aj, mix in g['assort'].items():
            if sum(mix.values()) != CARTON_PCS:
                issues.append(f"{k[0]} {k[1]}: assort {aj} carton = {sum(mix.values())} pcs (expected {CARTON_PCS})")
            if len(g['acount'][aj]) != 1:
                issues.append(f"{k[0]} {k[1]}: assort {aj} has differing assort counts {g['acount'][aj]}")
    return G, issues


# ---- styling helpers -------------------------------------------------------------------
F = lambda sz=10, b=False, color=None: Font(name='Arial', size=sz, bold=b, color=color)
thin, med = Side(style='thin'), Side(style='medium')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
YELLOW = PatternFill('solid', fgColor='FFFF00')
GREEN = PatternFill('solid', fgColor='C6E0B4')
GREY = PatternFill('solid', fgColor='D9D9D9')


def put(ws, r, c, v=None, font=None, fill=None, fmt=None, border=BOX, align=CEN):
    cell = ws.cell(r, c)
    if v is not None:
        cell.value = v
    cell.font = font or F()
    cell.alignment = align
    if border:
        cell.border = border
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def header(ws, r, jan_label):
    ws.row_dimensions[r].height = 27
    for c, t in [(1, 'CTN'), (3, 'CTN QTY'), (4, 'CTN TYPE'), (5, 'Style'), (6, 'COLOR'), (7, jan_label),
                 (14, 'PER BLST'), (15, 'QTY PER CTN'), (16, 'TOTAL QTY'), (17, 'GROSS WEIGHT'), (18, 'NET WEIGHT'),
                 (19, 'TTL CTN GROSS WEIGHT'), (20, 'TTL CTN NET WEIGHT'), (21, 'TOTAL CBM')]:
        for rr in (r, r + 1):
            for cc in ((1, 2) if c == 1 else (c,)):
                put(ws, rr, cc, None, F(10, True), GREY)
        ws.cell(r, c).value = t
        if c != 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(r + 1, 1).value = 'NO'
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
    for c in range(8, 14):
        put(ws, r, c, None, F(10, True), GREY)
    ws.cell(r, 8).value = 'SIZE'
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=13)
    for s, c in SCOL.items():
        put(ws, r + 1, c, s, F(9, True), GREY)


def section(ws, r, jan_label, rows, tare):
    """rows: list of dict(qty, typ, style, color, jan, sizes{size:pcs}, perblst, gross, est). Returns (next_row, total_row)."""
    header(ws, r, jan_label)
    r += 2
    first = r
    for i, d in enumerate(rows):
        put(ws, r, 1, 1 if i == 0 else f'=B{r-1}+1', F(10, True))
        put(ws, r, 2, f'=A{r}-1+C{r}', F(10, True))
        put(ws, r, 3, d['qty'], F(10, True, '0000FF'))
        put(ws, r, 4, d['typ'])
        put(ws, r, 5, d['style'], F(9, True), GREEN)
        put(ws, r, 6, d['color'], F(9, True))
        put(ws, r, 7, d['jan'], F(9), fmt='@')
        for s, c in SCOL.items():
            put(ws, r, c, d['sizes'].get(s), F(10, False, '0000FF'))
        put(ws, r, 14, d['perblst'], F(9, True))
        put(ws, r, 15, f'=SUM(H{r}:M{r})', F(10, True))
        put(ws, r, 16, f'=O{r}*C{r}', F(10, True))
        put(ws, r, 17, d['gross'], F(10, True, '0000FF'), YELLOW if (d['est'] or d['gross'] is None) else None, '0.00')
        put(ws, r, 18, f'=IF(Q{r}="","",Q{r}-{tare})', F(10, True), None, '0.00')
        put(ws, r, 19, f'=IF(Q{r}="",0,Q{r}*C{r})', F(10, True), None, '0.00')
        put(ws, r, 20, f'=IF(Q{r}="",0,R{r}*C{r})', F(10, True), None, '0.00')
        put(ws, r, 21, f'=C{r}*55*37*32/1000000', F(10, True), None, '0.000')
        r += 1
    last = r - 1
    put(ws, r, 1, 'TOTAL', F(10, True), GREY); put(ws, r, 2, None, F(10, True), GREY)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    put(ws, r, 3, f'=SUM(C{first}:C{last})', F(10, True), GREEN)
    put(ws, r, 4, None, F(10, True), GREY)
    put(ws, r, 5, 'CTNS', F(10, True), GREY)
    for c in (6, 7, 14, 17, 18):
        put(ws, r, c, None, F(10, True), GREY)
    for s, c in SCOL.items():
        put(ws, r, c, f'=SUMPRODUCT({L(c)}{first}:{L(c)}{last},$C{first}:$C{last})', F(10, True), GREY, '#,##0;;')
    put(ws, r, 15, f'=SUM(O{first}:O{last})', F(10, True), GREY)
    for c in (16, 19, 20):
        put(ws, r, c, f'=SUM({L(c)}{first}:{L(c)}{last})', F(10, True), GREY, '#,##0.00' if c != 16 else '#,##0')
    put(ws, r, 21, f'=SUM(U{first}:U{last})', F(10, True), GREY, '0.000')
    return r + 2, r


def title(ws, text4, text3='PACKING LIST'):
    for r, (t, sz) in enumerate([('CREATIVE COLLECTIONS LTD.', 14), ('Nishatnagar,Tongi,Gazipur,Dhaka.', 11), (text3, 10), (text4, 10)], 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=21)
        c = ws.cell(r, 1, t); c.font = F(sz, True); c.alignment = Alignment(horizontal='center', vertical='center')


def make_sheet(wb, dest, style, color_code, g, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    color = f"{color_code}:{COLOR_EN.get(color_code, '')}"
    has_a, has_s = bool(g['assort']), any(q for _, q in g['solid'].values())
    ws.cell(5, 1, 'CARTON MEASUREMENT  :').font = F(10, True)
    ws.cell(5, 6, CARTON_TXT).font = F()
    r = 7
    tot_rows = {'assort': None, 'solid': None}
    # ---- assort section
    if has_a:
        title(ws, f'Assort Pack- {dest}')
        rows = []
        for aj, mix in g['assort'].items():
            n = list(g['acount'][aj])[0]
            mixt = tuple((s, mix[s]) for s in SIZES if s in mix)
            gross = ASSORT_GROSS_KNOWN.get((style, mixt))
            est = False
            if gross is None:
                gross = est_assort_gross(style, mixt); est = gross is not None
            rows.append(dict(qty=n, typ='B', style=style, color=color, jan=aj, sizes=dict(mix),
                             perblst=f'{sum(mix.values())} Pcs X 1 Ball', gross=gross, est=est))
        r, tot_rows['assort'] = section(ws, 8 if False else r, 'AssortⅠ' if len(g['assort']) == 1 else 'Assort JAN', rows, ASSORT_TARE)
    # ---- solid section
    if has_s:
        if has_a:
            for rr, (t, sz) in zip((r, r + 1), [('PACKING LIST', 10), (f'Solid Pack- {dest}', 10)]):
                ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=21)
                c = ws.cell(rr, 1, t); c.font = F(sz, True); c.alignment = Alignment(horizontal='center')
            r += 3
        else:
            title(ws, f'Solid Pack- {dest}')
        known = SOLID_GROSS_KNOWN.get(style)
        un = unit_net(style)
        rows = []
        for s in SIZES:
            if s not in g['solid']:
                continue
            jan, q = g['solid'][s]
            if not q:
                continue
            full, rem = divmod(q, CARTON_PCS)
            for qty, pcs in ((full, CARTON_PCS), (1 if rem else 0, rem)):
                if not qty:
                    continue
                if known and s in known and pcs == CARTON_PCS:
                    gross, est = known[s], False
                elif un:
                    gross, est = round(un[s] * pcs + SOLID_TARE, 2), True
                else:
                    gross, est = None, False
                rows.append(dict(qty=qty, typ='A', style=style, color=color, jan=jan, sizes={s: pcs}, perblst=None,
                                 gross=gross, est=est))
        r, tot_rows['solid'] = section(ws, r, 'Solid JAN', rows, SOLID_TARE)
    # ---- summary (size)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
    put(ws, r, 6, 'SUMMARY', F(10, True), GREY)
    for c in range(7, 15):
        put(ws, r, c, None, F(10, True), GREY)
    r += 1
    put(ws, r, 6, 'SIZE', F(10, True), GREY); put(ws, r, 7, None, border=BOX, fill=GREY)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    for s, c in SCOL.items():
        put(ws, r, c, s, F(10, True), GREY)
    put(ws, r, 14, 'TOTAL', F(10, True), GREY)
    ro, rs, rx, rp = r + 1, r + 2, r + 3, r + 4
    for rr, t in ((ro, 'ORDER  QTY'), (rs, 'SHIP  QTY'), (rx, 'EX/SHORT'), (rp, '%')):
        put(ws, rr, 6, t, F(10, True)); put(ws, rr, 7); ws.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=7)
    for s, c in SCOL.items():
        col = L(c)
        put(ws, ro, c, g['order'].get(s) or None, F(10, False, '0000FF'), fmt='#,##0;;')
        parts = [f'{col}{tot_rows[k]}' for k in ('assort', 'solid') if tot_rows[k]]
        put(ws, rs, c, '=' + '+'.join(parts), F(10, True), fmt='#,##0;;')
        put(ws, rx, c, f'={col}{rs}-{col}{ro}', F(10, True), fmt='#,##0;-#,##0;0')
        put(ws, rp, c, f'=IFERROR({col}{rx}/{col}{ro},0)', F(10), fmt='0.0%')
    for rr in (ro, rs, rx):
        put(ws, rr, 14, f'=SUM(H{rr}:M{rr})', F(10, True), fmt='#,##0')
    put(ws, rp, 14, f'=IFERROR(N{rx}/N{ro},0)', F(10, True), fmt='0.0%')
    r = rp + 3
    # ---- summary (totals)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    for c in range(6, 10):
        put(ws, r, c, None, F(10, True), GREY)
    ws.cell(r, 6).value = 'SUMMARY'
    tr = [tot_rows[k] for k in ('assort', 'solid') if tot_rows[k]]
    for i, (lab, col, unit, fmt) in enumerate([('TOTAL CTN QTY', 'C', 'CTN', '#,##0'), ('TTL CTN GROSS WEIGHT', 'S', 'KGS', '#,##0.00'),
                                               ('TTL CTN NET WEIGHT', 'T', 'KGS', '#,##0.00'), ('TOTAL CBM', 'U', 'CBM', '0.000')], 1):
        rr = r + i
        put(ws, rr, 6, lab, F(10, True), align=Alignment(horizontal='left', vertical='center')); put(ws, rr, 7)
        ws.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=7)
        put(ws, rr, 8, '=' + '+'.join(f'{col}{t}' for t in tr), F(10, True), fmt=fmt)
        put(ws, rr, 9, unit, F(10, True))
    # legend
    lr = r + 6
    ws.cell(lr, 6, 'Blue = input (from PO / reference); yellow = weight is an estimate or still to be entered (replace with actual weighed carton weight).').font = F(8)
    # widths / print
    for c, w in {1: 7.6, 2: 7.6, 3: 7.6, 4: 7.6, 5: 13.6, 6: 14.5, 7: 15.1, 14: 14.5, 15: 9, 16: 9, 17: 9, 18: 9, 19: 10.5, 20: 10.5, 21: 9.5}.items():
        ws.column_dimensions[L(c)].width = w
    for c in range(8, 14):
        ws.column_dimensions[L(c)].width = 9.5
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    return ws



NOTES = [
 ("RULES USED (derived from PO + the Vietnam and Malaysia reference lists)", True),
 ("1. Carton = 24 pcs. Solid carton = one size only (24 pcs). Assort carton = the PO assort ratio (V x pcs/ball), always 24 pcs; number of assort cartons = PO col W.", False),
 ("2. PO row with assort JAN (cols R-Y) = assort cartons. PO row without assort JAN = solid qty (col AB). PO total (col AC) per size = solid + assort pcs - used as ORDER QTY.", False),
 ("3. Solid qty is split into full 24-pc cartons; any remainder (<24) becomes one partial carton of that size (only USA BE43AA7S has remainders).", False),
 ("4. Carton type A = solid, B = assort. Numbering restarts at 1 in each section (same as the references). CBM = ctn x 55 x 37 x 32 / 1,000,000.", False),
 ("5. Net weight = gross - 1.21 (solid) or gross - 1.5 (assort), as in the references. Style + colour (068 Navy / 069 Dark Navy) + size JAN come straight from the PO.", False),
 ("", False),
 ("CHECKS", True),
 ("PO is internally consistent: for every destination/style/size, solid + assort = PO total, and every assort mix sums to 24 pcs. SHIP = ORDER on every sheet (EX/SHORT = 0).", False),
 ("Vietnam BE43WA7S reproduces the reference carton counts (21 / 23 / 16 / 13 solid cartons of S / M / L / XL = 73 ctns, 1,752 pcs).", False),
 ("Malaysia BE43WA7S reproduces the reference assort block (27 ctns, 8 S / 8 M / 4 L / 4 XL) and solid cartons for M, L, XL (24 / 20 / 14).", False),
 ("", False),
 ("DIFFERENCES FROM THE REFERENCE FILES - please confirm", True),
 ("a) Vietnam reference ships 1,812 pcs vs PO 1,752: three extra loose cartons of 20 pcs (S, M, L). Not derivable from the PO, so not reproduced here.", False),
 ("b) Malaysia reference has 18 solid S cartons (432 pcs) vs PO 408 pcs (17 cartons) = +24 S. Not reproduced here.", False),
 ("c) Malaysia reference uses assort JAN 2810511208578, which is the Thailand assort JAN in the PO. Malaysia's PO assort JAN is 2810711207142 - used here.", False),
 ("d) Reference typos not copied: title 'LTD.3', and TOTAL CBM labelled 'LBS' (shown as CBM here).", False),
 ("", False),
 ("WEIGHTS (yellow cells)", True),
 ("Known from the references: BE43WA7S solid S 11.20 / M 11.57 / L 12.20 / XL 12.71 kg and the Malaysia assort carton 11.50 kg. These are filled in without highlight.", False),
 ("Estimated (yellow): BE43WA7S XS and XXL solid, and assort cartons with a different size mix - scaled from the per-piece weights of the known cartons. Replace with weighed values.", False),
 ("Blank (yellow): all BE43AA7S (flare culottes) sheets - there is no reference weight for that style. Culottes sheets are included because they are in the PO; delete them if Creative Collections does not produce that style.", False),
]


def make_notes(wb):
    ws = wb.create_sheet('Rules & Checks', 0)
    ws.column_dimensions['A'].width = 150
    ws.sheet_view.showGridLines = False
    for i, (t, h) in enumerate(NOTES, 1):
        c = ws.cell(i, 1, t)
        c.font = F(11 if h else 10, h)
        c.alignment = Alignment(wrap_text=True, vertical='top')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('po'); ap.add_argument('out')
    ap.add_argument('--dest', default=''); ap.add_argument('--style', default='')
    a = ap.parse_args()
    rows = read_po(a.po)
    G, issues = build_groups(rows)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    want_d = [d.strip() for d in a.dest.split(',') if d.strip()]
    want_s = [s.strip() for s in a.style.split(',') if s.strip()]
    order = sorted(G, key=lambda k: (k[1] != 'BE43WA7S', k[1], k[0] not in ('Vietnam', 'Malaysia'), k[0]))
    for k in order:
        d, st, cc = k
        if (want_d and d not in want_d) or (want_s and st not in want_s):
            continue
        make_sheet(wb, d, st, cc, G[k], f'{d} {st}')
    make_notes(wb)
    wb.save(a.out)
    for i in issues:
        print('CHECK:', i)
    print('sheets:', wb.sheetnames, '| data issues:', len(issues))


if __name__ == '__main__':
    main()
