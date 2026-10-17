"""
Abercrombie & Fitch (A&F) buyer: PO PDF  ->  packing list (Excel + PDF).

Pure logic, no Streamlit (UI is in anf_ui.py).
The Excel copies fonts / fills / borders from anf_template.xlsx (the Style 355260231-5318 reference sheet).

Packing rule (from the reference): LOOSE pack, one size per carton, `carton std` pcs (20) in a full carton,
the remainder of every size goes into one extra (small) carton.
"""
from __future__ import annotations

import io
import re
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as CL

HERE = Path(__file__).parent
TEMPLATE = HERE / "anf_template.xlsx"

DEF_STD = 20
DEF_NET_PER_PC = 0.60          # kg - placeholder, reference sheet uses 0.59 - 0.89 for adult sizes
DEF_TARE = 1.0                 # kg per carton (reference column "per cnt wt")
DEF_BIG = (64.0, 39.0, 24.0)   # cm, reference "L 64 X W 39 X H 24"
DEF_SMALL = (64.0, 39.0, 13.0)  # cm, reference "L 64 X W 39 X H 13"

T_FIRST_SIZE, T_LAST_SIZE = 5, 43      # template size columns E..AQ
T_BODY0, T_BODY1, T_TOTAL = 11, 66, 67  # template body rows / total row


# --------------------------------------------------------------------------- PO PDF
@dataclass
class PoData:
    po: str
    po_code: str
    tech_pack: str
    description: str
    carton_std: int
    pack_type: str
    ship_date: str
    total_units: int
    brand: str
    lines: pd.DataFrame
    warnings: list[str] = field(default_factory=list)


_LINE = re.compile(
    r"^(?P<color>.+?)\s+(?P<sku>\d+(?:-\d+){3,})\s+(?P<short>\d+)\s+(?P<ean>\d{10,14})\s+"
    r"(?:(?P<shade>[A-Za-z0-9]+)\s+)?(?P<ticket>\d+\.\d{2})\s+[A-Z]{3}\s+(?P<cost>\d+\.\d{2})\s+"
    r"(?P<size>.+?)\s+(?P<qty>\d+)\s*$")


def _page_text(data: bytes) -> list[str]:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def parse_po(data: bytes) -> PoData:
    pages = _page_text(data)
    text = "\n".join(pages)

    def grab(pat, default=""):
        m = re.search(pat, text)
        return m.group(1).strip() if m else default

    rows, seen = [], set()
    for ln in text.splitlines():
        m = _LINE.match(ln.strip())
        if not m:
            continue
        key = m.group("sku")
        if key in seen:                      # header block repeats on every page
            continue
        seen.add(key)
        rows.append({"color": m.group("color").strip(), "sku": key, "short_sku": m.group("short"),
                     "size": re.sub(r"\s+", " ", m.group("size").strip().upper()),
                     "qty": int(m.group("qty")), "cost": float(m.group("cost"))})
    if not rows:
        raise ValueError("No SKU lines found. Is this an A&F 'PO SKU Summary' PDF?")
    lines = pd.DataFrame(rows)

    po = PoData(
        po=grab(r"Purchase Order:\s*(\d+)"),
        po_code=grab(r"PO Code:\s*(\S+)"),
        tech_pack=grab(r"Tech Pack:\s*(\S+)"),
        description=grab(r"PO Description:\s*(.+)"),
        carton_std=int(grab(r"Carton Std:\s*(\d+)", str(DEF_STD))),
        pack_type=grab(r"Pack Type:\s*(\S+)"),
        ship_date=grab(r"Ship Date:\s*([\d/]+)"),
        total_units=int(grab(r"Total Ordered Units:\s*(\d+)", "0") or 0),
        brand=grab(r"Brand:\s*(.+?)\s+Seconds"),
        lines=lines,
    )
    got = int(lines["qty"].sum())
    if po.total_units and got != po.total_units:
        po.warnings.append(f"Parsed {got:,} pcs but the PO header says {po.total_units:,} - check the SKU lines.")
    if po.pack_type and po.pack_type.upper() != "LOOSE":
        po.warnings.append(f"Pack type is {po.pack_type}, not LOOSE - this layout packs one size per carton.")
    return po


def size_key(label: str):
    """5/6 SHORT < 5/6 < 5/6 LONG < 7/8 SHORT ...  (matches the S-R-L order of the reference)."""
    m = re.match(r"(\d+)(?:/(\d+))?\s*(.*)", label)
    var = {"SHORT": 0, "": 1, "REG": 1, "REGULAR": 1, "LONG": 2}
    if m:
        return (0, int(m.group(1)), var.get(m.group(3).strip(), 3), label)
    return (1, 0, 0, label)


def ordered_sizes(lines: pd.DataFrame) -> list[str]:
    return sorted(lines["size"].unique(), key=size_key)


# --------------------------------------------------------------------------- plan
@dataclass
class Plan:
    meta: dict
    sizes: list[str]
    colors: list[str]
    order: dict            # (color, size) -> order qty
    ship: dict             # (color, size) -> ship qty
    rows: list[dict]       # carton rows
    summary: dict
    wt: dict
    tare: float
    std: int
    big: tuple
    small: tuple


def make_plan(po: PoData, ship_override: dict | None = None, wt: dict | None = None, tare: float = DEF_TARE,
              std: int | None = None, big=DEF_BIG, small=DEF_SMALL, style: str | None = None,
              description: str | None = None) -> Plan:
    std = int(std or po.carton_std or DEF_STD)
    ship_override = ship_override or {}
    sizes = ordered_sizes(po.lines)
    colors = list(dict.fromkeys(po.lines["color"]))
    wt = {s: float((wt or {}).get(s, DEF_NET_PER_PC)) for s in sizes}
    order = {(r.color, r.size): int(r.qty) for r in po.lines.itertuples()}
    ship = {k: int(ship_override.get(k, v)) if ship_override.get(k) is not None else v for k, v in order.items()}

    rows = []
    for c in colors:
        for s in sizes:
            q = ship.get((c, s), 0)
            if q <= 0:
                continue
            full, rem = divmod(q, std)
            if full:
                rows.append({"color": c, "size": s, "qty": std, "ctns": full, "small": False})
            if rem:
                rows.append({"color": c, "size": s, "qty": rem, "ctns": 1, "small": True})
    nxt = 1
    for r in rows:
        r["from"], r["to"] = nxt, nxt + r["ctns"] - 1
        nxt = r["to"] + 1
        r["pcs"] = r["qty"] * r["ctns"]
        r["wpp"] = wt[r["size"]]
        r["net"] = r["wpp"] * r["qty"]
        r["gross"] = r["net"] + tare
    total_ctns = sum(r["ctns"] for r in rows)
    n_small = sum(r["ctns"] for r in rows if r["small"])
    n_big = total_ctns - n_small
    tot_o, tot_s = sum(order.values()), sum(ship.values())
    summary = {
        "order": tot_o, "ship": tot_s, "excess": tot_s - tot_o,
        "pct": (tot_s - tot_o) / tot_o * 100 if tot_o else 0.0,
        "ctns": total_ctns, "big": n_big, "small": n_small,
        "gross": sum(r["gross"] * r["ctns"] for r in rows),
        "net": sum(r["net"] * r["ctns"] for r in rows),
        "cbm": (n_big * big[0] * big[1] * big[2] + n_small * small[0] * small[1] * small[2]) / 1_000_000,
    }
    meta = {"buyer": "ABERCROMBIE", "style": style or po.tech_pack, "po": po.po,
            "description": description or po.po_code, "po_desc": po.description, "ship_date": po.ship_date}
    return Plan(meta, sizes, colors, order, ship, rows, summary, wt, tare, std, tuple(big), tuple(small))


def file_stem(plan: Plan) -> str:
    color = plan.colors[0] if len(plan.colors) == 1 else "MULTI"
    return f"Style__{plan.meta['style']}-{color}-PO__{plan.meta['po']}_-QTY__{plan.summary['ship']}_Pcs__{plan.std}_"


# --------------------------------------------------------------------------- previews
def main_df(plan: Plan) -> pd.DataFrame:
    out = []
    for r in plan.rows:
        d = {"Case from": r["from"], "Case to": r["to"], "Color": r["color"]}
        for s in plan.sizes:
            d[s] = r["qty"] if s == r["size"] else None
        d.update({"Qty/Ctn": r["qty"], "Total Ctn": r["ctns"], "Total Pcs": r["pcs"],
                  "Gr.wt/ctn": round(r["gross"], 2), "Net wt/ctn": round(r["net"], 2),
                  "Total Gr.wt": round(r["gross"] * r["ctns"], 2), "Total Net wt": round(r["net"] * r["ctns"], 2),
                  "Pc wt (kg)": r["wpp"], "Ctn wt (kg)": plan.tare})
        out.append(d)
    df = pd.DataFrame(out)
    s = plan.summary
    tot = {c: None for c in df.columns}
    tot.update({"Color": "TOTAL", "Total Ctn": s["ctns"], "Total Pcs": s["ship"],
                "Total Gr.wt": round(s["gross"], 2), "Total Net wt": round(s["net"], 2)})
    return pd.concat([df, pd.DataFrame([tot])], ignore_index=True)


def breakdown_df(plan: Plan, color: str) -> pd.DataFrame:
    sizes = [s for s in plan.sizes if (color, s) in plan.order]
    o = [plan.order[(color, s)] for s in sizes]
    sh = [plan.ship[(color, s)] for s in sizes]
    ex = [a - b for a, b in zip(sh, o)]
    pc = [round(e / x * 100, 2) if x else 0 for e, x in zip(ex, o)]
    df = pd.DataFrame([o, sh, ex, pc], columns=sizes, index=["ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE %"])
    df["TOTAL"] = [sum(o), sum(sh), sum(ex), round(sum(ex) / sum(o) * 100, 2) if sum(o) else 0]
    return df


def summary_df(plan: Plan) -> pd.DataFrame:
    s = plan.summary
    return pd.DataFrame([
        ("Order qty", f"{s['order']:,} pcs"), ("Pack qty (ship)", f"{s['ship']:,} pcs"),
        ("Excess / short qty", f"{s['excess']:+,} pcs"), ("Percentage", f"{s['pct']:.2f}%"),
        ("Total cartons", f"{s['ctns']:,}  (big {s['big']:,} + small {s['small']:,})"),
        ("Total gross weight", f"{s['gross']:,.2f} kg"), ("Total net weight", f"{s['net']:,.2f} kg"),
        ("Volume", f"{s['cbm']:,.3f} CBM"),
    ], columns=["Item", "Value"])


# --------------------------------------------------------------------------- Excel
def _cps(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def build_xlsx(plan: Plan) -> bytes:
    tpl = openpyxl.load_workbook(TEMPLATE)
    ts = tpl.active
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ts.title
    wb.loaded_theme = tpl.loaded_theme

    n, R, Bk = len(plan.sizes), len(plan.rows), len(plan.colors)
    dc = n - (T_LAST_SIZE - T_FIRST_SIZE + 1)         # column shift for everything right of the size block
    d = R - (T_BODY1 - T_BODY0 + 1)                   # row shift after the body
    stride = 6                                        # rows per colour block in the summary
    ex = stride * (Bk - 1)

    SZ0 = T_FIRST_SIZE                                # first size col (E)
    SZ1 = SZ0 + n - 1
    cAR = 44 + dc
    cAS, cAT, cAV, cAW, cAX, cAY, cAZ, cBA = cAR + 1, cAR + 2, cAR + 4, cAR + 5, cAR + 6, cAR + 7, cAR + 8, cAR + 9
    first, last = T_BODY0, T_BODY0 + R - 1
    tot = last + 1
    r69 = 69 + d
    blk0 = 71 + d                                     # first summary block header row

    def role(k):
        if k == 0:
            return T_FIRST_SIZE
        if k == n - 1:
            return T_LAST_SIZE
        return T_FIRST_SIZE + 1

    def colmap(c):
        if c < T_FIRST_SIZE:
            return [(c, c)]
        if c > T_LAST_SIZE:
            return [(c + dc, c)]
        if c == T_FIRST_SIZE:
            return [(SZ0 + k, role(k)) for k in range(n)]
        return []

    def newrow(r):
        if r < T_BODY0:
            return r
        if r <= T_TOTAL:
            return r + d if r >= T_TOTAL else None
        if r <= 75:
            return r + d
        return r + d + ex

    # -- header block (rows 1-10)
    def copy_row(tr, nr, values=True):
        for c in range(1, 54):
            for k, (nc, sc) in enumerate(colmap(c)):
                s, dd = ts.cell(tr, sc), ws.cell(nr, nc)
                _cps(s, dd)
                v = s.value
                if values and v is not None and not (isinstance(v, str) and v.startswith("=")):
                    if c != T_FIRST_SIZE or k == 0:
                        dd.value = v

    for r in range(1, 11):
        copy_row(r, r)
    # -- body
    for i in range(R):
        copy_row(T_BODY0 if i == 0 else T_BODY0 + 1, T_BODY0 + i, values=False)
    # -- total row, rows between total and summary
    copy_row(T_TOTAL, tot)
    for r in (68, 69, 70):
        copy_row(r, r + d)
    # -- summary blocks
    for b in range(Bk):
        for tr in range(71, 76):
            copy_row(tr, tr + d + stride * b)
        if b:                                           # spacer between blocks
            copy_row(70, 70 + d + stride * b)
    # -- footer
    for r in range(76, 83):
        copy_row(r, r + d + ex)

    # ---------------------------------------------------------------- header values
    m = plan.meta
    ws["C4"], ws["C5"], ws["C6"], ws["C7"] = m["buyer"], str(m["style"]), str(m["po"]), m["description"]

    # ---------------------------------------------------------------- size headers
    long_label = max((len(s) for s in plan.sizes), default=3)
    rot = Alignment(horizontal="center", vertical="center", textRotation=90)
    for k, s in enumerate(plan.sizes):
        for rr in (10, blk0 + stride * 0):
            pass
        ws.cell(10, SZ0 + k).value = s
        ws.cell(10, SZ0 + k).alignment = rot
    ws.row_dimensions[10].height = max(30, 7.5 * long_label + 12)

    # ---------------------------------------------------------------- body rows
    col_of = {s: SZ0 + k for k, s in enumerate(plan.sizes)}
    a_, z_ = CL(SZ0), CL(SZ1)
    for i, r in enumerate(plan.rows):
        nr = T_BODY0 + i
        ws[f"A{nr}"] = 1 if i == 0 else f"=C{nr - 1}+1"
        ws[f"B{nr}"] = "-"
        ws[f"C{nr}"] = f"=A{nr}-1+{CL(cAS)}{nr}"
        ws[f"D{nr}"] = r["color"]
        ws.cell(nr, col_of[r["size"]]).value = r["qty"]
        ws.cell(nr, cAR).value = f"=SUM({a_}{nr}:{z_}{nr})"
        ws.cell(nr, cAS).value = r["ctns"]
        ws.cell(nr, cAT).value = f"={CL(cAS)}{nr}*{CL(cAR)}{nr}"
        ws.cell(nr, cAV).value = f"={CL(cAW)}{nr}+{CL(cBA)}{nr}"
        ws.cell(nr, cAW).value = f"={CL(cAZ)}{nr}*{CL(cAR)}{nr}"
        ws.cell(nr, cAX).value = f"={CL(cAV)}{nr}*{CL(cAS)}{nr}"
        ws.cell(nr, cAY).value = f"={CL(cAW)}{nr}*{CL(cAS)}{nr}"
        ws.cell(nr, cAZ).value = r["wpp"]
        ws.cell(nr, cBA).value = plan.tare
        ws.row_dimensions[nr].height = ts.row_dimensions[T_BODY0].height or 18.75

    # ---------------------------------------------------------------- total row
    ws[f"A{tot}"] = "TOTAL"
    for c in (cAS, cAT, cAX, cAY):
        ws.cell(tot, c).value = f"=SUM({CL(c)}{first}:{CL(c)}{last})"
    for c in (cAV, cAW):
        ws.cell(tot, c).value = None

    # ---------------------------------------------------------------- summary blocks
    ws.cell(r69, 1).value = "SUMMARY : COLOR & SIZEWISE BREAKDOWN"
    ws.cell(r69, cAX).value = f"={CL(cAX)}{tot}"
    ws.cell(r69, cAY).value = f"={CL(cAY)}{tot}"
    rows_by = {}
    for i, r in enumerate(plan.rows):
        rows_by.setdefault((r["color"], r["size"]), []).append(T_BODY0 + i)
    order_cells, ship_cells = [], []
    for b, color in enumerate(plan.colors):
        r71 = blk0 + stride * b
        ro, rs, rx, rp = r71 + 1, r71 + 2, r71 + 3, r71 + 4
        ws.cell(r71, 3).value = "SIZE" if Bk == 1 else f"SIZE  ({color})"
        ws.cell(r71, cAR).value = "TOTAL "
        ws.cell(ro, 3).value, ws.cell(rs, 3).value = "ORDER QTY", "SHIP QTY"
        ws.cell(rx, 3).value, ws.cell(rp, 3).value = "EXS/SHT QTY", "PERCENTAGE"
        for k, s in enumerate(plan.sizes):
            c = SZ0 + k
            ws.cell(r71, c).value = s
            ws.cell(r71, c).alignment = rot
            if (color, s) not in plan.order:
                continue
            ws.cell(ro, c).value = plan.order[(color, s)]
            cells = [f"{CL(cAT)}{x}" for x in rows_by.get((color, s), [])]
            ws.cell(rs, c).value = "=" + ("+".join(cells) if cells else "0")
            ws.cell(rx, c).value = f"={CL(c)}{rs}-{CL(c)}{ro}"
            ws.cell(rp, c).value = f"={CL(c)}{rx}/{CL(c)}{ro}*100"
        for rr in (ro, rs, rx):
            ws.cell(rr, cAR).value = f"=SUM({a_}{rr}:{z_}{rr})"
            ws.cell(rr, cAS).value = "PCS"
        ws.cell(rp, cAR).value = f"={CL(cAR)}{rx}/{CL(cAR)}{ro}*100"
        ws.cell(rp, cAS).value = "%"
        ws.row_dimensions[r71].height = max(21.75, 7.5 * long_label + 12)
        for rr in (ro, rs, rx, rp):
            ws.row_dimensions[rr].height = 21.75
        order_cells.append(f"{CL(cAR)}{ro}")
        ship_cells.append(f"{CL(cAR)}{rs}")

    # top-right block (rows 4-7)
    lab = max(6, cAR - 10)
    for off, (txt, formula, unit) in enumerate([
            ("ORDER QTY", "=" + "+".join(order_cells), "PCS"),
            ("PACK QTY", "=" + "+".join(ship_cells), "PCS"),
            ("EXCESS/SHORT QTY", f"={CL(cAS)}5-{CL(cAS)}4", "PCS"),
            ("PERCENTAGE", f"={CL(cAS)}6/{CL(cAS)}4*100", "%")]):
        r = 4 + off
        src_lab = ts.cell(r, 22)
        _cps(src_lab, ws.cell(r, lab))
        ws.cell(r, lab).value = txt
        for c in range(lab + 1, cAR):                  # keep the box continuous
            _cps(ts.cell(r, 23), ws.cell(r, c))
        for src_c, dst_c, val in ((44, cAR, ":"), (45, cAS, formula), (46, cAT, unit)):
            _cps(ts.cell(r, src_c), ws.cell(r, dst_c))
            ws.cell(r, dst_c).value = val
        ws.cell(r, cAS).number_format = "0.00" if txt == "PERCENTAGE" else "General"

    # ---------------------------------------------------------------- footer
    f0 = 77 + d + ex
    L_, W_, Hb = plan.big
    _, _, Hs = plan.small
    ws.cell(f0, 1).value = f"CARTON MEASUREMENT : L {L_:g} X W {W_:g} X H {Hb:g} CM"
    ws.cell(f0 + 1, 1).value = f"CARTON MEASUREMENT : L {plan.small[0]:g} X W {plan.small[1]:g} X H {Hs:g} CM"
    cbm_c = min(13, SZ1)
    ws.cell(f0, cbm_c).value = (f"=({L_:g}*{W_:g}*{Hb:g}*{CL(cAR)}{f0 + 1}"
                                f"+{plan.small[0]:g}*{plan.small[1]:g}*{Hs:g}*{CL(cAR)}{f0 + 2})/1000000")
    ws.cell(f0, cbm_c).number_format = '0.000" CBM"'
    if SZ1 >= cbm_c + 2:
        ws.merge_cells(start_row=f0, start_column=cbm_c, end_row=f0, end_column=cbm_c + 2)
    labc = cAR - 4
    ws.cell(f0 + 1, labc).value = "Total  Carton BIG= "
    ws.cell(f0 + 1, cAR).value = f"={CL(cAS)}{tot}-{CL(cAR)}{f0 + 2}"
    ws.cell(f0 + 2, labc).value = "Total  Carton Small="
    ws.cell(f0 + 2, cAR).value = plan.summary["small"]
    for rr in (f0 + 1, f0 + 2):
        for c in (labc, cAR):
            _cps(ts.cell(78, 40 if c == labc else 44), ws.cell(rr, c))
    ws.cell(f0 + 4, 1).value = "T. GROSS WEIGHT :"
    ws.cell(f0 + 4, 4).value = f"={CL(cAX)}{r69}"
    ws.cell(f0 + 4, 5).value = "KGS"
    ws.cell(f0 + 5, 1).value = "T. NET      WEIGHT :"
    ws.cell(f0 + 5, 4).value = f"={CL(cAY)}{r69}"
    ws.cell(f0 + 5, 5).value = "KGS"
    for rr in (f0 + 4, f0 + 5):
        ws.cell(rr, 4).number_format = "0.00"
        for c, sc in ((1, 1), (4, 4), (5, 5)):
            _cps(ts.cell(81 if rr == f0 + 4 else 82, sc), ws.cell(rr, c))

    # ---------------------------------------------------------------- merges
    def mc(c, end):
        if c < T_FIRST_SIZE:
            return c
        if c <= T_LAST_SIZE:
            return SZ1 if (end and c == T_LAST_SIZE) else min(c, SZ1)
        return c + dc

    for rng in ts.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        if r1 in range(T_BODY0, T_TOTAL):
            continue
        if r1 == T_TOTAL or 68 <= r1 <= 70:
            nr1 = nr2 = r1 + d
            blocks = [0]
        elif 71 <= r1 <= 75:
            nr1, nr2 = r1 + d, r2 + d
            blocks = range(Bk)
        elif r1 >= 76:
            nr1, nr2 = r1 + d + ex, r2 + d + ex
            blocks = [0]
        else:
            nr1, nr2 = r1, r2
            blocks = [0]
        for b in blocks:
            a = (nr1 + stride * b, mc(c1, False))
            z = (nr2 + stride * b, mc(c2, True))
            if a != z:
                ws.merge_cells(start_row=a[0], start_column=a[1], end_row=z[0], end_column=z[1])

    # ---------------------------------------------------------------- widths / sheet setup
    def tw(c):
        for d_ in ts.column_dimensions.values():
            if d_.min and d_.max and d_.min <= c <= d_.max and d_.width:
                return d_.width
        return None

    ws.sheet_format.defaultColWidth = ts.sheet_format.defaultColWidth or 9.14
    for j in range(1, 54 + dc):
        if j < SZ0:
            w = tw(j)
        elif j <= SZ1:
            w = 7.0 if long_label > 6 else 6.5
        else:
            w = tw(j - dc)
        if w:
            ws.column_dimensions[CL(j)].width = w
    for r in list(range(1, 10)) + [tot, tot + 1, r69, r69 + 1]:
        h = ts.row_dimensions[min(r, 82)].height
        if h and r not in (10,):
            ws.row_dimensions[r].height = h

    ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ts.page_setup.paperSize or 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_area = f"A1:{CL(51 + dc)}{f0 + 5}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- PDF
def build_pdf(plan: Plan) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    n = len(plan.sizes)
    m = plan.meta
    fixed = [20, 22] + [13.5] * n + [12, 12, 14, 14, 14, 16, 16]
    scale = 400.0 / sum(fixed)
    widths = [w * scale * mm for w in fixed]
    fs = max(4.8, min(7.5, 7.5 * scale))
    sty = ParagraphStyle("c", fontName="Helvetica", fontSize=fs, leading=fs + 1.5, alignment=1)
    stb = ParagraphStyle("b", parent=sty, fontName="Helvetica-Bold")
    lab = ParagraphStyle("l", fontName="Helvetica", fontSize=8.5, leading=11)
    big = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=18, leading=24, alignment=1, spaceAfter=2)
    mid = ParagraphStyle("m", fontName="Helvetica-Bold", fontSize=13, leading=17, alignment=1, spaceAfter=6)
    P_ = lambda t, s=sty: Paragraph(str(t), s)
    wrap = lambda s: s.replace(" ", "<br/>") if " " in s else s

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A3), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=f"Packing list PO {m['po']}")
    S = plan.summary
    story = [Paragraph("CREATIVE COLLECTIONS LTD.U-1B", big), Paragraph("PACKING LIST DETAILS", mid)]
    info = Table([[Paragraph(f"<b>BUYER</b> : {m['buyer']}<br/><b>STYLE</b> : {m['style']}<br/><b>P.O NO</b> : {m['po']}<br/>"
                             f"<b>DESCRIPTION</b> : {m['description']}", lab),
                   Paragraph(f"<b>ORDER QTY</b> : {S['order']:,} PCS<br/><b>PACK QTY</b> : {S['ship']:,} PCS<br/>"
                             f"<b>EXCESS/SHORT QTY</b> : {S['excess']:+,} PCS<br/><b>PERCENTAGE</b> : {S['pct']:.2f} %", lab)]],
                 colWidths=[sum(widths) / 2] * 2)
    info.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.7, colors.black), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.grey)]))
    story += [info, Spacer(1, 4 * mm)]

    h1 = ["CASE LABEL NO.", "COLOR"] + ["SIZE"] + [""] * (n - 1) + ["QTY/<br/>CTN", "TOTAL<br/>CTN", "TOTAL<br/>PCS",
                                                                     "Grs.wt<br/>per ctn", "Net.wt<br/>per ctn",
                                                                     "Total<br/>Grs.wt", "Total<br/>net.wt"]
    h2 = ["", ""] + [wrap(s) for s in plan.sizes] + [""] * 7
    data = [[P_(c, stb) for c in h1], [P_(c, stb) for c in h2]]
    for r in plan.rows:
        row = [P_(f"{r['from']} - {r['to']}"), P_(r["color"])]
        row += [P_(r["qty"] if s == r["size"] else "") for s in plan.sizes]
        row += [P_(r["qty"]), P_(r["ctns"], stb), P_(r["pcs"]), P_(f"{r['gross']:.2f}"), P_(f"{r['net']:.2f}"),
                P_(f"{r['gross'] * r['ctns']:,.2f}"), P_(f"{r['net'] * r['ctns']:,.2f}")]
        data.append(row)
    tot = [P_("TOTAL", stb)] + [P_("")] * (n + 8)
    tot[n + 3] = P_(f"{S['ctns']:,}", stb)
    tot[n + 4] = P_(f"{S['ship']:,}", stb)
    tot[n + 7] = P_(f"{S['gross']:,.2f}", stb)
    tot[n + 8] = P_(f"{S['net']:,.2f}", stb)
    data.append(tot)
    t = Table(data, colWidths=widths, repeatRows=2)
    last = len(fixed) - 1
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (n + 3 - 0, 2), (n + 3, len(data) - 2), colors.HexColor("#FFFF99")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("SPAN", (2, 0), (1 + n, 0)),
        *[("SPAN", (c, 0), (c, 1)) for c in [0, 1] + list(range(2 + n, last + 1))],
        ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    story += [t, Spacer(1, 6 * mm), Paragraph("SUMMARY : COLOR &amp; SIZEWISE BREAKDOWN", ParagraphStyle(
        "s", fontName="Helvetica-Bold", fontSize=10, spaceAfter=3))]

    for color in plan.colors:
        bd = breakdown_df(plan, color)
        hdr = [P_(color if len(plan.colors) > 1 else "SIZE", stb)] + [P_(wrap(c), stb) for c in bd.columns]
        body = [[P_(i, stb)] + [P_(f"{v:g}" if isinstance(v, float) else v) for v in row]
                for i, row in zip(bd.index, bd.values.tolist())]
        w0 = 24 * mm
        wc = min(14 * mm, (sum(widths) - w0) / len(bd.columns))
        bt = Table([hdr] + body, colWidths=[w0] + [wc] * len(bd.columns), hAlign="LEFT")
        bt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
                                ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story += [bt, Spacer(1, 3 * mm)]

    story += [Spacer(1, 2 * mm), Paragraph(
        f"CARTON MEASUREMENT : L {plan.big[0]:g} X W {plan.big[1]:g} X H {plan.big[2]:g} CM &nbsp;&nbsp;&nbsp; BIG = {S['big']:,}<br/>"
        f"CARTON MEASUREMENT : L {plan.small[0]:g} X W {plan.small[1]:g} X H {plan.small[2]:g} CM &nbsp;&nbsp;&nbsp; SMALL = {S['small']:,}"
        f" &nbsp;&nbsp;&nbsp; VOLUME = {S['cbm']:,.3f} CBM<br/><b>T. GROSS WEIGHT :</b> {S['gross']:,.2f} KGS<br/>"
        f"<b>T. NET WEIGHT :</b> {S['net']:,.2f} KGS", lab)]
    doc.build(story)
    return buf.getvalue()
