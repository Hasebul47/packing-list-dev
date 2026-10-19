"""
Haddad buyer: PO status sheet (XLS/XLSX)  ->  packing list (Excel + PDF).

Pure logic, no Streamlit. The UI lives in haddad_ui.py.
The Excel output copies fonts / fills / borders / logo from haddad_template.xlsx
(the LVB-WEU213 reference sheet), so it looks the same as the reference.
"""
from __future__ import annotations

import io
import re
import datetime as dt
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter as CL

HERE = Path(__file__).parent
TEMPLATE = HERE / "haddad_template.xlsx"

LBS = 2.2046226218          # kg -> lbs
DEF_NET_PER_PC = 0.25625    # kg/pc  -> (7.00 gross - 0.85 tare) / 24 pcs, the reference carton
DEF_TARE = 0.85             # kg per carton
DEF_DIMS = (20.0, 13.0, 6.5)  # carton L x W x H, inch (reference)

NEEDED = ["Ref#", "Style", "Style Name", "Color", "Fashion Color", "Label", "PO#", "PO Qty",
          "Status", "Ship Date", "Size Configuration", "Pack Ratio", "Master Box Quantity"]


# --------------------------------------------------------------------------- input
def load_status(data: bytes, name: str = "status.xls") -> pd.DataFrame:
    """Read the Haddad PO status sheet. Header row is detected (it is row 2 in the sample)."""
    bio = io.BytesIO(data)
    engine = "xlrd" if name.lower().endswith(".xls") else None
    raw = pd.read_excel(bio, header=None, dtype=str, engine=engine)
    hdr = None
    for i in range(min(15, len(raw))):
        vals = {str(v).strip() for v in raw.iloc[i].tolist()}
        if "Ref#" in vals and "PO#" in vals:
            hdr = i
            break
    if hdr is None:
        raise ValueError("Could not find the header row (a row containing 'Ref#' and 'PO#').")
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(c).strip() for c in raw.iloc[hdr]]
    df = df.loc[:, ~df.columns.duplicated()]          # '%', 'Duty Rate [%]' ... repeat 6x: keep first
    missing = [c for c in NEEDED if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns in the status sheet: {missing}")
    df = df[NEEDED].copy()
    for c in NEEDED:
        df[c] = df[c].fillna("").astype(str).str.strip()
    df = df[(df["PO#"] != "") & (df["Ref#"] != "")]
    df["PO Qty"] = pd.to_numeric(df["PO Qty"], errors="coerce")
    df["Master Box Quantity"] = pd.to_numeric(df["Master Box Quantity"], errors="coerce")
    bad = df["PO Qty"].isna() | df["Master Box Quantity"].isna() | (df["Master Box Quantity"] <= 0)
    df.attrs["dropped"] = int(bad.sum())
    df = df[~bad].reset_index(drop=True)
    df["ShipDT"] = pd.to_datetime(df["Ship Date"], format="%m/%d/%y", errors="coerce")
    return df


def _split(s: str) -> list[str]:
    return [x for x in re.split(r"\s*-\s*", str(s).strip()) if x]


def ref_table(df: pd.DataFrame) -> pd.DataFrame:
    """Table 1: one row per Ref# with its order information."""
    def uniq(s, limit=4):
        u = list(dict.fromkeys(x for x in s if x))
        return ", ".join(u[:limit]) + (" …" if len(u) > limit else "")

    rows = []
    for ref, g in df.groupby("Ref#", sort=True):
        cfg = g["Size Configuration"].unique()
        rat = g["Pack Ratio"].unique()
        mbq = g["Master Box Quantity"].unique()
        dates = g.sort_values("ShipDT")["Ship Date"].tolist()
        rows.append({
            "Ref#": ref,
            "Style": g["Style"].iloc[0],
            "Style name": g["Style Name"].iloc[0],
            "POs": len(g),
            "Total PO qty": int(g["PO Qty"].sum()),
            "Ship date": uniq(dates, 3),
            "Label": uniq(g["Label"]),
            "Size config": cfg[0] if len(cfg) == 1 else f"{len(cfg)} sets",
            "Pack ratio": rat[0] if len(rat) == 1 else "mixed",
            "Box qty": int(mbq[0]) if len(mbq) == 1 else "mixed",
        })
    return pd.DataFrame(rows)


def select_lines(df: pd.DataFrame, ref: str, ship_dates=None, include_bk: bool = True) -> pd.DataFrame:
    g = df[df["Ref#"] == ref]
    if ship_dates:
        g = g[g["Ship Date"].isin(ship_dates)]
    if not include_bk:
        g = g[g["Status"].str.upper() != "PO BK"]
    return g.sort_values(["ShipDT", "PO#"], kind="stable").reset_index(drop=True)


# --------------------------------------------------------------------------- sizes
_LETTERS = ["XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL"]


def size_key(s: str):
    u = s.upper()
    if u in _LETTERS:
        return (0, _LETTERS.index(u), "")
    m = re.match(r"(\d+)(.*)", u)
    if m:
        return (1, int(m.group(1)), m.group(2))
    return (2, 0, u)


def size_union(lines: pd.DataFrame) -> list[str]:
    seqs = [_split(c) for c in lines["Size Configuration"]]
    seen: list[str] = []
    for sq in seqs:
        for s in sq:
            if s not in seen:
                seen.append(s)
    if all(len(sq) == 1 for sq in seqs):
        seen.sort(key=size_key)
    return seen


def pack_pcs(cfg: str, ratio: str, mbq: float, po: str = "") -> dict[str, int]:
    """Pieces of each size in ONE carton = pack ratio scaled up to the master box quantity."""
    sizes = _split(cfg)
    r = [float(x) for x in _split(ratio)]
    if len(sizes) != len(r) or sum(r) <= 0:
        raise ValueError(f"PO {po}: size configuration '{cfg}' and pack ratio '{ratio}' do not match.")
    scale = mbq / sum(r)
    return {s: int(round(ri * scale)) for s, ri in zip(sizes, r)}


# --------------------------------------------------------------------------- calculation
@dataclass
class Result:
    ref: str
    sizes: list[str]
    rows: list[dict]
    summary: dict
    tare: float
    dims: tuple
    restart: bool = False
    warnings: list[str] = field(default_factory=list)


def compute(lines: pd.DataFrame, sizes: list[str], excess: dict, override: dict,
            wt: dict, tare: float, dims: tuple, restart: bool = False) -> Result:
    """
    excess  : PO# -> excess % (ship qty is X% above order qty)
    override: PO# -> order qty typed by the user (wins over the % when > 0)
    wt      : size -> net kg per piece
    Ship qty is always the status sheet PO Qty (cartons = PO Qty / master box qty).
    """
    rows, warns = [], []
    nxt = 1
    for _, r in lines.iterrows():
        po = r["PO#"]
        mbq = int(r["Master Box Quantity"])
        ctns = int(round(r["PO Qty"] / mbq))
        ship = ctns * mbq
        if abs(ship - r["PO Qty"]) > 0.5:
            warns.append(f"PO {po}: qty {r['PO Qty']:.0f} is not a multiple of box qty {mbq}.")
        pcs = pack_pcs(r["Size Configuration"], r["Pack Ratio"], mbq, po)
        ov = override.get(po)
        if ov and float(ov) > 0:
            order = int(round(float(ov)))
        else:
            ex = float(excess.get(po) or 0.0)
            order = int(round(ship / (1 + ex / 100.0)))
        short = ship - order
        net_kg = sum(n * float(wt.get(s, 0.0)) for s, n in pcs.items())
        gross_kg = net_kg + tare
        c_from = 1 if restart else nxt
        c_to = c_from + ctns - 1
        nxt = c_to + 1
        rows.append({
            "po": po, "style": r["Style"], "label": r["Label"], "ref": r["Ref#"],
            "color": f"{r['Fashion Color']} - {r['Color']}".strip(" -"),
            "ctns": ctns, "order": order, "ship": ship, "short": short,
            "short_pct": (short / order) if order else 0.0,
            "pcs": pcs, "ppk": sum(pcs.values()),
            "net_kg": net_kg, "gross_kg": gross_kg,
            "c_from": c_from, "c_to": c_to, "dims": dims,
            "cbm": ctns * dims[0] * dims[1] * dims[2] * 2.54 ** 3 / 1_000_000,
        })
    tot_order = sum(x["order"] for x in rows)
    tot_ship = sum(x["ship"] for x in rows)
    gk = sum(x["gross_kg"] * x["ctns"] for x in rows)
    nk = sum(x["net_kg"] * x["ctns"] for x in rows)
    summary = {
        "total_order": tot_order, "total_ship": tot_ship,
        "excess_qty": tot_ship - tot_order,
        "excess_pct": (tot_ship - tot_order) / tot_order if tot_order else 0.0,
        "ctns": sum(x["ctns"] for x in rows),
        "gross_kg": gk, "gross_lbs": gk * LBS, "net_kg": nk, "net_lbs": nk * LBS,
        "cbm": sum(x["cbm"] for x in rows),
    }
    ref = rows[0]["ref"] if rows else ""
    return Result(ref, sizes, rows, summary, tare, dims, restart, warns)


def compute_rows(lines: pd.DataFrame, sizes: list[str], rows_in: list[dict], wt: dict, tare: float,
                 restart: bool = False) -> tuple[Result, list[str]]:
    """Carton table edited by the user -> Result.
    rows_in keys: po, ctns, pcs (dict size->pcs per carton), dims (L, W, H inches), order (int or None).
    Style / label / ref / colour come from the status-sheet line of the chosen PO#."""
    meta = {r["PO#"]: r for _, r in lines.iterrows()}
    rows, problems = [], []
    nxt = 1
    for n, r in enumerate(rows_in, 1):
        po = r.get("po")
        if po not in meta:
            if po:
                problems.append(f"Row {n}: pick a PO# from the list - row ignored.")
            continue
        try:
            ctns = int(float(r.get("ctns")))
            dims = tuple(float(x) for x in r.get("dims"))
        except (TypeError, ValueError):
            problems.append(f"Row {n} (PO {po}): cartons and carton size must be numbers - row ignored.")
            continue
        pcs = {}
        for sz, q in (r.get("pcs") or {}).items():
            try:
                q = int(float(q))
            except (TypeError, ValueError):
                q = 0
            if q > 0:
                pcs[sz] = q
        if ctns <= 0 or not pcs or min(dims) <= 0:
            problems.append(f"Row {n} (PO {po}): cartons, pieces and carton size must be above 0 - row ignored.")
            continue
        src = meta[po]
        ppk = sum(pcs.values())
        ship = ctns * ppk
        try:
            order = int(round(float(r.get("order")))) if r.get("order") not in (None, "") and float(r.get("order")) > 0 else ship
        except (TypeError, ValueError):
            order = ship
        short = ship - order
        net_kg = sum(q * float(wt.get(sz, 0.0)) for sz, q in pcs.items())
        c_from = 1 if restart else nxt
        c_to = c_from + ctns - 1
        nxt = c_to + 1
        rows.append({
            "po": po, "style": src["Style"], "label": src["Label"], "ref": src["Ref#"],
            "color": f"{src['Fashion Color']} - {src['Color']}".strip(" -"),
            "ctns": ctns, "order": order, "ship": ship, "short": short,
            "short_pct": (short / order) if order else 0.0,
            "pcs": pcs, "ppk": ppk, "net_kg": net_kg, "gross_kg": net_kg + tare,
            "c_from": c_from, "c_to": c_to, "dims": dims,
            "cbm": ctns * dims[0] * dims[1] * dims[2] * 2.54 ** 3 / 1_000_000,
        })
    tot_order = sum(x["order"] for x in rows)
    tot_ship = sum(x["ship"] for x in rows)
    gk = sum(x["gross_kg"] * x["ctns"] for x in rows)
    nk = sum(x["net_kg"] * x["ctns"] for x in rows)
    summary = {
        "total_order": tot_order, "total_ship": tot_ship, "excess_qty": tot_ship - tot_order,
        "excess_pct": (tot_ship - tot_order) / tot_order if tot_order else 0.0,
        "ctns": sum(x["ctns"] for x in rows), "gross_kg": gk, "gross_lbs": gk * LBS,
        "net_kg": nk, "net_lbs": nk * LBS, "cbm": sum(x["cbm"] for x in rows),
    }
    ref = rows[0]["ref"] if rows else ""
    dims0 = rows[0]["dims"] if rows else DEF_DIMS
    return Result(ref, sizes, rows, summary, tare, dims0, restart, []), problems


def main_df(res: Result) -> pd.DataFrame:
    """Preview of the main table (same columns / order as the Excel sheet)."""
    out = []
    for x in res.rows:
        d = {"Ctn from": x["c_from"], "Ctn to": x["c_to"], "Ctns": x["ctns"], "Style": x["style"],
             "Label": x["label"], "PO NO.": x["po"], "REF.NO": x["ref"], "Color": x["color"],
             "Order qty": x["order"]}
        for s in res.sizes:
            d[s] = x["pcs"].get(s)
        d.update({
            "Pcs/PPK": x["ppk"], "Ship qty": x["ship"], "Short/Excess": x["short"],
            "Short/Excess %": round(x["short_pct"] * 100, 2),
            "Ctn GRS (lbs)": round(x["gross_kg"] * LBS, 2), "Ctn NET (lbs)": round(x["net_kg"] * LBS, 2),
            "Total GRS (lbs)": round(x["gross_kg"] * LBS * x["ctns"], 2),
            "Total NET (lbs)": round(x["net_kg"] * LBS * x["ctns"], 2),
            "CBM": round(x["cbm"], 3), "Ctn L (in)": x["dims"][0], "Ctn W (in)": x["dims"][1], "Ctn H (in)": x["dims"][2],
        })
        out.append(d)
    df = pd.DataFrame(out)
    s = res.summary
    tot = {c: None for c in df.columns}
    tot.update({"Color": "TOTAL", "Ctns": s["ctns"], "Order qty": s["total_order"], "Ship qty": s["total_ship"],
                "Short/Excess": s["excess_qty"], "Short/Excess %": round(s["excess_pct"] * 100, 2),
                "Total GRS (lbs)": round(s["gross_lbs"], 2), "Total NET (lbs)": round(s["net_lbs"], 2),
                "CBM": round(s["cbm"], 3)})
    return pd.concat([df, pd.DataFrame([tot])], ignore_index=True)


def summary_df(res: Result) -> pd.DataFrame:
    s = res.summary
    return pd.DataFrame([
        ("Total order qty", f"{s['total_order']:,} pcs"),
        ("Total ship qty", f"{s['total_ship']:,} pcs"),
        ("Short/excess qty", f"{s['excess_qty']:+,} pcs"),
        ("Short/excess %", f"{s['excess_pct'] * 100:.2f}%"),
        ("Total cartons", f"{s['ctns']:,}"),
        ("Total gross weight", f"{s['gross_kg']:,.1f} kg / {s['gross_lbs']:,.0f} lbs"),
        ("Total net weight", f"{s['net_kg']:,.1f} kg / {s['net_lbs']:,.0f} lbs"),
        ("Volume", f"{s['cbm']:,.3f} CBM"),
    ], columns=["Item", "Value"])


def file_stem(res: Result, lines: pd.DataFrame) -> str:
    d = lines["ShipDT"].dropna().min() if len(lines) else None
    tag = f"{d.month}_{d.day}" if d is not None and not pd.isna(d) else "NA"
    return f"PACKING_LIST_{res.ref}_-_USA_{tag}"


# --------------------------------------------------------------------------- Excel
def _cps(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def build_xlsx(res: Result, header: dict | None = None) -> bytes:
    """
    Writes the reference layout with live formulas.
    header keys (all optional): date, container, invoice, warehouse
    """
    header = header or {}
    tpl = openpyxl.load_workbook(TEMPLATE)
    ts = tpl.active
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ts.title
    wb.loaded_theme = tpl.loaded_theme          # keep the template's theme colours

    n, P = len(res.sizes), len(res.rows)
    dc, dr = n - 4, P - 4
    K = LBS

    def role(k):                                # which template size column styles size #k
        if n == 4:
            return 10 + k
        if k == 0:
            return 10
        if k == n - 1:
            return 13
        return 11

    def colmap(c):                              # template col -> [(new col, source col)]
        if c <= 9:
            return [(c, c)]
        if c >= 14:
            return [(c + dc, c)]
        if c == 10:
            return [(10 + k, role(k)) for k in range(n)]
        return []

    def newrow(r):
        if r <= 15:
            return r
        if r == 20:
            return 16 + P
        return r + dr

    sub = 16 + P                                 # subtotal row
    first, last = 16, 15 + P

    # ---- copy header block, subtotal row and summary block (rows 1-15, 20-29)
    for r in list(range(1, 16)) + list(range(20, 30)):
        nr = newrow(r)
        for c in range(1, 28):
            for k, (nc, sc) in enumerate(colmap(c)):
                s, d = ts.cell(r, sc), ws.cell(nr, nc)
                _cps(s, d)
                if c == 10:
                    take = (k == 0) if n != 4 else True
                else:
                    take = True
                v = s.value
                if take and v is not None and not (isinstance(v, str) and v.startswith("=")):
                    d.value = v

    # ---- body rows
    for i, x in enumerate(res.rows):
        nr = 16 + i
        tr = 16 if i == 0 else (19 if i == P - 1 else 17)
        for c in range(1, 28):
            for nc, sc in colmap(c):
                _cps(ts.cell(tr, sc), ws.cell(nr, nc))
        cN, cO, cP, cQ, cR, cS, cT, cU, cV = (CL(14 + dc + j) for j in range(9))
        cW, cX, cY = CL(23 + dc), CL(24 + dc), CL(25 + dc)
        s1, s2 = CL(10), CL(10 + n - 1)
        if res.restart:
            ws[f"A{nr}"] = 1
        else:
            ws[f"A{nr}"] = 1 if i == 0 else f"=B{nr - 1}+1"
        ws[f"B{nr}"] = f"=A{nr}+C{nr}-1"
        ws[f"C{nr}"] = x["ctns"]
        ws[f"D{nr}"] = x["style"]
        ws[f"E{nr}"] = x["label"]
        ws[f"F{nr}"] = x["po"]
        ws[f"G{nr}"] = x["ref"]
        ws[f"H{nr}"] = x["color"]
        ws[f"I{nr}"] = x["order"]
        for k, sz in enumerate(res.sizes):
            v = x["pcs"].get(sz)
            if v:
                ws.cell(nr, 10 + k).value = v
        ws[f"{cN}{nr}"] = f"=SUM({s1}{nr}:{s2}{nr})"
        ws[f"{cO}{nr}"] = f"=C{nr}*{cN}{nr}"
        ws[f"{cP}{nr}"] = f"={cO}{nr}-I{nr}"
        ws[f"{cQ}{nr}"] = f"={cP}{nr}/I{nr}"
        ws[f"{cR}{nr}"] = f"=({x['net_kg']:.5f}+{res.tare:g})*{K}"
        ws[f"{cS}{nr}"] = f"={x['net_kg']:.5f}*{K}"
        ws[f"{cT}{nr}"] = f"={cR}{nr}*C{nr}"
        ws[f"{cU}{nr}"] = f"={cS}{nr}*C{nr}"
        ws[f"{cV}{nr}"] = f"=C{nr}*({cW}{nr}*2.54)*({cX}{nr}*2.54)*({cY}{nr}*2.54)/1000000"
        ws[f"{cW}{nr}"], ws[f"{cX}{nr}"], ws[f"{cY}{nr}"] = x["dims"]

    cN, cO, cP, cQ, cR, cS, cT, cU, cV = (CL(14 + dc + j) for j in range(9))

    # ---- size names
    for k, sz in enumerate(res.sizes):
        ws.cell(15, 10 + k).value = sz

    # ---- subtotal row
    for col in ("C", "I", cO, cP, cT, cU, cV):
        ws[f"{col}{sub}"] = f"=SUBTOTAL(9,{col}{first}:{col}{last})"
    ws[f"{cQ}{sub}"] = f"={cP}{sub}/I{sub}"

    # ---- summary block (template rows 22-29)
    r22 = 22 + dr
    ws[f"D{r22}"] = f"=I{sub}"
    ws[f"D{r22 + 1}"] = f"={cO}{sub}"
    ws[f"D{r22 + 2}"] = f"=D{r22 + 1}-D{r22}"
    ws[f"D{r22 + 3}"] = f"=D{r22 + 2}/D{r22}"
    ws[f"D{r22 + 4}"] = f"=C{sub}"
    ws[f"D{r22 + 5}"] = f"={cT}{sub}"                    # lbs
    ws[f"F{r22 + 5}"] = f"=D{r22 + 5}/{K}"               # kgs
    ws[f"D{r22 + 6}"] = f"={cU}{sub}"
    ws[f"F{r22 + 6}"] = f"=D{r22 + 6}/{K}"
    ws[f"D{r22 + 7}"] = f"={cV}{sub}"

    # ---- header inputs
    d = header.get("date")
    ws["I11"] = d if d else dt.date.today()
    if header.get("container"):
        ws["I9"] = header["container"]
    if header.get("invoice"):
        ws["I10"] = header["invoice"]
    if header.get("warehouse"):
        ws[f"{CL(16 + dc)}13"] = header["warehouse"]

    # ---- merges
    def mc(c, end):
        if c <= 10:
            return c
        if c <= 13:
            return (10 + n - 1) if (end and c == 13) else min(c, 10 + n - 1)
        return c + dc

    for rng in ts.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        a = (newrow(r1), mc(c1, False))
        b = (newrow(r2), mc(c2, True))
        if a != b:
            ws.merge_cells(start_row=a[0], start_column=a[1], end_row=b[0], end_column=b[1])

    # ---- widths / heights
    def tw(c):
        for d_ in ts.column_dimensions.values():
            if d_.min and d_.max and d_.min <= c <= d_.max and d_.width:
                return d_.width
        return None

    if ts.sheet_format.defaultColWidth:
        ws.sheet_format.defaultColWidth = ts.sheet_format.defaultColWidth
    for j in range(1, 28 + dc):
        if j <= 9:
            w = tw(j)
        elif j < 10 + n:
            w = max(tw(10) or 5.7, 6.5)
        else:
            w = tw(j - dc)
        if w:
            ws.column_dimensions[CL(j)].width = w
    for r in range(1, 30):
        h = ts.row_dimensions[r].height
        if h:
            ws.row_dimensions[newrow(r) if r not in range(16, 20) else 16].height = h
    for i in range(P):
        ws.row_dimensions[16 + i].height = ts.row_dimensions[16].height or 39.95

    # ---- logo (cropped like the template)
    try:
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        if ts._images:
            im = PILImage.open(io.BytesIO(ts._images[0]._data()))
            w, h = im.size
            im = im.crop((0, 0, w, int(h * (1 - 0.26029))))
            buf = io.BytesIO()
            im.save(buf, format="PNG")
            buf.seek(0)
            xi = XLImage(buf)
            xi.width, xi.height = 400, 81.5
            xi.anchor = OneCellAnchor(_from=AnchorMarker(col=7, colOff=1796362, row=1, rowOff=34108),
                                      ext=XDRPositiveSize2D(3810000, 775948))
            ws.add_image(xi)
    except Exception:
        pass

    # ---- sheet / print setup
    ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ts.page_setup.paperSize or 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{CL(25 + dc)}{29 + dr}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- PDF
def build_pdf(res: Result, header: dict | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    header = header or {}
    ts = openpyxl.load_workbook(TEMPLATE).active
    clean = lambda v: re.sub(r"\s{2,}", "\n", str(v or "")).strip()
    factory, addr = clean(ts["C4"].value), clean(ts["C5"].value)
    ship_to = clean(ts["C8"].value)
    d = header.get("date") or dt.date.today()
    n = len(res.sizes)

    fixed = [22, 15, 17, 15, 18, 24, 46, 16] + [11] * n + [14, 16, 17, 17, 18, 18, 18, 18, 15, 10, 10, 10]
    avail = 400.0
    scale = avail / sum(fixed)
    widths = [w * scale * mm for w in fixed]
    fs = max(4.6, min(7.0, 7.0 * scale))

    sty = ParagraphStyle("c", fontName="Helvetica", fontSize=fs, leading=fs + 1.5, alignment=1)
    stb = ParagraphStyle("b", parent=sty, fontName="Helvetica-Bold")
    P_ = lambda t, s=sty: Paragraph(str(t), s)
    big = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=18, leading=24, alignment=1, spaceAfter=8)
    lab = ParagraphStyle("l", fontName="Helvetica", fontSize=8, leading=10)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A3), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f"Packing list {res.ref}")
    story = [Paragraph("HADDAD BRAND PACKING LIST", big)]

    info = [[Paragraph(f"<b>FACTORY</b><br/>{factory.replace(chr(10), '<br/>')}<br/>{addr.replace(chr(10), '<br/>')}", lab),
             Paragraph(f"<b>SHIPPING ADDRESS</b><br/>{ship_to.replace(chr(10), '<br/>')}", lab),
             Paragraph(f"<b>DATE</b> {d:%Y-%m-%d}<br/><b>CONTAINER NO</b> {header.get('container', '')}<br/>"
                       f"<b>INVOICE NO</b> {header.get('invoice', '')}<br/><b>WAREHOUSE CODE</b> {header.get('warehouse', '')}<br/>"
                       f"<b>FROM</b> BANGLADESH &nbsp;&nbsp; <b>TO</b> USA", lab)]]
    it = Table(info, colWidths=[sum(widths) / 3] * 3)
    it.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.black), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [it, Spacer(1, 4 * mm)]

    # header rows
    h1 = ["CARTON NO\nFrom - To", "TOTAL\nCTN", "STYLE NO.", "LABEL", "PO NO.", "REF.NO", "COLOR", "ORDER\nQTY"] + \
         ["PACK QUANTITY"] + [""] * (n - 1) + \
         ["UNITS PER\nPPK (PCS)", "SHIP QTY\n(PCS)", "SHORT/\nEXCESS", "SHORT/\nEXCESS %", "PER CTN\nGRS (LBS)",
          "PER CTN\nNET (LBS)", "TOTAL\nGRS (LBS)", "TOTAL\nNET (LBS)", "TOTAL\nCBM", "CTNS MEASUREMENT (INCH)", "", ""]
    h2 = [""] * 8 + list(res.sizes) + [""] * 9 + ["L", "W", "H"]
    data = [[P_(c.replace("\n", "<br/>"), stb) for c in h1], [P_(c, stb) for c in h2]]
    for x in res.rows:
        row = [P_(f"{x['c_from']} - {x['c_to']}"), P_(x["ctns"]), P_(x["style"]), P_(x["label"]), P_(x["po"]),
               P_(x["ref"]), P_(x["color"]), P_(f"{x['order']:,}")]
        row += [P_(x["pcs"].get(s, "") or "") for s in res.sizes]
        row += [P_(x["ppk"]), P_(f"{x['ship']:,}"), P_(f"{x['short']:+,}"), P_(f"{x['short_pct'] * 100:.2f}%"),
                P_(f"{x['gross_kg'] * LBS:.2f}"), P_(f"{x['net_kg'] * LBS:.2f}"),
                P_(f"{x['gross_kg'] * LBS * x['ctns']:,.1f}"), P_(f"{x['net_kg'] * LBS * x['ctns']:,.1f}"),
                P_(f"{x['cbm']:.3f}"), P_(f"{x['dims'][0]:g}"), P_(f"{x['dims'][1]:g}"), P_(f"{x['dims'][2]:g}")]
        data.append(row)
    s = res.summary
    tot = [P_("", stb) for _ in range(8 + n + 12)]
    tot[1] = P_(f"{s['ctns']:,}", stb)
    tot[7] = P_(f"{s['total_order']:,}", stb)
    base = 8 + n
    tot[base + 1] = P_(f"{s['total_ship']:,}", stb)
    tot[base + 2] = P_(f"{s['excess_qty']:+,}", stb)
    tot[base + 3] = P_(f"{s['excess_pct'] * 100:.2f}%", stb)
    tot[base + 6] = P_(f"{s['gross_lbs']:,.1f}", stb)
    tot[base + 7] = P_(f"{s['net_lbs']:,.1f}", stb)
    tot[base + 8] = P_(f"{s['cbm']:.3f}", stb)
    data.append(tot)

    t = Table(data, colWidths=widths, repeatRows=2)
    nl = len(data)
    last_c = len(fixed) - 1
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, nl - 1), (-1, nl - 1), colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("SPAN", (8, 0), (7 + n, 0)),
        ("SPAN", (last_c - 2, 0), (last_c, 0)),
        *[("SPAN", (c, 0), (c, 1)) for c in list(range(0, 8)) + list(range(8 + n, last_c - 2))],
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story += [t, Spacer(1, 5 * mm)]

    sm = summary_df(res)
    st = Table([[P_(a, ParagraphStyle("sl", parent=lab, fontName="Helvetica-Bold")), P_(b, ParagraphStyle("sv", parent=lab, alignment=0))]
                for a, b in sm.itertuples(index=False)], colWidths=[50 * mm, 70 * mm], hAlign="LEFT")
    st.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFEFEF"))]))
    story.append(st)
    doc.build(story)
    return buf.getvalue()
