

"""
PO -> Packing List Streamlit App
=================================
Upload a VF/Vans-style PO PDF, review the extracted data, and download
a packing-list-style Excel workbook (one sheet per PO line).

HOW TO RUN:
    pip install streamlit pdfplumber openpyxl reportlab pandas
    streamlit run VF_Packing_app.py
"""
import io
import json
import re
import uuid
import datetime

import pandas as pd
import pdfplumber
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

# ============================================================
# PART 1 - PDF EXTRACTION
# ============================================================

LINE_PATTERN = re.compile(
    r"""
    (?P<po_line_no>\d{15})
    \s+
    (?P<style>[A-Z0-9]+)
    \s+
    (?P<description>[A-Z0-9 \-/]+?)
    \s+
    (?P<order_qty>[\d,]+)
    \s+
    (?:PIECE|EACH)
    \s+
    (?P<unit_price>[\d.]+)
    \s+
    (?P<line_cost>[\d,]+\.\d{2})
    """,
    re.IGNORECASE | re.VERBOSE,
)

SUBLINE_START_PATTERN = re.compile(
    r"(?P<subline_no>\d{12,13})\s+(?P<style>VN[A-Z0-9]+)",
    re.IGNORECASE,
)


def extract_sublines(full_text):
    text = full_text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    matches = list(SUBLINE_START_PATTERN.finditer(text))
    sublines = []

    for i, start_match in enumerate(matches):
        block_start = start_match.start()
        block_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        block = text[block_start:block_end]

        if not re.search(r"UnitOfMeasureCode\s+(?:PC|EA)\b", block, re.I):
            continue

        qty_cost = re.search(
            r"(?P<style>VN[A-Z0-9]+)\s+.*?"
            r"(?P<qty>[\d,]+)\s+(?P<unit_cost>\d+(?:\.\d+)?)\s+.*?"
            r"UnitOfMeasureCode",
            block, re.IGNORECASE | re.DOTALL,
        )
        color = re.search(
            r"UnitOfMeasureCode\s+(?:PC|EA)\s+Color\s+(?P<color>.*?)\s+Size",
            block, re.IGNORECASE | re.DOTALL,
        )
        size = re.search(
            r"Size\s+(?P<size>.*?)\s+Dimension\s+1",
            block, re.IGNORECASE | re.DOTALL,
        )
        sku = re.search(r"SKU\s+Number\s+(?P<sku>\S+)", block, re.IGNORECASE)
        upc = re.search(r"UPC\s+Number\s+(?P<upc>\d{12,13})", block, re.IGNORECASE)
        pack = re.search(
            r"Items\s+Per\s+Outer\s+(?:Pack\s+)?(?P<pack>[\d,]+)",
            block, re.IGNORECASE | re.DOTALL,
        )

        if not all((qty_cost, color, size, sku, upc, pack)):
            continue

        sublines.append({
            "subline_no": start_match.group("subline_no"),
            "style": start_match.group("style"),
            "qty": int(qty_cost.group("qty").replace(",", "")),
            "unit_cost": float(qty_cost.group("unit_cost").replace(",", "")),
            "color": re.sub(r"\s+", " ", color.group("color")).strip(),
            "size": re.sub(r"\s+", " ", size.group("size")).strip(),
            "sku": sku.group("sku"),
            "upc": upc.group("upc"),
            "items_per_outer_pack": int(pack.group("pack").replace(",", "")),
        })

    return sublines


CRD_PATTERN = re.compile(r"Brand Requested CRD\s+(\d{4}-\d{2}-\d{2})")
COUNTRY_PATTERN = re.compile(r"([A-Z]{3,})\s*\na\s*\na\s*\nBrand Buyer Code")


def extract_data(pdf_file):
    with pdfplumber.open(pdf_file) as pdf:
        full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)

    lines = []
    for m in LINE_PATTERN.finditer(full_text):
        lines.append({
            "po_line_no": m.group("po_line_no"),
            "style": m.group("style"),
            "description": re.sub(r"\s+", " ", m.group("description")).strip(),
            "order_qty": int(m.group("order_qty").replace(",", "")),
            "unit_price": float(m.group("unit_price").replace(",", "")),
            "line_cost": float(m.group("line_cost").replace(",", "")),
        })

    sublines = extract_sublines(full_text)

    crd_matches = CRD_PATTERN.findall(full_text)
    for i, l in enumerate(lines):
        l["crd"] = crd_matches[i] if i < len(crd_matches) else None

    country_match = COUNTRY_PATTERN.search(full_text)
    destination_country = country_match.group(1) if country_match else None
    for l in lines:
        l["destination_country"] = destination_country

    return {"lines": lines, "sublines": sublines}


@st.cache_data(show_spinner=False, max_entries=10)
def extract_data_cached(pdf_bytes: bytes):
    return extract_data(io.BytesIO(pdf_bytes))


_HAS_FRAGMENT = hasattr(st, "fragment")


def _fragment(fn):
    return st.fragment(fn) if _HAS_FRAGMENT else fn


# ============================================================
# PART 2 - CARTON ROW COMPUTATION
# ============================================================

def compute_carton_rows_v3(sublines, max_by_size, mode="solid_mixed"):
    if not sublines:
        return []

    size_order = [s["size"] for s in sublines]
    full_rows = []
    partial_rows = []
    remainder_pool = []

    for s in sublines:
        size = s["size"]
        qty = int(s["qty"])
        pack = int(max_by_size.get(size) or s.get("items_per_outer_pack") or 1)
        if pack < 1:
            pack = 1

        full_ctn = qty // pack
        rem = qty % pack

        if full_ctn > 0:
            full_rows.append({
                "sizes": {size: pack},
                "total_ctn": full_ctn,
                "ctn_pcs": pack,
                "color": s["color"],
                "upc": s["upc"],
                "mixed": False,
            })

        if rem > 0:
            if mode == "solid_no_leftover":
                pass
            elif mode == "solid_all":
                partial_rows.append({
                    "sizes": {size: rem},
                    "total_ctn": 1,
                    "ctn_pcs": rem,
                    "color": s["color"],
                    "upc": s["upc"],
                    "mixed": False,
                })
            else:
                remainder_pool.append({
                    "size": size,
                    "rem": rem,
                    "color": s["color"],
                    "upc": s["upc"],
                    "max": pack,
                })

    size_to_full = {}
    for r in full_rows:
        for sz in r["sizes"]:
            size_to_full.setdefault(sz, []).append(r)

    size_to_partial = {}
    for r in partial_rows:
        for sz in r["sizes"]:
            size_to_partial.setdefault(sz, []).append(r)

    solid_block = []
    for sz in size_order:
        for r in size_to_full.get(sz, []):
            solid_block.append(r)
        for r in size_to_partial.get(sz, []):
            solid_block.append(r)

    if mode in ("solid_no_leftover", "solid_all"):
        return solid_block

    mixed_rows = []
    idx = 0
    while idx < len(remainder_pool):
        carton = {}
        carton_total = 0
        carton_cap = None
        progressed = False

        while idx < len(remainder_pool):
            entry = remainder_pool[idx]
            size = entry["size"]
            size_max = entry["max"]
            proposed_cap = size_max if carton_cap is None else min(carton_cap, size_max)

            if carton_total > proposed_cap:
                break
            space = proposed_cap - carton_total
            if space <= 0:
                break

            take = min(entry["rem"], space)
            if take > 0:
                carton[size] = carton.get(size, 0) + take
                carton_total += take
                entry["rem"] -= take
                carton_cap = proposed_cap
                progressed = True

            if entry["rem"] == 0:
                idx += 1
            else:
                break

            if carton_total >= carton_cap:
                break

        if carton:
            mixed_rows.append({
                "sizes": dict(carton),
                "total_ctn": 1,
                "ctn_pcs": carton_total,
                "color": None,
                "upc": None,
                "mixed": True,
            })
        if not progressed:
            break

    return solid_block + mixed_rows


# ============================================================
# PART 3 - EXCEL BUILDING
# ============================================================

FONT = "Arial"
bold = Font(name=FONT, bold=True, size=10)
normal = Font(name=FONT, size=10)
title_font = Font(name=FONT, bold=True, size=13)
thin = Side(style="thin")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
yellow = PatternFill("solid", fgColor="FFFF00")
grey = PatternFill("solid", fgColor="D9D9D9")


def sheet_name_for(line):
    raw = f"{line['style']}"
    return re.sub(r'[\[\]:*?/\\]', '-', raw)[:31]


def build_sheet(ws, line, sublines):
    """
    Packing-list sheet.
    Layout rules (centered design):
    - Title rows are merged across the FULL table width (A .. last table column)
    - SIZE / N.W. / N.N.W. / EMPTY CTN block is centered over the table
    - Summary block (SIZE / CUT / ORDER / SHIP / EXS-SHT / % + G Total + weights)
      sits under the table, its size columns aligned with the main size columns,
      and its left/right margins balanced so it reads as centered
    - Every cell is center/center; only bottom labels are right-aligned
    """
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    # ---------- styles ----------
    FONT = "Arial"
    bold = Font(name=FONT, bold=True, size=10)
    normal = Font(name=FONT, size=10)
    title_font = Font(name=FONT, bold=True, size=13)
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    yellow = PatternFill("solid", fgColor="FFFF00")
    grey = PatternFill("solid", fgColor="D9D9D9")

    # ---------- helpers ----------
    def style_range(r1, c1, r2, c2, border=None, fill=None, font=None, align=None):
        """Apply style to every cell of a (possibly merged) range."""
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                cell = ws.cell(row=rr, column=cc)
                if border is not None:
                    cell.border = border
                if fill is not None:
                    cell.fill = fill
                if font is not None:
                    cell.font = font
                if align is not None:
                    cell.alignment = align

    def put_merged(r1, c1, r2, c2, value=None, font=bold, fill=None,
                   border=box, align=center):
        """Merge a range, write value in top-left, style all cells."""
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        ws.cell(row=r1, column=c1, value=value)
        style_range(r1, c1, r2, c2, border=border, fill=fill, font=font, align=align)

    # ---------- unique sizes ----------
    sizes = []
    seen = set()
    for s in sublines:
        sz = s["size"]
        if sz not in seen:
            sizes.append(sz)
            seen.add(sz)
    n_sizes = len(sizes)

    # ---------- column map (main table) ----------
    COL_CASE       = 1
    COL_CARTON     = 2
    COL_COLOR      = 3
    COL_UPC        = 4
    COL_SIZE_START = 5
    COL_CTN_PCS    = 5 + n_sizes
    COL_TOTAL_CTN  = COL_CTN_PCS + 1
    COL_TOTAL_PCS  = COL_CTN_PCS + 2
    COL_GRS_PR     = COL_CTN_PCS + 3
    COL_NET_PR     = COL_CTN_PCS + 4
    COL_TOTAL_GRS  = COL_CTN_PCS + 5
    COL_TOTAL_NET  = COL_CTN_PCS + 6
    n_cols = COL_TOTAL_NET            # last column of the table

    # ---------- column widths (uniform: NO widened size columns) ----------
    W_SIZE = 10
    widths = {1: 18, 2: 14, 3: 14, 4: 16}
    for i in range(n_sizes):
        widths[COL_SIZE_START + i] = W_SIZE
    for col in range(COL_CTN_PCS, n_cols + 1):
        widths[col] = 13
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    def centered_start(block_cols):
        """First column so a block of `block_cols` columns is centered
        (by real column width) under the full table A..n_cols."""
        best, best_gap = 1, None
        for s0 in range(1, n_cols - block_cols + 2):
            lw = sum(widths[c] for c in range(1, s0))
            rw = sum(widths[c] for c in range(s0 + block_cols, n_cols + 1))
            gap = abs(lw - rw)
            if best_gap is None or gap < best_gap:
                best, best_gap = s0, gap
        return best

    # =========================================================
    # 1. Title – merged over the WHOLE table width, centered
    # =========================================================
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws["A1"] = "CREATIVE COLLECTIONS LTD."
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    ws["A2"] = "PACKING LIST DETAILS."
    ws["A2"].font = Font(name=FONT, bold=True, size=11)
    ws["A2"].alignment = center

    # =========================================================
    # 2. Weight spec table – centered over the table
    #    label (2 merged cols) + one column per size
    # =========================================================
    spec_header_row = 4
    spec_c0 = centered_start(n_sizes + 2)          # first column of the block
    spec_first = spec_c0 + 2                       # first size column
    spec_last = spec_c0 + 1 + n_sizes              # last size column

    put_merged(spec_header_row, spec_c0, spec_header_row, spec_c0 + 1,
               "SIZE", font=bold, fill=grey)
    for i, sz in enumerate(sizes):
        c = ws.cell(row=spec_header_row, column=spec_first + i, value=sz)
        c.font = bold
        c.border = box
        c.fill = grey
        c.alignment = center

    for offset, label in enumerate(["N.W.", "N.N.W.", "EMPTY CTN"]):
        r = spec_header_row + 1 + offset
        put_merged(r, spec_c0, r, spec_c0 + 1, label, font=bold)
        for i in range(n_sizes):
            c = ws.cell(row=r, column=spec_first + i, value=1.0)
            c.border = box
            c.alignment = center
            c.font = normal
            if label == "N.W.":
                c.fill = yellow

    # =========================================================
    # 3. Header info block
    #    left : labels + values (values merged so long text never clips)
    #    right: qty block pinned to the right edge of the table
    # =========================================================
    info_start = 9
    val_c1, val_c2 = 2, COL_SIZE_START + 3          # B .. H

    def info_left(row, label, value):
        c = ws.cell(row=row, column=1, value=label)
        c.font = bold
        c.alignment = left
        ws.merge_cells(start_row=row, start_column=val_c1, end_row=row, end_column=val_c2)
        v = ws.cell(row=row, column=val_c1, value=value)
        v.font = normal
        v.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    info_left(info_start,     "BUYER",       "VANS")
    info_left(info_start + 1, "LOT NO",      line.get("style", ""))
    info_left(info_start + 2, "P.O NO",      line.get("po_line_no", ""))
    info_left(info_start + 3, "DESCRIPTION", line.get("description", ""))

    # long merged text -> give the row enough height
    avail_chars = int(sum(widths[c] for c in range(val_c1, val_c2 + 1)) * 0.95)
    for k, key in enumerate(["style", "po_line_no", "description"], start=1):
        txt = str(line.get(key, "") or "")
        n_lines = max(1, -(-len(txt) // max(avail_chars, 1)))
        if n_lines > 1:
            ws.row_dimensions[info_start + k].height = 14 * n_lines

    # right block: [label x2][value][unit]   [CRD label][CRD value]
    lab_c = n_cols - 6          # label (merged over 2 cols)
    val_c = n_cols - 4
    unit_c = n_cols - 3
    crd_l = n_cols - 1
    crd_v = n_cols

    def info_right(row, label):
        put_merged(row, lab_c, row, lab_c + 1, label, font=bold, border=None)
        v = ws.cell(row=row, column=val_c)
        v.font = normal
        v.alignment = center
        return v

    ws.cell(row=info_start, column=val_c, value=line.get("order_qty", 0))
    order_cell = info_right(info_start, "ORDER QTY")
    pack_qty_cell = info_right(info_start + 1, "PACK QTY")
    excess_cell   = info_right(info_start + 2, "EXCESS/SHORT QTY")
    pct_cell      = info_right(info_start + 3, "PERCENTAGE")
    for row, unit in enumerate(["PCS", "PCS", "PCS", "%"], start=info_start):
        c = ws.cell(row=row, column=unit_c, value=unit)
        c.font = normal
        c.alignment = center

    c = ws.cell(row=info_start, column=crd_l, value="CRD")
    c.font = bold
    c.alignment = center
    c = ws.cell(row=info_start, column=crd_v, value=line.get("crd", ""))
    c.font = normal
    c.alignment = center

    c = ws.cell(row=info_start + 1, column=crd_l, value="COUNTRY")
    c.font = bold
    c.alignment = center
    c = ws.cell(row=info_start + 1, column=crd_v, value=line.get("destination_country", ""))
    c.font = normal
    c.alignment = center

    # =========================================================
    # 4. Main carton table header
    # =========================================================
    header_row = info_start + 6

    headers = (
        ["CASE LABEL NO.", "CARTON NO.", "COLOR", "UPC Number"]
        + sizes
        + ["CTN PCS", "TOTAL CTN", "TOTAL PCS",
           "Grs.wt.pr ctn", "Net.wt.pr ctn", "total.Grs.wt", "total.net.wt"]
    )

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = bold
        c.border = box
        c.fill = grey
        c.alignment = center
    ws.row_dimensions[header_row].height = 30

    # =========================================================
    # 5. Carton data rows
    # =========================================================
    pdf_max_by_size = {s["size"]: int(s.get("items_per_outer_pack") or 1) for s in sublines}
    carton_rows = compute_carton_rows_v3(sublines, pdf_max_by_size, mode="solid_mixed")

    data_start = header_row + 1
    running_carton = 0

    for idx, row_data in enumerate(carton_rows):
        r = data_start + idx

        c = ws.cell(row=r, column=COL_CASE, value=idx + 1)
        c.border = box
        c.alignment = center

        total_ctn = int(row_data["total_ctn"])
        start_ctn = running_carton + 1
        end_ctn = running_carton + total_ctn
        running_carton = end_ctn
        carton_label = str(start_ctn) if start_ctn == end_ctn else f"{start_ctn}-{end_ctn}"
        c = ws.cell(row=r, column=COL_CARTON, value=carton_label)
        c.border = box
        c.alignment = center

        color_val = row_data["color"] if row_data["color"] is not None else "MIXED"
        c = ws.cell(row=r, column=COL_COLOR, value=color_val)
        c.border = box
        c.alignment = center
        if row_data.get("mixed"):
            c.fill = yellow

        c = ws.cell(row=r, column=COL_UPC, value=row_data.get("upc") or "")
        c.border = box
        c.alignment = center

        for i, sz in enumerate(sizes):
            val = row_data["sizes"].get(sz)
            c = ws.cell(row=r, column=COL_SIZE_START + i, value=val if val else None)
            c.border = box
            c.alignment = center

        first_sz = get_column_letter(COL_SIZE_START)
        last_sz  = get_column_letter(COL_SIZE_START + n_sizes - 1)
        c = ws.cell(row=r, column=COL_CTN_PCS,
                    value=f"=SUM({first_sz}{r}:{last_sz}{r})")
        c.border = box
        c.alignment = center
        c.font = bold

        c = ws.cell(row=r, column=COL_TOTAL_CTN, value=total_ctn)
        c.border = box
        c.alignment = center

        ctn_pcs_l = get_column_letter(COL_CTN_PCS)
        tot_ctn_l = get_column_letter(COL_TOTAL_CTN)
        c = ws.cell(row=r, column=COL_TOTAL_PCS,
                    value=f"={tot_ctn_l}{r}*{ctn_pcs_l}{r}")
        c.border = box
        c.alignment = center

        # Gross weight per carton (spec block may be shifted -> use spec_first/last)
        nw_row = spec_header_row + 1
        empty_row = spec_header_row + 3
        first_spec = get_column_letter(spec_first)
        last_spec  = get_column_letter(spec_last)

        grs_formula = (
            f"=SUMPRODUCT({first_sz}{r}:{last_sz}{r},"
            f"{first_spec}{nw_row}:{last_spec}{nw_row})"
            f"+{first_spec}{empty_row}"
        )
        c = ws.cell(row=r, column=COL_GRS_PR, value=grs_formula)
        c.border = box
        c.alignment = center

        grs_l = get_column_letter(COL_GRS_PR)
        c = ws.cell(row=r, column=COL_NET_PR,
                    value=f"={grs_l}{r}-{first_spec}{empty_row}")
        c.border = box
        c.alignment = center

        c = ws.cell(row=r, column=COL_TOTAL_GRS,
                    value=f"={tot_ctn_l}{r}*{grs_l}{r}")
        c.border = box
        c.alignment = center

        net_l = get_column_letter(COL_NET_PR)
        c = ws.cell(row=r, column=COL_TOTAL_NET,
                    value=f"={tot_ctn_l}{r}*{net_l}{r}")
        c.border = box
        c.alignment = center

    data_end = data_start + len(carton_rows) - 1
    total_row = data_end + 1

    # =========================================================
    # 6. TOTAL row
    # =========================================================
    for col in range(1, n_cols + 1):
        ws.cell(row=total_row, column=col).border = box
        ws.cell(row=total_row, column=col).alignment = center

    ws.cell(row=total_row, column=COL_COLOR, value="TOTAL").font = bold

    tot_ctn_l = get_column_letter(COL_TOTAL_CTN)
    for i in range(n_sizes):
        col = COL_SIZE_START + i
        col_l = get_column_letter(col)
        c = ws.cell(
            row=total_row, column=col,
            value=f"=SUMPRODUCT({col_l}{data_start}:{col_l}{data_end},"
                  f"{tot_ctn_l}{data_start}:{tot_ctn_l}{data_end})"
        )
        c.font = bold
        c.border = box
        c.alignment = center

    c = ws.cell(row=total_row, column=COL_TOTAL_CTN,
                value=f"=SUM({tot_ctn_l}{data_start}:{tot_ctn_l}{data_end})")
    c.font = bold
    c.border = box
    c.alignment = center

    tot_pcs_l = get_column_letter(COL_TOTAL_PCS)
    c = ws.cell(row=total_row, column=COL_TOTAL_PCS,
                value=f"=SUM({tot_pcs_l}{data_start}:{tot_pcs_l}{data_end})")
    c.font = bold
    c.border = box
    c.alignment = center

    for col in (COL_TOTAL_GRS, COL_TOTAL_NET):
        col_l = get_column_letter(col)
        c = ws.cell(row=total_row, column=col,
                    value=f"=SUM({col_l}{data_start}:{col_l}{data_end})")
        c.font = bold
        c.border = box
        c.alignment = center

    # =========================================================
    # 7. Summary block
    #    label  : C:D merged
    #    sizes  : exactly under the main size columns
    #    G Total: under CTN PCS
    #    GW     : TOTAL CTN : TOTAL PCS   (merged, rows merged too)
    #    NW     : Grs.wt.pr ctn : Net.wt.pr ctn (merged, rows merged too)
    #    -> block spans C .. n_cols-2, i.e. balanced margins either side
    # =========================================================
    sum_header = total_row + 3
    SUM_LABEL_C1, SUM_LABEL_C2 = COL_COLOR, COL_UPC          # C:D

    gt_col = COL_CTN_PCS
    gw_c1, gw_c2 = COL_TOTAL_CTN, COL_TOTAL_PCS
    nw_c1, nw_c2 = COL_GRS_PR, COL_NET_PR

    put_merged(sum_header, SUM_LABEL_C1, sum_header, SUM_LABEL_C2,
               "SIZE", font=bold, fill=grey)

    for i, sz in enumerate(sizes):
        c = ws.cell(row=sum_header, column=COL_SIZE_START + i, value=sz)
        c.font = bold
        c.border = box
        c.fill = grey
        c.alignment = center

    c = ws.cell(row=sum_header, column=gt_col, value="G Total")
    c.font = bold
    c.border = box
    c.fill = grey
    c.alignment = center

    put_merged(sum_header, gw_c1, sum_header, gw_c2,
               "Grand Total gross weight kg", font=bold, fill=grey)
    put_merged(sum_header, nw_c1, sum_header, nw_c2,
               "Grand Total Netweight kg", font=bold, fill=grey)
    ws.row_dimensions[sum_header].height = 30

    rows_needed = ["CUT QTY", "ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE"]
    cut_r   = sum_header + 1 + rows_needed.index("CUT QTY")
    order_r = sum_header + 1 + rows_needed.index("ORDER QTY")
    ship_r  = sum_header + 1 + rows_needed.index("SHIP QTY")
    exs_r   = sum_header + 1 + rows_needed.index("EXS/SHT QTY")
    pct_r   = sum_header + 1 + rows_needed.index("PERCENTAGE")

    first_l = get_column_letter(COL_SIZE_START)
    last_l  = get_column_letter(COL_SIZE_START + n_sizes - 1)
    gt_l    = get_column_letter(gt_col)

    for i, label in enumerate(rows_needed):
        r = sum_header + 1 + i
        put_merged(r, SUM_LABEL_C1, r, SUM_LABEL_C2, label, font=bold)

        for j, sz in enumerate(sizes):
            col = COL_SIZE_START + j
            col_l = get_column_letter(col)
            cell = ws.cell(row=r, column=col)
            cell.border = box
            cell.alignment = center
            cell.font = normal

            if label == "ORDER QTY":
                cell.value = next((s["qty"] for s in sublines if s["size"] == sz), 0)
            elif label == "CUT QTY":
                cell.fill = yellow
            elif label == "SHIP QTY":
                cell.value = f"={col_l}{total_row}"
            elif label == "EXS/SHT QTY":
                cell.value = f'=IF({col_l}{ship_r}="","",{col_l}{ship_r}-{col_l}{order_r})'
            elif label == "PERCENTAGE":
                cell.value = f'=IF({col_l}{exs_r}="","",IFERROR({col_l}{exs_r}/{col_l}{order_r},0))'
                cell.number_format = "0.00%"

        gt_cell = ws.cell(row=r, column=gt_col)
        gt_cell.border = box
        gt_cell.font = bold
        gt_cell.alignment = center

        if label in ("CUT QTY", "EXS/SHT QTY"):
            gt_cell.value = f'=IF(COUNT({first_l}{r}:{last_l}{r})=0,"",SUM({first_l}{r}:{last_l}{r}))'
        elif label in ("ORDER QTY", "SHIP QTY"):
            gt_cell.value = f"=SUM({first_l}{r}:{last_l}{r})"
        else:
            gt_cell.value = f'=IF({gt_l}{exs_r}="","",IFERROR({gt_l}{exs_r}/{gt_l}{order_r},0))'
            gt_cell.number_format = "0.00%"

    # Link top info block
    pack_qty_cell.value = f"={gt_l}{ship_r}"
    excess_cell.value   = f"={gt_l}{exs_r}"
    pct_cell.value = (
        f'=IF({gt_l}{exs_r}="","-",'
        f'TEXT(IFERROR({gt_l}{exs_r}/{gt_l}{order_r},0),"0.00%"))'
    )

    # Grand total weights (merged block, centered)
    first_sum = sum_header + 1
    last_sum  = sum_header + len(rows_needed)

    put_merged(first_sum, gw_c1, last_sum, gw_c2,
               f"={get_column_letter(COL_TOTAL_GRS)}{total_row}", font=bold)
    ws.cell(row=first_sum, column=gw_c1).number_format = "0.00"
    put_merged(first_sum, nw_c1, last_sum, nw_c2,
               f"={get_column_letter(COL_TOTAL_NET)}{total_row}", font=bold)
    ws.cell(row=first_sum, column=nw_c1).number_format = "0.00"

    # =========================================================
    # 8. Bottom weights + CTN MEAS + CBM  (labels RIGHT, values CENTER)
    # =========================================================
    wt_row = last_sum + 2

    c = ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :")
    c.font = bold
    c.alignment = right
    c = ws.cell(row=wt_row, column=2,
                value=f"={get_column_letter(COL_TOTAL_GRS)}{total_row}")
    c.font = bold
    # c.border = box
    c.alignment = center

    c = ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :")
    c.font = bold
    c.alignment = right
    c = ws.cell(row=wt_row + 1, column=2,
                value=f"={get_column_letter(COL_TOTAL_NET)}{total_row}")
    c.font = bold
    # c.border = box
    c.alignment = center

    c = ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :")
    c.font = bold
    c.alignment = right
    for k, txt in enumerate(["L 24'' X W 16'' X H 10''", "L 24'' X W 16'' X H 6''"]):
        put_merged(wt_row + 2 + k, 2, wt_row + 2 + k, 4, txt,
                   font=normal, fill=yellow)

    c = ws.cell(row=wt_row + 4, column=1, value="CBM :")
    c.font = bold
    c.alignment = right

    cbm_per_cell = ws.cell(row=wt_row + 4, column=2, value=0.06295)  # ≈ 24*16*10 / 61000
    cbm_per_cell.fill = yellow
    # cbm_per_cell.border = box
    cbm_per_cell.alignment = center
    cbm_per_cell.number_format = "0.00000"

    c = ws.cell(row=wt_row + 4, column=3,
                value=f"=B{wt_row + 4}*{get_column_letter(COL_TOTAL_CTN)}{total_row}")
    c.font = bold
    # c.border = box
    c.alignment = center
    c.number_format = "0.00000"

    # =========================================================
    # 9. Signature block
    # =========================================================
    sig_row = wt_row + 7

    sig_al = Alignment(horizontal="center", vertical="center", wrap_text=False)
    for col in (2, 5):
        ws.cell(row=sig_row, column=col, value="_" * 14).alignment = sig_al
    ws.cell(row=sig_row + 1, column=2, value="INCHARGE").font = bold
    ws.cell(row=sig_row + 1, column=2).alignment = sig_al
    ws.cell(row=sig_row + 1, column=5, value="FM").font = bold
    ws.cell(row=sig_row + 1, column=5).alignment = sig_al

    # Print setup: one page wide, landscape, centered on the page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True









def build_workbook(data):
    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()
    for line in data["lines"]:
        sublines = [s for s in data["sublines"] if s["style"] == line["style"]]
        if not sublines:
            continue
        name = sheet_name_for(line)
        base_name, i = name, 1
        while name in used_names:
            i += 1
            name = f"{base_name[:28]}_{i}"
        used_names.add(name)
        ws = wb.create_sheet(title=name)
        build_sheet(ws, line, sublines)
    return wb


def build_edited_pdf(export_state):
    line = export_state["line"]
    full_sizes = export_state["full_sizes"]
    rows = export_state["rows"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=5 * mm, rightMargin=5 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("CREATIVE COLLECTIONS LTD", styles["Title"]))
    elements.append(Paragraph("PACKING LIST DETAILS", styles["Heading3"]))
    elements.append(Spacer(1, 6))

    info_data = [
        ["BUYER", "VANS", "LOT NO", line["style"], "P.O NO", line["po_line_no"]],
        ["ORDER QTY", line["order_qty"], "EXCESS/SHORT QTY", export_state["exs_total"],
         "CRD", line.get("crd") or ""],
        ["COUNTRY", line.get("destination_country") or "", "DESCRIPTION",
         line["description"], "", ""],
    ]
    info_table = Table(info_data, colWidths=[65, 90, 90, 140, 65, 90])
    info_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (0, -1), rl_colors.whitesmoke),
        ("BACKGROUND", (2, 0), (2, -1), rl_colors.whitesmoke),
        ("BACKGROUND", (4, 0), (4, -1), rl_colors.whitesmoke),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 8))

    headers = ["Ctn No", "Color", "UPC"] + full_sizes + [
        "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"
    ]
    table_data = [headers]
    for rdata in rows:
        table_data.append([
            rdata["ctn_no"], rdata["color"], rdata["upc"],
            *[rdata["sizes"].get(sz, "") for sz in full_sizes],
            rdata["ctn_pcs"], rdata["total_ctn"], rdata["total_pcs"],
            f'{rdata["grs_per_ctn"]:.2f}', f'{rdata["net_per_ctn"]:.2f}',
            f'{rdata["total_grs"]:.2f}', f'{rdata["total_net"]:.2f}',
        ])
    total_row = (
        ["", "TOTAL", ""]
        + [export_state["size_packet_totals"].get(sz, 0) for sz in full_sizes]
        + ["", export_state["total_ctn"], export_state["total_pcs"], "", "",
           f'{export_state["total_grs"]:.2f}', f'{export_state["total_net"]:.2f}']
    )
    table_data.append(total_row)

    main_table = Table(table_data, repeatRows=1)
    main_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
        ("BACKGROUND", (0, -1), (-1, -1), rl_colors.lightgrey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    elements.append(main_table)
    elements.append(Spacer(1, 10))

    def fmt_pct(v):
        return "" if v is None else f"{v * 100:.2f}%"

    summary_data = [["SIZE"] + full_sizes + ["G TOTAL"]]
    summary_data.append(
        ["CUT QTY"] + [export_state["cut_qty"].get(sz, "") for sz in full_sizes]
        + [export_state["cut_total"]]
    )
    summary_data.append(
        ["ORDER QTY"] + [export_state["order_qty_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["order_total"]]
    )
    summary_data.append(
        ["SHIP QTY"] + [export_state["ship_qty_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["ship_total"]]
    )
    summary_data.append(
        ["EXS/SHT QTY"] + [export_state["exs_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["exs_total"]]
    )
    summary_data.append(
        ["PERCENTAGE"] + [fmt_pct(export_state["pct_by_size"].get(sz)) for sz in full_sizes]
        + [fmt_pct(export_state["pct_total"])]
    )

    summary_table = Table(summary_data, repeatRows=1)
    summary_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 10))

    wt_data = [
        ["GROSS WEIGHT :", f'{export_state["total_grs"]:.2f}'],
        ["NET WEIGHT :", f'{export_state["total_net"]:.2f}'],
        ["CTN MEAS :", f'{export_state["ctn_meas_1"]} {export_state["ctn_meas_2"]}'.strip()],
        ["CBM :", f'{export_state["cbm_total"]:.8f}'],
    ]
    wt_table = Table(wt_data, colWidths=[100, 150])
    wt_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
    ]))
    elements.append(wt_table)
    elements.append(Spacer(1, 20))

    sign_table = Table([["INCHARGE", "", "FM"]], colWidths=[120, 120, 120])
    sign_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (0, 0), 0.75, rl_colors.black),
        ("LINEABOVE", (2, 0), (2, 0), 0.75, rl_colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(sign_table)

    doc.build(elements)
    buf.seek(0)
    return buf


# ============================================================
# PART 4 - STREAMLIT UI & PREVIEW
# ============================================================


@_fragment

# def render_packing_list_preview(line, sublines):
#     style_key = line["style"]
#     full_sizes = list(dict.fromkeys([s["size"] for s in sublines])) # deduplicate while preserving order
#     default_color = next((s["color"] for s in sublines if s.get("color")), "MIXED")
#     pdf_max_by_size = {s["size"]: int(s.get("items_per_outer_pack") or 1) for s in sublines}

#     # ---- Weight spec table (native data_editor) ----
#     st.markdown("##### Weight spec (N.W. / N.N.W. / EMPTY CTN)")
#     spec_df = pd.DataFrame(
#         [[1.0] * len(full_sizes) for _ in range(3)],
#         index=["N.W.", "N.N.W.", "EMPTY CTN"],
#         columns=full_sizes,
#     )
#     spec_edit = st.data_editor(
#         spec_df, use_container_width=True, key=f"spec_v10_{style_key}",
#     )
#     nw_by_size = spec_edit.loc["N.W."].to_dict()
#     nnw_by_size = spec_edit.loc["N.N.W."].to_dict()
#     empty_by_size = spec_edit.loc["EMPTY CTN"].to_dict()
#     empty_ctn_wt = float(next(iter(empty_by_size.values()), 0.0))

#     # ---- Order quantity by size (HORIZONTAL: sizes as columns, no visible headers) ----
#     st.markdown("##### Order quantity by size")

#     ZWSP = "\u200b"  # zero-width space — unique per column, invisible in UI
#     order_data = {
#         ZWSP: ["SIZE", "COLOR", "UPC", "ORDER QTY"],
#     }
#     for i, s in enumerate(sublines):
#         order_data[f"{ZWSP}{i + 1}"] = [s["size"], s["color"], s["upc"], s["qty"]]

#     order_df = pd.DataFrame(order_data)
#     st.dataframe(order_df, use_container_width=True, hide_index=True)

#     # ---- MAX per carton (native data_editor) ----
#     st.markdown("##### Packing qty (MAX per carton)")
#     max_df = pd.DataFrame(
#         [[pdf_max_by_size.get(sz, 1) for sz in full_sizes]],
#         index=["MAX"], columns=full_sizes,
#     )
#     max_edit = st.data_editor(
#         max_df, use_container_width=True, key=f"pack_max_v6_{style_key}",
#         column_config={sz: st.column_config.NumberColumn(min_value=1, step=1) for sz in full_sizes},
#     )
#     max_by_size = {sz: int(max_edit.loc["MAX", sz]) for sz in full_sizes}

#     # ---- Packing mode ----
#     pack_mode = st.radio(
#         "Packing mode",
#         ["solid_no_leftover", "solid_all", "solid_mixed"],
#         index=2,
#         format_func=lambda k: {
#             "solid_no_leftover": "Solid, no leftover",
#             "solid_all": "Solid-all",
#             "solid_mixed": "Solid + mixed (current)",
#         }[k],
#         key=f"pack_mode_v2_{style_key}",
#         horizontal=True,
#     )

#     # ---- Build row state ----
#     rows_key = f"carton_rows_v10_{style_key}"
#     max_sig_key = f"max_sig_v10_{style_key}"
#     max_sig = tuple(sorted(max_by_size.items())) + (pack_mode,)
#     if rows_key not in st.session_state or st.session_state.get(max_sig_key) != max_sig:
#         base = compute_carton_rows_v3(sublines, max_by_size, mode=pack_mode)
#         st.session_state[rows_key] = [
#             {
#                 "id": uuid.uuid4().hex,
#                 "color": r["color"] or default_color,
#                 "upc": r["upc"] if not r["mixed"] else "MIXED",
#                 "sizes": {sz: int(r["sizes"].get(sz, 0) or 0) for sz in full_sizes},
#                 "total_ctn": int(r["total_ctn"]),
#                 "mixed": bool(r["mixed"]),
#             }
#             for r in base
#         ]
#         st.session_state[max_sig_key] = max_sig

#     rows = st.session_state[rows_key]

#     # ---- Main carton table (single native data_editor) ----
#     st.markdown("##### Main carton table (editable)")

#     def to_df(rows):
#         data = []
#         for r in rows:
#             rec = {"Delete": False, "Color": r["color"], "UPC": r["upc"]}
#             for sz in full_sizes:
#                 rec[sz] = int(r["sizes"].get(sz, 0) or 0)
#             rec["TOTAL CTN"] = int(r["total_ctn"])
#             data.append(rec)
#         return pd.DataFrame(data)

#     df = to_df(rows)
#     col_cfg = {"Delete": st.column_config.CheckboxColumn(width="small")}
#     for sz in full_sizes:
#         col_cfg[sz] = st.column_config.NumberColumn(min_value=0, step=1, format="%d")
#     col_cfg["TOTAL CTN"] = st.column_config.NumberColumn(min_value=0, step=1, format="%d")

#     edited = st.data_editor(
#         df,
#         use_container_width=True,
#         hide_index=True,
#         num_rows="dynamic",
#         key=f"grid_v10_{style_key}",
#         column_config=col_cfg,
#     )

#     new_rows = []
#     for i, (_, rec) in enumerate(edited.iterrows()):
#         old_id = rows[i]["id"] if i < len(rows) else uuid.uuid4().hex
#         new_rows.append({
#             "id": old_id,
#             "color": rec["Color"],
#             "upc": rec["UPC"],
#             "sizes": {sz: int(rec[sz] or 0) for sz in full_sizes},
#             "total_ctn": int(rec["TOTAL CTN"] or 0),
#             "mixed": str(rec["UPC"]).strip().upper() == "MIXED",
#         })

#     keep = [r for r, (_, rec) in zip(new_rows, edited.iterrows()) if not bool(rec["Delete"])]
#     if len(keep) != len(new_rows) and keep:
#         st.session_state[rows_key] = keep
#         rows = keep
#     else:
#         st.session_state[rows_key] = new_rows
#         rows = new_rows

#     # ---- Computed totals ----
#     size_packet_totals = {
#         sz: sum(int(r["sizes"].get(sz, 0)) * int(r["total_ctn"]) for r in rows)
#         for sz in full_sizes
#     }
#     total_ctn = sum(int(r["total_ctn"]) for r in rows)
#     total_pcs = sum(
#         sum(int(r["sizes"].get(sz, 0)) for sz in full_sizes) * int(r["total_ctn"])
#         for r in rows
#     )
#     grs_list = [
#         sum(int(r["sizes"].get(sz, 0)) * float(nw_by_size.get(sz, 0.0)) for sz in full_sizes)
#         + empty_ctn_wt
#         for r in rows
#     ]
#     net_list = [
#         sum(int(r["sizes"].get(sz, 0)) * float(nw_by_size.get(sz, 0.0)) for sz in full_sizes)
#         for r in rows
#     ]
#     total_grs = sum(g * int(r["total_ctn"]) for g, r in zip(grs_list, rows))
#     total_net = sum(n * int(r["total_ctn"]) for n, r in zip(net_list, rows))

#     # ---- Summary table (native st.table) ----
#     order_by = {s["size"]: s["qty"] for s in sublines}
#     ship_by = size_packet_totals
#     exs_by = {sz: ship_by.get(sz, 0) - order_by.get(sz, 0) for sz in full_sizes}
#     pct_by = {
#         sz: (exs_by[sz] / order_by[sz] * 100 if order_by.get(sz) else 0)
#         for sz in full_sizes
#     }
#     st.markdown("##### Summary")
#     st.table(pd.DataFrame({
#         "SIZE": full_sizes,
#         "ORDER QTY": [order_by.get(sz, 0) for sz in full_sizes],
#         "SHIP QTY": [ship_by.get(sz, 0) for sz in full_sizes],
#         "EXS/SHT QTY": [exs_by[sz] for sz in full_sizes],
#         "PERCENTAGE %": [round(pct_by[sz], 2) for sz in full_sizes],
#     }))

#     # ---- Live totals ----
#     m1, m2, m3, m4 = st.columns(4)
#     m1.metric("Total Cartons", total_ctn)
#     m2.metric("Total Pieces", total_pcs)
#     m3.metric("Total Gross Wt", f"{total_grs:.2f}")
#     m4.metric("Total Net Wt", f"{total_net:.2f}")

#     _exs_now = total_pcs - sum(order_by.values())
#     if _exs_now > 0:
#         st.success(f"Over-ship: **+{_exs_now}** pieces")
#     elif _exs_now < 0:
#         st.error(f"Short-ship: **{_exs_now}** pieces")
#     else:
#         st.info("Ship qty matches order qty exactly.")

#     # ---- CTN MEAS / CBM inputs ----
#     with st.expander("CTN MEAS / CBM", expanded=False):
#         d1, d2, d3 = st.columns(3)
#         ctn_l = d1.number_input("Length (in)", min_value=0.0, value=24.0, step=0.5, key=f"cl_{style_key}")
#         ctn_w = d2.number_input("Width (in)", min_value=0.0, value=16.0, step=0.5, key=f"cw_{style_key}")
#         ctn_h = d3.number_input("Height (in)", min_value=0.0, value=10.0, step=0.5, key=f"ch_{style_key}")
#         ctn_meas_1 = ""
#         ctn_meas_2 = f"L {ctn_l:g}'' X W {ctn_w:g}'' X H {ctn_h:g}''"
#         cbm_per_ctn = (ctn_l * ctn_w * ctn_h) / 61000

#     cbm_total = cbm_per_ctn * total_ctn

#     # ---- Export state ----
#     rows_export = []
#     cum = 0
#     for r, g, n in zip(rows, grs_list, net_list):
#         s_ctn = cum + 1
#         cum += int(r["total_ctn"])
#         rows_export.append({
#             "ctn_no": f"{s_ctn}-{cum}",
#             "color": r["color"],
#             "upc": r["upc"],
#             "sizes": r["sizes"],
#             "ctn_pcs": sum(r["sizes"].values()),
#             "total_ctn": int(r["total_ctn"]),
#             "grs_per_ctn": g,
#             "net_per_ctn": n,
#             "total_pcs": int(r["total_ctn"]) * sum(r["sizes"].values()),
#             "total_grs": int(r["total_ctn"]) * g,
#             "total_net": int(r["total_ctn"]) * n,
#         })

#     order_total = sum(order_by.values())
#     export_state = {
#         "line": line, "full_sizes": full_sizes, "rows": rows_export,
#         "total_ctn": total_ctn, "total_pcs": total_pcs,
#         "total_grs": total_grs, "total_net": total_net,
#         "size_packet_totals": size_packet_totals,
#         "order_qty_by_size": order_by, "ship_qty_by_size": ship_by,
#         "exs_by_size": exs_by, "pct_by_size": {sz: v / 100 for sz, v in pct_by.items()},
#         "order_total": order_total, "ship_total": total_pcs,
#         "exs_total": total_pcs - order_total,
#         "pct_total": ((total_pcs - order_total) / order_total if order_total else None),
#         "cut_qty": {sz: 0 for sz in full_sizes}, "cut_total": 0,
#         "ctn_meas_1": ctn_meas_1, "ctn_meas_2": ctn_meas_2,
#         "cbm_per_ctn": cbm_per_ctn, "cbm_total": cbm_total,
#         "nw_by_size": nw_by_size, "nnw_by_size": nnw_by_size, "empty_by_size": empty_by_size,
#     }

#     st.markdown("##### Download")
#     c1, c2 = st.columns(2)
#     with c1:
#         if st.button("Prepare Excel (.xlsx)", key=f"prep_xlsx_v10_{style_key}"):
#             wb = build_workbook({"lines": [line], "sublines": sublines})
#             b = io.BytesIO()
#             wb.save(b)
#             b.seek(0)
#             st.session_state[f"xlsx_bytes_v10_{style_key}"] = b.getvalue()
#         if st.session_state.get(f"xlsx_bytes_v10_{style_key}"):
#             st.download_button(
#                 f"Download combined (.xlsx)",
#                 data=st.session_state[f"xlsx_bytes_v10_{style_key}"],
#                 file_name=f"packing_list_combined.xlsx",
#                 mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#                 key=f"dl_xlsx_v10_{style_key}",
#             )
#     with c2:
#         if st.button("Prepare PDF", key=f"prep_pdf_v10_{style_key}"):
#             pb = build_edited_pdf(export_state)
#             st.session_state[f"pdf_bytes_v10_{style_key}"] = pb.getvalue()
#         if st.session_state.get(f"pdf_bytes_v10_{style_key}"):
#             st.download_button(
#                 f"Download combined (.pdf)",
#                 data=st.session_state[f"pdf_bytes_v10_{style_key}"],
#                 file_name=f"packing_list_combined.pdf",
#                 mime="application/pdf",
#                 key=f"dl_pdf_v10_{style_key}",
#             )


@_fragment
def render_packing_list_preview(line, sublines):
    style_key = line["style"]
    
    # Group and merge data for the SAME size from multiple POs
    grouped = {}
    for s in sublines:
        sz = s["size"]
        if sz not in grouped:
            grouped[sz] = {"colors": [], "upcs": [], "qty": 0}
        # Append if not already present
        if s["color"] not in grouped[sz]["colors"]:
            grouped[sz]["colors"].append(s["color"])
        if s["upc"] not in grouped[sz]["upcs"]:
            grouped[sz]["upcs"].append(s["upc"])
        grouped[sz]["qty"] += s["qty"]

    full_sizes = list(grouped.keys())  # Preserves order of appearance
    default_color = next((s["color"] for s in sublines if s.get("color")), "MIXED")
    
    # Get max pack by size (using the first occurrence or fallback to 1)
    pdf_max_by_size = {}
    for s in sublines:
        sz = s["size"]
        if sz not in pdf_max_by_size:
            pdf_max_by_size[sz] = int(s.get("items_per_outer_pack") or 1)

    # ---- Weight spec table (native data_editor) ----
    st.markdown("##### Weight spec (N.W. / N.N.W. / EMPTY CTN)")
    spec_df = pd.DataFrame(
        [[1.0] * len(full_sizes) for _ in range(3)],
        index=["N.W.", "N.N.W.", "EMPTY CTN"],
        columns=full_sizes,
    )
    spec_edit = st.data_editor(
        spec_df, use_container_width=True, key=f"spec_v10_{style_key}",
    )
    nw_by_size = spec_edit.loc["N.W."].to_dict()
    nnw_by_size = spec_edit.loc["N.N.W."].to_dict()
    empty_by_size = spec_edit.loc["EMPTY CTN"].to_dict()
    empty_ctn_wt = float(next(iter(empty_by_size.values()), 0.0))

    # ---- Order quantity by size (HORIZONTAL, MERGED) ----
    st.markdown("##### Order quantity by size")

    ZWSP = "\u200b"
    order_data = {
        ZWSP: ["SIZE", "COLOR", "UPC", "ORDER QTY"],
    }
    
    for i, sz in enumerate(full_sizes):
        colors = grouped[sz]["colors"]
        upcs = grouped[sz]["upcs"]
        
        # Merge multiple colors/upcs with commas
        color_str = ", ".join(colors)
        upc_str = ", ".join(upcs)
        
        order_data[f"{ZWSP}{i + 1}"] = [sz, color_str, upc_str, grouped[sz]["qty"]]

    order_df = pd.DataFrame(order_data)
    st.dataframe(order_df, use_container_width=True, hide_index=True)

    # ---- MAX per carton (native data_editor) ----
    st.markdown("##### Packing qty (MAX per carton)")
    max_df = pd.DataFrame(
        [[pdf_max_by_size.get(sz, 1) for sz in full_sizes]],
        index=["MAX"], columns=full_sizes,
    )
    max_edit = st.data_editor(
        max_df, use_container_width=True, key=f"pack_max_v6_{style_key}",
        column_config={sz: st.column_config.NumberColumn(min_value=1, step=1) for sz in full_sizes},
    )
    max_by_size = {sz: int(max_edit.loc["MAX", sz]) for sz in full_sizes}

    # ---- Packing mode ----
    pack_mode = st.radio(
        "Packing mode",
        ["solid_no_leftover", "solid_all", "solid_mixed"],
        index=2,
        format_func=lambda k: {
            "solid_no_leftover": "Solid, no leftover",
            "solid_all": "Solid-all",
            "solid_mixed": "Solid + mixed (current)",
        }[k],
        key=f"pack_mode_v2_{style_key}",
        horizontal=True,
    )

    # ---- Build row state ----
    rows_key = f"carton_rows_v10_{style_key}"
    max_sig_key = f"max_sig_v10_{style_key}"
    max_sig = tuple(sorted(max_by_size.items())) + (pack_mode,)
    
    # Note: We pass the merged sublines (or just rebuild based on max_by_size)
    if rows_key not in st.session_state or st.session_state.get(max_sig_key) != max_sig:
        # Reconstruct simplified sublines for the packing logic
        merged_sublines = [{"size": sz, "qty": grouped[sz]["qty"], "items_per_outer_pack": pdf_max_by_size[sz], "color": grouped[sz]["colors"][0], "upc": grouped[sz]["upcs"][0]} for sz in full_sizes]
        
        base = compute_carton_rows_v3(merged_sublines, max_by_size, mode=pack_mode)
        st.session_state[rows_key] = [
            {
                "id": uuid.uuid4().hex,
                "color": r["color"] or default_color,
                "upc": r["upc"] if not r["mixed"] else "MIXED",
                "sizes": {sz: int(r["sizes"].get(sz, 0) or 0) for sz in full_sizes},
                "total_ctn": int(r["total_ctn"]),
                "mixed": bool(r["mixed"]),
            }
            for r in base
        ]
        st.session_state[max_sig_key] = max_sig

    rows = st.session_state[rows_key]

    # ---- Main carton table (single native data_editor) ----
    st.markdown("##### Main carton table (editable)")

    def to_df(rows):
        data = []
        for r in rows:
            rec = {"Delete": False, "Color": r["color"], "UPC": r["upc"]}
            for sz in full_sizes:
                rec[sz] = int(r["sizes"].get(sz, 0) or 0)
            rec["TOTAL CTN"] = int(r["total_ctn"])
            data.append(rec)
        return pd.DataFrame(data)

    df = to_df(rows)
    col_cfg = {"Delete": st.column_config.CheckboxColumn(width="small")}
    for sz in full_sizes:
        col_cfg[sz] = st.column_config.NumberColumn(min_value=0, step=1, format="%d")
    col_cfg["TOTAL CTN"] = st.column_config.NumberColumn(min_value=0, step=1, format="%d")

    edited = st.data_editor(
        df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key=f"grid_v10_{style_key}",
        column_config=col_cfg,
    )

    new_rows = []
    for i, (_, rec) in enumerate(edited.iterrows()):
        old_id = rows[i]["id"] if i < len(rows) else uuid.uuid4().hex
        new_rows.append({
            "id": old_id,
            "color": rec["Color"],
            "upc": rec["UPC"],
            "sizes": {sz: int(rec[sz] or 0) for sz in full_sizes},
            "total_ctn": int(rec["TOTAL CTN"] or 0),
            "mixed": str(rec["UPC"]).strip().upper() == "MIXED",
        })

    keep = [r for r, (_, rec) in zip(new_rows, edited.iterrows()) if not bool(rec["Delete"])]
    if len(keep) != len(new_rows) and keep:
        st.session_state[rows_key] = keep
        rows = keep
    else:
        st.session_state[rows_key] = new_rows
        rows = new_rows

    # ---- Computed totals ----
    size_packet_totals = {
        sz: sum(int(r["sizes"].get(sz, 0)) * int(r["total_ctn"]) for r in rows)
        for sz in full_sizes
    }
    total_ctn = sum(int(r["total_ctn"]) for r in rows)
    total_pcs = sum(
        sum(int(r["sizes"].get(sz, 0)) for sz in full_sizes) * int(r["total_ctn"])
        for r in rows
    )
    grs_list = [
        sum(int(r["sizes"].get(sz, 0)) * float(nw_by_size.get(sz, 0.0)) for sz in full_sizes)
        + empty_ctn_wt
        for r in rows
    ]
    net_list = [
        sum(int(r["sizes"].get(sz, 0)) * float(nw_by_size.get(sz, 0.0)) for sz in full_sizes)
        for r in rows
    ]
    total_grs = sum(g * int(r["total_ctn"]) for g, r in zip(grs_list, rows))
    total_net = sum(n * int(r["total_ctn"]) for n, r in zip(net_list, rows))

    # ---- Summary table (native st.table) ----
    order_by = {sz: grouped[sz]["qty"] for sz in full_sizes}
    ship_by = size_packet_totals
    exs_by = {sz: ship_by.get(sz, 0) - order_by.get(sz, 0) for sz in full_sizes}
    pct_by = {
        sz: (exs_by[sz] / order_by[sz] * 100 if order_by.get(sz) else 0)
        for sz in full_sizes
    }
    st.markdown("##### Summary")
    st.table(pd.DataFrame({
        "SIZE": full_sizes,
        "ORDER QTY": [order_by.get(sz, 0) for sz in full_sizes],
        "SHIP QTY": [ship_by.get(sz, 0) for sz in full_sizes],
        "EXS/SHT QTY": [exs_by[sz] for sz in full_sizes],
        "PERCENTAGE %": [round(pct_by[sz], 2) for sz in full_sizes],
    }))

    # ---- Live totals ----
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total Cartons", total_ctn)
    m2.metric("Total Pieces", total_pcs)
    m3.metric("Total Gross Wt", f"{total_grs:.2f}")
    m4.metric("Total Net Wt", f"{total_net:.2f}")

    _exs_now = total_pcs - sum(order_by.values())
    if _exs_now > 0:
        st.success(f"Over-ship: **+{_exs_now}** pieces")
    elif _exs_now < 0:
        st.error(f"Short-ship: **{_exs_now}** pieces")
    else:
        st.info("Ship qty matches order qty exactly.")

    # ---- CTN MEAS / CBM inputs ----
    with st.expander("CTN MEAS / CBM", expanded=False):
        d1, d2, d3 = st.columns(3)
        ctn_l = d1.number_input("Length (in)", min_value=0.0, value=24.0, step=0.5, key=f"cl_{style_key}")
        ctn_w = d2.number_input("Width (in)", min_value=0.0, value=16.0, step=0.5, key=f"cw_{style_key}")
        ctn_h = d3.number_input("Height (in)", min_value=0.0, value=10.0, step=0.5, key=f"ch_{style_key}")
        ctn_meas_1 = ""
        ctn_meas_2 = f"L {ctn_l:g}'' X W {ctn_w:g}'' X H {ctn_h:g}''"
        cbm_per_ctn = (ctn_l * ctn_w * ctn_h) / 61000

    cbm_total = cbm_per_ctn * total_ctn

    # ---- Export state ----
    rows_export = []
    cum = 0
    for r, g, n in zip(rows, grs_list, net_list):
        s_ctn = cum + 1
        cum += int(r["total_ctn"])
        rows_export.append({
            "ctn_no": f"{s_ctn}-{cum}",
            "color": r["color"],
            "upc": r["upc"],
            "sizes": r["sizes"],
            "ctn_pcs": sum(r["sizes"].values()),
            "total_ctn": int(r["total_ctn"]),
            "grs_per_ctn": g,
            "net_per_ctn": n,
            "total_pcs": int(r["total_ctn"]) * sum(r["sizes"].values()),
            "total_grs": int(r["total_ctn"]) * g,
            "total_net": int(r["total_ctn"]) * n,
        })

    order_total = sum(order_by.values())
    export_state = {
        "line": line, "full_sizes": full_sizes, "rows": rows_export,
        "total_ctn": total_ctn, "total_pcs": total_pcs,
        "total_grs": total_grs, "total_net": total_net,
        "size_packet_totals": size_packet_totals,
        "order_qty_by_size": order_by, "ship_qty_by_size": ship_by,
        "exs_by_size": exs_by, "pct_by_size": {sz: v / 100 for sz, v in pct_by.items()},
        "order_total": order_total, "ship_total": total_pcs,
        "exs_total": total_pcs - order_total,
        "pct_total": ((total_pcs - order_total) / order_total if order_total else None),
        "cut_qty": {sz: 0 for sz in full_sizes}, "cut_total": 0,
        "ctn_meas_1": ctn_meas_1, "ctn_meas_2": ctn_meas_2,
        "cbm_per_ctn": cbm_per_ctn, "cbm_total": cbm_total,
        "nw_by_size": nw_by_size, "nnw_by_size": nnw_by_size, "empty_by_size": empty_by_size,
    }

    # st.markdown("##### Download")
    # c1, c2 = st.columns(2)
    # with c1:
    #     if st.button("Prepare Excel (.xlsx)", key=f"prep_xlsx_v10_{style_key}"):
    #         wb = build_workbook({"lines": [line], "sublines": sublines})
    #         b = io.BytesIO()
    #         wb.save(b)
    #         b.seek(0)
    #         st.session_state[f"xlsx_bytes_v10_{style_key}"] = b.getvalue()
    #     if st.session_state.get(f"xlsx_bytes_v10_{style_key}"):
    #         st.download_button(
    #             f"Download combined (.xlsx)",
    #             data=st.session_state[f"xlsx_bytes_v10_{style_key}"],
    #             file_name=f"packing_list_combined.xlsx",
    #             mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    #             key=f"dl_xlsx_v10_{style_key}",
    #         )
    # with c2:
    #     if st.button("Prepare PDF", key=f"prep_pdf_v10_{style_key}"):
    #         pb = build_edited_pdf(export_state)
    #         st.session_state[f"pdf_bytes_v10_{style_key}"] = pb.getvalue()
    #     if st.session_state.get(f"pdf_bytes_v10_{style_key}"):
    #         st.download_button(
    #             f"Download combined (.pdf)",
    #             data=st.session_state[f"pdf_bytes_v10_{style_key}"],
    #             file_name=f"packing_list_combined.pdf",
    #             mime="application/pdf",
    #             key=f"dl_pdf_v10_{style_key}",
    #         )
    # st.markdown("##### Download")
    # c1, c2 = st.columns(2)
    # with c1:
    #     if st.button("Prepare Excel (.xlsx)", key=f"prep_xlsx_v10_{style_key}"):
    #         # ✅ NEW: Create a single workbook and a single sheet manually
    #         wb = Workbook()
    #         ws = wb.active
    #         ws.title = "Combined_Packing"  # Give it a valid name
            
    #         # Call build_sheet directly with the combined line and sublines
    #         build_sheet(ws, line, sublines)
            
    #         b = io.BytesIO()
    #         wb.save(b)
    #         b.seek(0)
    #         st.session_state[f"xlsx_bytes_v10_{style_key}"] = b.getvalue()
    #     if st.session_state.get(f"xlsx_bytes_v10_{style_key}"):
    #         st.download_button(
    #             f"Download combined (.xlsx)",
    #             data=st.session_state[f"xlsx_bytes_v10_{style_key}"],
    #             file_name=f"packing_list_combined.xlsx",
    #             mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    #             key=f"dl_xlsx_v10_{style_key}",
    #         )
    # with c2:
    #     if st.button("Prepare PDF", key=f"prep_pdf_v10_{style_key}"):
    #         pb = build_edited_pdf(export_state)
    #         st.session_state[f"pdf_bytes_v10_{style_key}"] = pb.getvalue()
    #     if st.session_state.get(f"pdf_bytes_v10_{style_key}"):
    #         st.download_button(
    #             f"Download combined (.pdf)",
    #             data=st.session_state[f"pdf_bytes_v10_{style_key}"],
    #             file_name=f"packing_list_combined.pdf",
    #             mime="application/pdf",
    #             key=f"dl_pdf_v10_{style_key}",
    #         )


def run(file=None):
    if file is None:
        st.title("PO to Packing List Converter")
        st.caption("Upload a VF/Vans Purchase Order PDF, pick one or more PO lines, edit the combined carton table, "
                   "then download the Excel or PDF.")
    uploaded_file = file if file is not None else st.file_uploader("Upload PO PDF", type=["pdf"])
    if uploaded_file is not None:
        with st.spinner("Reading PDF..."):
            try:
                data = extract_data_cached(uploaded_file.getvalue())
            except Exception as e:
                st.error(f"Failed to read the PDF: {e}")
                st.stop()
    
        if not data["lines"]:
            st.warning("No PO lines were found. This PDF's layout may differ from the expected format.")
            st.stop()
    
        _upload_sig = f"{uploaded_file.name}::{uploaded_file.size}"
        if st.session_state.get("_last_upload_sig") != _upload_sig:
            for _k in list(st.session_state.keys()):
                if _k.startswith(("carton_rows_v10_", "max_sig_v10_",
                                  "pack_max_v6_", "pack_mode_v2_",
                                  "xlsx_bytes_v10_", "pdf_bytes_v10_",
                                  "spec_v10_", "grid_v10_",
                                  "cl_", "cw_", "ch_")):
                    st.session_state.pop(_k, None)
            st.session_state.pop("po_lines_table", None)
            st.session_state.pop("preview_styles", None)
            st.session_state["_last_upload_sig"] = _upload_sig
    
        # ---- 1 summary ----
        st.subheader("1. PO summary")
        _tot_qty = sum(int(l.get("order_qty") or 0) for l in data["lines"])
        c1, c2, c3 = st.columns(3)
        c1.metric("PO lines", len(data["lines"]))
        c2.metric("Size / colour sublines", len(data["sublines"]))
        c3.metric("Order qty", f"{_tot_qty:,}")
    
        # ---- 2 choose a PO line (MULTI-SELECT) ----
        st.subheader("2. Choose PO line(s)")
        st.caption("Click the checkbox to select one or more PO lines. Their items will be combined into one single table below.")
    
        po_line_df = pd.DataFrame([
            {
                "PO Line No.": l["po_line_no"],
                "Style": l["style"],
                "Description": l["description"],
                "Order Qty": l["order_qty"],
                "Unit Price": l["unit_price"],
                "Line Cost": l["line_cost"],
                "CRD": l.get("crd"),
                "Country": l.get("destination_country"),
            }
            for l in data["lines"]
        ])
        
        selection = st.dataframe(
            po_line_df,
            use_container_width=True,
            on_select="rerun",
            selection_mode="multi-row", 
            key="po_lines_table",
        )
    
        selected_idxs = []
        if selection is not None and hasattr(selection, "selection"):
            selected_idxs = selection.selection.rows
    
        selected_lines = [data["lines"][i] for i in selected_idxs if i < len(data["lines"])]
        st.session_state["preview_styles"] = [l["style"] for l in selected_lines]
    
        if selected_lines:
            st.info(f"Selected **{len(selected_lines)}** PO line(s): {', '.join(l['style'] for l in selected_lines)}")
        else:
            st.caption("No PO lines selected yet.")
    
        # ---- Sublines expander ----
        with st.expander("Sublines (size / colour / UPC)", expanded=False):
            if selected_lines:
                filtered_sublines = [
                    s for s in data["sublines"] 
                    if s["style"] in st.session_state["preview_styles"]
                ]
                st.dataframe(pd.DataFrame(filtered_sublines), use_container_width=True)
            else:
                st.dataframe(pd.DataFrame(data["sublines"]), use_container_width=True)
    
        # ---- 3 single combined carton table ----
        if selected_lines:
            st.subheader("3. Carton table (combined)")
            
            # Create a dummy "line" object so the table renders once
            combined_line = {
                "style": " + ".join([l["style"] for l in selected_lines]),
                "po_line_no": " | ".join([l["po_line_no"] for l in selected_lines]),
                "description": " | ".join([l["description"] for l in selected_lines]),
                "order_qty": sum([l["order_qty"] for l in selected_lines]),
                "crd": selected_lines[0].get("crd"),
                "destination_country": selected_lines[0].get("destination_country"),
            }
            
            # Combine sublines from all selected POs
            combined_sublines = [
                s for s in data["sublines"] 
                if s["style"] in st.session_state["preview_styles"]
            ]
            
            # Render once
            render_packing_list_preview(combined_line, combined_sublines)
        else:
            st.subheader("3. Carton table")
            st.caption("Select one or more PO lines above to see and edit the combined carton table.")
     
        # ---- 4 download ----
        st.subheader("4. Download")
        col1, col2, col3, col4, col5 = st.columns(5)
        
        with col1:
            # ✅ FIX: Return the buffer object itself, not bytes
            def get_full_wb():
                wb = build_workbook(data)
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                return buf  
                
            st.download_button(
                "⬇ Download all PO lines (Excel)",
                data=get_full_wb(),
                file_name="packing_list_all_PO_lines.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                key="dl_full_xlsx",
            )
                
        with col2:
            if selected_lines:
                # ✅ FIX: Return the buffer object itself
                def get_selected_wb():
                    selected_data = {
                        "lines": selected_lines,
                        "sublines": [s for s in data["sublines"] if s["style"] in st.session_state["preview_styles"]],
                    }
                    wb = build_workbook(selected_data)
                    buf = io.BytesIO()
                    wb.save(buf)
                    buf.seek(0)
                    return buf
                    
                st.download_button(
                    "⬇ Download selected PO line(s) (Excel)",
                    data=get_selected_wb(),
                    file_name="packing_list_selected.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_selected_xlsx",
                )
            else:
                st.caption("Select PO line(s) above to download just those sheets.")
                
        with col3:
            if selected_lines:
                combined_style_key = " + ".join([l["style"] for l in selected_lines])
                
                # ✅ FIX: Return the buffer object itself
                def get_combined_xlsx():
                    combined_line = {
                        "style": combined_style_key,
                        "po_line_no": " | ".join([l["po_line_no"] for l in selected_lines]),
                        "description": " | ".join([l["description"] for l in selected_lines]),
                        "order_qty": sum([l["order_qty"] for l in selected_lines]),
                        "crd": selected_lines[0].get("crd"),
                        "destination_country": selected_lines[0].get("destination_country"),
                    }
                    combined_sublines = [
                        s for s in data["sublines"] 
                        if s["style"] in st.session_state["preview_styles"]
                    ]
                    
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Combined_Packing"
                    build_sheet(ws, combined_line, combined_sublines)
                    
                    b = io.BytesIO()
                    wb.save(b)
                    b.seek(0)
                    return b
                    
                st.download_button(
                    "⬇ Download combined (.xlsx)",
                    data=get_combined_xlsx(),
                    file_name="packing_list_combined.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"dl_xlsx_v10_{combined_style_key}",
                )
            else:
                st.caption("Select PO line(s) to prepare combined Excel.")
                
        with col4:
            if selected_lines:
                combined_style_key = " + ".join([l["style"] for l in selected_lines])
                
                # ✅ FIX: Return the buffer object itself
                def get_combined_pdf():
                    combined_line = {
                        "style": combined_style_key,
                        "po_line_no": " | ".join([l["po_line_no"] for l in selected_lines]),
                        "description": " | ".join([l["description"] for l in selected_lines]),
                        "order_qty": sum([l["order_qty"] for l in selected_lines]),
                        "crd": selected_lines[0].get("crd"),
                        "destination_country": selected_lines[0].get("destination_country"),
                    }
                    
                    export_state = {
                        "line": combined_line, 
                        "full_sizes": [s["size"] for s in data["sublines"] if s["style"] in st.session_state["preview_styles"]], 
                        "rows": [],
                        "total_ctn": 0, "total_pcs": 0, "total_grs": 0, "total_net": 0,
                        "size_packet_totals": {}, "order_qty_by_size": {}, "ship_qty_by_size": {},
                        "exs_by_size": {}, "pct_by_size": {}, "cut_qty": {}, "cut_total": 0,
                        "order_total": 0, "ship_total": 0, "exs_total": 0, "pct_total": 0,
                        "ctn_meas_1": "", "ctn_meas_2": "", "cbm_per_ctn": 0, "cbm_total": 0,
                        "nw_by_size": {}, "nnw_by_size": {}, "empty_by_size": {}
                    }
                    
                    pb = build_edited_pdf(export_state)
                    pb.seek(0)
                    return pb
                    
                st.download_button(
                    "⬇ Download combined (.pdf)",
                    data=get_combined_pdf(),
                    file_name="packing_list_combined.pdf",
                    mime="application/pdf",
                    key=f"dl_pdf_v10_{combined_style_key}",
                )
            else:
                st.caption("Select PO line(s) to prepare combined PDF.")
    
        with col5:
            st.download_button(
                label="⬇ Download extracted JSON",
                data=json.dumps(data, indent=2).encode("utf-8"),
                file_name=f"po_extracted_{len(data['lines'])}_lines.json",
                mime="application/json",
                key="dl_po_json",
            )
    else:
        st.info("Upload a PDF to get started.")
        st.stop()

if __name__ == "__main__":
    st.set_page_config(page_title="VF PO -> Packing List", layout="wide", page_icon="📦")
    run()