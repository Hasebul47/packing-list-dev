

"""
Carter's / OshKosh B'Gosh PO -> Packing List
============================================
Reads the e-PLM "v14 Prod Purchase Order" PDF and emits one sheet per JESTA
line (PO-line) in the standard OSHKOSH B'GOSH packing list layout.

Three packing modes are supported and picked per line in the UI:
  * solid_size   - one size per carton (full + one partial)
  * assort_size  - mixed prepack, CTN PER BLSTR x blisters
  * single_mixed - everything in one carton

Standalone run:  streamlit run carters_app.py
Module use:      from carters_app import run;  run(pdf_file)
"""
import datetime
import io
import re
from collections import OrderedDict

import pdfplumber
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L

import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
if str(BASE) not in sys.path:
    sys.path.insert(0, str(BASE))

import grid_editor
import sys as _sys
from pathlib import Path as _Path
_ROOT = str(_Path(__file__).resolve().parent.parent)
if _ROOT not in _sys.path:                     # pl_kit.py lives in the project root
    _sys.path.insert(0, _ROOT)
import pl_kit as kit                           # noqa: E402

# ============================================================
# CONSTANTS / DEFAULTS
# ============================================================
SIZES = ["3M", "6M", "9M", "12M", "18M", "24M"]

DIM_DEFAULT = {
    "DIM-BLANK":  (23.5, 15,  6.5, 0.98),
    "DIM-EC002":  (23.5, 15,  6.5, 0.98),
    "DIM-RH8465": (23.5, 15,  6.5, 0.98),
    "MIX/RT981":  (23.5, 15, 13.0, 1.26),
    "MIX/RH8465": (23.5, 15, 13.0, 1.26),
}
DIM_ORDER = ["DIM-BLANK", "DIM-EC002", "DIM-RH8465", "MIX/RT981", "MIX/RH8465"]

NETWT_DEFAULT   = {"3M": 0.052, "6M": 0.055, "9M": 0.060,
                   "12M": 0.065, "18M": 0.070, "24M": 0.075}
N_NETWT_DEFAULT = {"3M": 0.030, "6M": 0.035, "9M": 0.040,
                   "12M": 0.045, "18M": 0.050, "24M": 0.055}
PACK_DEFAULT    = {s: 90 for s in SIZES}

INCH2M3 = 2.54 ** 3 / 1_000_000

# ============================================================
# PDF PARSING
# ============================================================
PO_RE     = re.compile(r"Purchase\s+order\s*#\s*:\s*(\d+)", re.I)
STYLE_RE  = re.compile(r"Style\s*:\s*(\S+)")
SEASON_RE = re.compile(r"Season\s*:\s*([A-Z0-9]+)\s*-\s*(\w+)")
COLOUR_RE = re.compile(r"(\d{3})\s+(\w+)\s+([A-Z]{2})\s+\(\d{3}\)")
VID_RE    = re.compile(r"Vendor\s*:\s*(\d+)", re.I)
FID_RE    = re.compile(r"Manufacturing\s+Factory\s*:\s*(\d+)", re.I)
LINE_RE   = re.compile(r"\bLine\s*:?\s*(\d{1,4})\b")

SIZE_HEADER_RE = re.compile(
    r"3M(?:/EC002)?\s*"
    r"6M(?:/EC002)?\s*"
    r"9M(?:/EC002)?\s*"
    r"12M(?:/EC002)?\s*"
    r"18M(?:/EC002)?\s*"
    r"24M(?:/EC002)?",
    re.S)

RATIO_RE = re.compile(r"Ratio[\s]*([\d\s]{6,40})")
UNITS_RE = re.compile(r"Units[\s]*([\d\s]{6,60})")

DEST_RE = re.compile(
    r"Ship\s+(?:"
    r"The\s+William\s+Carter\s+to\s*:?\s*Company|"
    r"LF\s+Logistics[^!]*?to\s*:?\s*Co\.\s*,?\s*Ltd\.\s*-?\s*DC\s*#?\s*\d+"
    r")"
    r"([^!]*?)"
    r"(?:Logistics\s+Plan|Mode\s*:)", re.S)


# ============================================================
# HELPERS
# ============================================================
def split_digits(s: str, n: int):
    L_ = len(s)
    if not (n <= L_ <= 4 * n):
        return None
    results = []

    def rec(rem, parts):
        if len(parts) == n:
            if not rem:
                results.append([int(p) for p in parts])
            return
        slots_left = n - len(parts) - 1
        for take in range(1, 5):
            if take > len(rem):
                break
            rem_after = len(rem) - take
            if rem_after < slots_left or rem_after > 4 * slots_left:
                continue
            rec(rem[take:], parts + [rem[:take]])

    rec(s, [])
    if not results:
        return None

    import statistics
    def score(parts):
        try:
            return (statistics.pstdev(parts), -sum(parts))
        except statistics.StatisticsError:
            return (0.0, -sum(parts))

    results.sort(key=score)
    return results[0]


def _cleanup(text: str) -> str:
    text = re.sub(r"([a-zA-Z])(\d)", r"\1 \2", text)
    text = re.sub(r"(\d)([A-Z][a-z])", r"\1 \2", text)
    return text


def parse_size_qtys(block: str):
    m = SIZE_HEADER_RE.search(block)
    if not m:
        return None
    tail = block[m.end(): m.end() + 400]
    i = tail.find("$")
    if i >= 0:
        tail = tail[:i]
    for line in tail.splitlines():
        parts = line.strip().split()
        if len(parts) == 6 and all(p.isdigit() for p in parts):
            return [int(p) for p in parts]
    for line in tail.splitlines():
        line = line.strip()
        if not line:
            continue
        digits = re.sub(r"\D", "", line)
        if len(digits) >= 6:
            res = split_digits(digits, 6)
            if res:
                return res
        break
    return None


def parse_prepack(block: str):
    ratio = units = None
    rm = RATIO_RE.search(block)
    um = UNITS_RE.search(block)
    if rm:
        raw = rm.group(1).strip()
        parts = raw.split()
        if len(parts) >= 6:
            ratio = [int(parts[i]) for i in range(6)]
        else:
            d = raw.replace(" ", "")
            if d.isdigit():
                ratio = split_digits(d, 6)
    if um:
        raw = um.group(1).strip()
        parts = raw.split()
        if len(parts) >= 6:
            units = [int(parts[i]) for i in range(6)]
        else:
            d = raw.replace(" ", "")
            if d.isdigit():
                units = split_digits(d, 6)
    return ratio, units


def parse_dest(block: str) -> str:
    m = DEST_RE.search(block)
    if m:
        d = re.sub(r"\s+", " ", m.group(1)).strip(" ,-")
        return _cleanup(d)
    m = re.search(r"Ship\s+(.+?)(?:Logistics\s+Plan|Mode\s*:)", block, re.S)
    if m:
        return _cleanup(re.sub(r"\s+", " ", m.group(1)).strip(" ,-"))
    return ""


# ============================================================
# PDF PARSER
# ============================================================
def parse_po(pdf_bytes: bytes):
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        text = "\n".join((p.extract_text() or "") for p in pdf.pages)

    def grab(pat, default="", flags=0):
        m = re.search(pat, text, flags)
        return m.group(1).strip() if m else default

    head = {
        "po":     grab(r"Purchase\s+order\s*#\s*:\s*(\d+)", "", re.I),
        "style":  grab(r"Style\s*:\s*(\S+)"),
        "vid":    grab(r"Vendor\s*:\s*(\d+)", "", re.I),
        "fid":    grab(r"Manufacturing\s+Factory\s*:\s*(\d+)", "", re.I),
        "colour": "", "cdesc": "", "season": "", "gender": "",
        "brand":  "OSHKOSH B'GOSH", "item": "",
    }
    cm = COLOUR_RE.search(text)
    if cm:
        head["colour"] = cm.group(1)
        head["cdesc"]  = f"{cm.group(2)} ({cm.group(1)})"
    sm = SEASON_RE.search(text)
    if sm:
        head["season"] = sm.group(1)
        head["gender"] = sm.group(2)
    dm = re.search(r"Style\s*:\s*\S+(\S.*?)(?:Season\s*:|\n)", text)
    if dm:
        head["item"] = re.sub(r"\s+", " ", dm.group(1)).strip()

    marks = list(LINE_RE.finditer(text))
    raw_blocks = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        block = text[m.end():end]
        raw_blocks.append((m.group(1), block))

    merged = []
    for ln, blk in raw_blocks:
        has_ship = "Ship" in blk
        has_size = bool(SIZE_HEADER_RE.search(blk) or RATIO_RE.search(blk))
        if not has_ship and not has_size and merged:
            merged[-1] = (merged[-1][0], merged[-1][1] + "\n" + blk)
        else:
            merged.append((ln, blk))

    lines = []
    for ln, block in merged:
        dest = parse_dest(block)
        qty = parse_size_qtys(block)
        ratio, units = parse_prepack(block)
        if qty is None and units is not None:
            qty = units
        has_ratio = ratio is not None
        if qty is None:
            qty = [0] * 6
        if ratio is None:
            ratio = [1] * 6
        lines.append({
            "line":      ln,
            "dest":      dest,
            "qty":       dict(zip(SIZES, qty)),
            "ratio":     ratio,
            "has_ratio": has_ratio,
            "raw":       block,
        })
    return head, lines


# ============================================================
# CARTON LOGIC
# ============================================================
def _pad_sizes(d: dict) -> dict:
    return {s: int(d.get(s, 0) or 0) for s in SIZES}


def build_solid_rows(qty, pack):
    rows = []
    for s in SIZES:
        q = int(qty.get(s, 0) or 0)
        if q <= 0:
            continue
        p = max(int(pack.get(s, 90)), 1)
        full, rem = divmod(q, p)
        if full:
            rows.append({"kind": "full",    "sizes": _pad_sizes({s: p}),   "ctn": full})
        if rem:
            rows.append({"kind": "partial", "sizes": _pad_sizes({s: rem}), "ctn": 1})
    return rows


def build_assort_rows(qty, ratio, blisters):
    per_ctn = [ratio[i] * blisters for i in range(6)]
    if not any(per_ctn):
        return []
    full = min((int(qty.get(SIZES[i], 0)) // per_ctn[i])
               for i in range(6) if per_ctn[i] > 0)
    rows = []
    if full:
        rows.append({"kind": "full",
                     "sizes": _pad_sizes({SIZES[i]: per_ctn[i] for i in range(6)}),
                     "ctn": full})
    rem = {SIZES[i]: int(qty.get(SIZES[i], 0)) - full * per_ctn[i] for i in range(6)}
    if sum(rem.values()) > 0:
        bl = min((rem[SIZES[i]] // ratio[i]) for i in range(6) if ratio[i] > 0)
        part = {SIZES[i]: ratio[i] * bl for i in range(6)}
        left = {SIZES[i]: rem[SIZES[i]] - part[SIZES[i]] for i in range(6)}
        for s in SIZES:
            part[s] += left[s]
        if sum(part.values()) > 0:
            rows.append({"kind": "partial",
                         "sizes": _pad_sizes(part),
                         "ctn": 1,
                         "blisters": bl})
    return rows


def build_single_row(qty):
    return [{"kind": "single",
             "sizes": _pad_sizes({s: int(qty.get(s, 0) or 0) for s in SIZES}),
             "ctn": 1}]


def carton_rows(mode, qty, pack, ratio, blisters):
    if mode == "solid_size":
        return build_solid_rows(qty, pack)
    if mode == "assort_size":
        return build_assort_rows(qty, ratio, blisters)
    return build_single_row(qty)


def carton_weights(rows, netwt, n_netwt, dim):
    Li, Wi, Hi, tare = DIM_DEFAULT.get(dim, (23.5, 15, 6.5, 0.98))
    m3 = Li * Wi * Hi * INCH2M3
    out = []
    for r in rows:
        sz = _pad_sizes(r["sizes"])
        pcs   = sum(sz[s] for s in SIZES)
        net   = round(sum(sz[s] * netwt[s]   for s in SIZES), 4)
        n_net = round(sum(sz[s] * n_netwt[s] for s in SIZES), 4)
        gross = round(net + tare, 2)
        out.append({**r, "sizes": sz, "pcs": pcs, "net": net, "n_net": n_net,
                    "gross": gross, "m3": m3, "tare": tare, "netwt": dict(netwt), "n_netwt": dict(n_netwt),
                    "L": Li, "W": Wi, "H": Hi})
    return out


# ============================================================
# EXCEL
# ============================================================
def build_workbook(head, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sh in sheets:
        name = f"{head['po']}-{sh['line']}"
        base, i = name, 1
        while name in used:
            i += 1
            name = f"{base}_{i}"
        used.add(name)
        _fill_sheet(wb.create_sheet(title=name[:31]), head, sh)
    return wb


def _fill_sheet(ws, head, sh):
    TNR, ARI, VER = "Times New Roman", "Arial", "Verdana"
    thin = Side(style="thin")
    CEN  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left",   vertical="center")

    def put(r, c, v=None, name=ARI, sz=10, al=CEN, fmt=None, bold=True):
        cell = ws.cell(row=r, column=c)
        if v is not None:
            cell.value = v
        cell.font = Font(name=name, size=sz, bold=bold)
        cell.alignment = al
        if fmt:
            cell.number_format = fmt
        return cell

    def merge(r1, c1, r2, c2):
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    def box(r1, c1, r2, c2, side=thin):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = Border(
                    left=side, right=side, top=side, bottom=side)

    put(1, 1, "PACKING LIST", TNR, 12, al=LEFT)
    put(2, 1, "Creative Collections Ltd (UNIT) 01A", TNR, 11, al=LEFT)
    put(3, 1, "69 & 107 , Nishat Nagor, Tongi, Gazipur. 1711", TNR, 11, al=LEFT)
    put(4, 1, head["brand"], TNR, 12, al=LEFT)

    put(6, 1, "FROM - VID #", TNR, 11, al=LEFT); put(6, 4, head["vid"], TNR, 11, al=LEFT)
    put(7, 1, "FROM - FID #", TNR, 11, al=LEFT); put(7, 4, head["fid"], TNR, 11, al=LEFT)
    put(8, 1, "PO", TNR, 11, al=LEFT);           put(8, 4, head["po"], TNR, 11, al=LEFT)
    put(8, 17, "DATE", TNR, 11, al=LEFT);        put(8, 20, sh["date"], TNR, 11, al=LEFT)
    put(9, 1, "STYLE NUMBER", TNR, 11, al=LEFT); put(9, 4, head["style"], TNR, 11, al=LEFT)
    put(9, 17, "TOTAL CARTON", TNR, 11, al=LEFT)
    put(9, 20, sh["total_ctn"], TNR, 11, al=LEFT); put(9, 21, "CTN", TNR, 11, al=LEFT)
    put(10, 1, "ITEM", TNR, 11, al=LEFT);        put(10, 4, head["item"], TNR, 11, al=LEFT)
    put(10, 7, sh["dim"], TNR, 11, al=LEFT)
    put(10, 17, "TOTAL GROSS WEIGHT", TNR, 11, al=LEFT)
    put(10, 20, sh["gw"], TNR, 11, al=LEFT);     put(10, 21, "KGS", TNR, 11, al=LEFT)
    put(11, 1, "COLOUR", TNR, 11, al=LEFT);      put(11, 4, head["cdesc"], TNR, 11, al=LEFT)
    put(11, 17, "TOTAL NET WEIGHT", TNR, 11, al=LEFT)
    put(11, 20, sh["nw"], TNR, 11, al=LEFT);     put(11, 21, "KGS", TNR, 11, al=LEFT)
    put(12, 1, "DESCRIPTION", TNR, 11, al=LEFT); put(12, 4, head["item"], TNR, 11, al=LEFT)
    put(12, 17, "TOTAL CBM", TNR, 11, al=LEFT)
    put(12, 20, round(sh["cbm"], 5), TNR, 11, al=LEFT, fmt="0.00000")
    put(12, 21, "m3", TNR, 11, al=LEFT)
    put(13, 1, "JESTA  P.O #", TNR, 11, al=LEFT)
    put(13, 4, f"{head['po']}-{sh['line']}", TNR, 11, al=LEFT)

    mode_lbl = {"solid_size":  "SOLID COLOR SOLID SIZE",
                "assort_size": "SOLID COLOR ASSORT SIZE",
                "single_mixed":"SOLID COLOR SOLID SIZE"}[sh["mode"]]
    put(17, 4, mode_lbl, TNR, 11, al=LEFT)

    R_H1, R_H2, R_D0 = 19, 20, 21
    C_SL, C_NO1, C_NO2, C_TCTN = 1, 4, 5, 6
    C_SZ0, C_SZL = 7, 12
    if sh["mode"] == "assort_size":
        C_B1, C_B2, C_B3, C_PC, C_PCS = C_SZL+1, C_SZL+2, C_SZL+3, C_SZL+4, C_SZL+5
    else:
        C_B1 = C_B2 = C_B3 = None
        C_PC  = C_SZL + 1
        C_PCS = C_SZL + 2
    C_GW  = C_PCS + 1
    C_NW  = C_GW  + 1
    C_NNW = C_NW  + 1
    C_TGW = C_NNW + 1
    C_TNW = C_TGW + 1
    C_TNN = C_TNW + 1
    C_L   = C_TNN + 1
    C_W   = C_L   + 1
    C_H   = C_W   + 1
    C_M3  = C_H   + 1
    C_LAST = C_M3

    put(R_H1, C_SL,   "CARTON S/L", TNR)
    put(R_H1, C_NO1,  "CARTON NO",  TNR); merge(R_H1, C_NO1, R_H1, C_NO2)
    put(R_H1, C_TCTN, "TOTAL CTN",  TNR)
    merge(R_H1, C_SZ0, R_H1, C_SZL)
    put(R_H1, C_SZ0, "SIZE", TNR)
    for i, s in enumerate(SIZES):
        put(R_H2, C_SZ0 + i, s, VER, fmt="@")
    if sh["mode"] == "assort_size":
        merge(R_H1, C_B1, R_H1, C_B3)
        put(R_H1, C_B1, "CTN PER BLSTR", TNR)
    put(R_H1, C_PC,  "PER CTN",   TNR)
    put(R_H1, C_PCS, "TOTAL PCS", TNR)
    put(R_H1, C_GW,  "GROSS WEIGHT", TNR)
    put(R_H1, C_NW,  "NET WEIGHT",   TNR)
    put(R_H1, C_NNW, "N.N.W",        TNR)
    put(R_H1, C_TGW, "TTL G.W.T",    TNR)
    put(R_H1, C_TNW, "TTL N.W.T",    TNR)
    put(R_H1, C_TNN, "TTL N.N.W",    TNR)
    merge(R_H1, C_L, R_H1, C_M3)
    put(R_H1, C_L, "CARTON MEASUREMENT IN INCH", TNR)
    put(R_H2, C_L, "L", TNR); put(R_H2, C_W, "W", TNR)
    put(R_H2, C_H, "H", TNR); put(R_H2, C_M3, "m3", TNR)
    box(R_H1, C_SL, R_H2, C_LAST)

    r = R_D0
    cno = 1
    for row in sh["rows"]:
        ctn = int(row["ctn"])
        put(r, C_SL,   "-", TNR, fmt="@")
        put(r, C_NO1,  cno if r == R_D0 else f"={L(C_NO2)}{r - 1}+1", TNR, fmt="0")
        put(r, C_NO2,  f"={L(C_NO1)}{r}+{L(C_TCTN)}{r}-1", TNR, fmt="0")
        put(r, C_TCTN, ctn, TNR, fmt="0")
        for i, s in enumerate(SIZES):
            v = row["sizes"].get(s, 0)
            put(r, C_SZ0 + i, v if v else None, VER)
        if sh["mode"] == "assort_size":
            bl = row.get("blisters", 0)
            if bl:
                put(r, C_B1, sh["ratio"][0], TNR, 9)
                put(r, C_B2, "X", TNR, 9)
                put(r, C_B3, bl, TNR, 9)
        # LIVE formulas: change a size qty / carton count and pcs, weights and m3 recalculate
        _a, _z = L(C_SZ0), L(C_SZ0 + len(SIZES) - 1)
        _nw = row.get("netwt") or {}
        _nn = row.get("n_netwt") or {}
        _t_nw = "+".join(f"{L(C_SZ0 + i)}{r}*{float(_nw.get(s, 0.0)):.6g}" for i, s in enumerate(SIZES))
        _t_nn = "+".join(f"{L(C_SZ0 + i)}{r}*{float(_nn.get(s, 0.0)):.6g}" for i, s in enumerate(SIZES))
        put(r, C_PC,  f"=SUM({_a}{r}:{_z}{r})", TNR, fmt="0")
        put(r, C_PCS, f"={L(C_TCTN)}{r}*{L(C_PC)}{r}", TNR, fmt="0")
        put(r, C_NW,  f"=ROUND({_t_nw},4)", TNR, fmt="0.00")
        put(r, C_GW,  f"=ROUND({L(C_NW)}{r}+{float(row.get('tare', 0.0)):.6g},2)", TNR, fmt="0.00")
        put(r, C_NNW, f"=ROUND({_t_nn},4)", TNR, fmt="0.00")
        put(r, C_TGW, f"=ROUND({L(C_GW)}{r}*{L(C_TCTN)}{r},2)", TNR, fmt="0.00")
        put(r, C_TNW, f"=ROUND({L(C_NW)}{r}*{L(C_TCTN)}{r},2)", TNR, fmt="0.00")
        put(r, C_TNN, f"=ROUND({L(C_NNW)}{r}*{L(C_TCTN)}{r},2)", TNR, fmt="0.00")
        put(r, C_L,   row["L"], TNR, fmt="0.0")
        put(r, C_W,   row["W"], TNR, fmt="0.0")
        put(r, C_H,   row["H"], TNR, fmt="0.0")
        put(r, C_M3,  f"=ROUND({L(C_L)}{r}*{L(C_W)}{r}*{L(C_H)}{r}*{INCH2M3!r}*{L(C_TCTN)}{r},8)", TNR, fmt="0.00000")
        cno += ctn
        r += 1

    R_TOT = r
    put(R_TOT, C_SL, "TOTAL", TNR)
    for i in range(6):
        c = C_SZ0 + i
        put(R_TOT, c, f"=SUM({L(c)}{R_D0}:{L(c)}{R_TOT-1})", TNR, fmt="0")
    put(R_TOT, C_TCTN, f"=SUM({L(C_TCTN)}{R_D0}:{L(C_TCTN)}{R_TOT-1})", TNR)
    put(R_TOT, C_PCS,  f"=SUM({L(C_PCS)}{R_D0}:{L(C_PCS)}{R_TOT-1})", TNR)
    for c in (C_GW, C_NW, C_NNW, C_TGW, C_TNW, C_TNN):
        put(R_TOT, c, f"=SUM({L(c)}{R_D0}:{L(c)}{R_TOT-1})", TNR, fmt="0.00")
    put(R_TOT, C_M3, f"=SUM({L(C_M3)}{R_D0}:{L(C_M3)}{R_TOT-1})", TNR, fmt="0.00000")
    box(R_TOT, C_SL, R_TOT, C_LAST)
    # header block totals follow the TOTAL row (live)
    put(9,  20, f"={L(C_TCTN)}{R_TOT}", TNR, 11, al=LEFT)
    put(10, 20, f"={L(C_TGW)}{R_TOT}", TNR, 11, al=LEFT, fmt="0.00")
    put(11, 20, f"={L(C_TNW)}{R_TOT}", TNR, 11, al=LEFT, fmt="0.00")
    put(12, 20, f"={L(C_M3)}{R_TOT}", TNR, 11, al=LEFT, fmt="0.00000")

    R_SU = R_TOT + 2
    put(R_SU, C_NO1, "SUMMARY BREAKDOWN", TNR, 11, al=LEFT)
    R_HD = R_SU + 2
    put(R_HD, C_NO1, "SIZE", TNR)
    for i, s in enumerate(SIZES):
        put(R_HD, C_SZ0 + i, s, TNR)
    put(R_HD, C_SZ0 + 6, "TOTAL", TNR)
    order_by_size = sh.get("order_qty") or {}
    r_o, r_s, r_x, r_p = R_HD + 1, R_HD + 2, R_HD + 3, R_HD + 4
    T = C_SZ0 + 6                                    # TOTAL column
    for rr, lbl in ((r_o, "ORDER QTY"), (r_s, "SHIP QTY"), (r_x, "SHORT / EXCESS"), (r_p, "PERCENTAGE")):
        put(rr, C_NO1, lbl, TNR, al=LEFT)
    for j in range(6):
        c = C_SZ0 + j
        cl = L(c)
        put(r_o, c, int(order_by_size.get(SIZES[j], 0) or 0), TNR)                       # PO order (input)
        # packed = size pcs per carton x number of cartons, summed over all rows (live)
        put(r_s, c, f"=SUMPRODUCT({cl}{R_D0}:{cl}{R_TOT-1},${L(C_TCTN)}${R_D0}:${L(C_TCTN)}${R_TOT-1})", TNR)
        put(r_x, c, f"={cl}{r_s}-{cl}{r_o}", TNR)                                        # excess / short
        put(r_p, c, f"=IF({cl}{r_o}=0,0,ROUND({cl}{r_s}/{cl}{r_o},4))", TNR, fmt="0.00%")
    for rr in (r_o, r_s, r_x):
        put(rr, T, f"=SUM({L(C_SZ0)}{rr}:{L(C_SZL)}{rr})", TNR)
    put(r_p, T, f"=IF({L(T)}{r_o}=0,0,ROUND({L(T)}{r_s}/{L(T)}{r_o},4))", TNR, fmt="0.00%")
    box(R_HD, C_NO1, R_HD + 4, C_SZ0 + 6)

    ws.column_dimensions["A"].width = 26
    for c in range(2, 6):
        ws.column_dimensions[L(c)].width = 6
    ws.column_dimensions[L(C_TCTN)].width = 9
    for i in range(6):
        ws.column_dimensions[L(C_SZ0 + i)].width = 6
    for c in range(C_PC, C_LAST + 1):
        ws.column_dimensions[L(c)].width = 11

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================
# STREAMLIT UI
# ============================================================
@st.cache_data(show_spinner=False, max_entries=5)
def _parse_cached(b: bytes):
    return parse_po(b)


def _render_header(head, sh):
    """Render the header info block (matching the Excel sheet header)."""
    left_items = [
        ("FROM - VID #",   head["vid"] or "—"),
        ("FROM - FID #",   head["fid"] or "—"),
        ("PO",             head["po"]),
        ("STYLE NUMBER",   head["style"]),
        ("ITEM",           head["item"] or "—"),
        ("COLOUR",         head["cdesc"] or "—"),
        ("DESCRIPTION",    head["item"] or "—"),
        ("JESTA P.O #",    f"{head['po']}-{sh['line']}"),
    ]
    right_items = [
        ("DATE",                str(sh["date"])),
        ("TOTAL CARTON",        f"{sh['total_ctn']} CTN"),
        ("TOTAL GROSS WEIGHT",  f"{sh['gw']:.2f} KGS"),
        ("TOTAL NET WEIGHT",    f"{sh['nw']:.2f} KGS"),
        ("TOTAL CBM",           f"{sh['cbm']:.5f} m3"),
        ("DIM CLASS",           sh["dim"]),
    ]

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("##### PO Header")
        for lbl, val in left_items:
            st.markdown(f"**{lbl}:**  {val}")
    with c2:
        st.markdown("##### Packing Info")
        for lbl, val in right_items:
            st.markdown(f"**{lbl}:**  {val}")


def _render_carton_table(sh):
    """Render the full carton table with proper headers like the Excel layout."""
    rows = sh["rows"]
    if not rows:
        st.info("No carton rows to display.")
        return

    # Build the display table matching the Excel column layout
    ctn_no = 1
    table_rows = []
    for row in rows:
        ctn = int(row["ctn"])
        r = {
            "CARTON S/L":       "-",
            "CARTON NO FROM":   ctn_no,
            "CARTON NO TO":     ctn_no + ctn - 1,
            "TOTAL CTN":        ctn,
        }
        for s in SIZES:
            v = row["sizes"].get(s, 0)
            r[s] = v if v else ""
        if sh["mode"] == "assort_size":
            bl = row.get("blisters", 0)
            r["CTN PER BLSTR"] = f"{sh['ratio'][0]} X {bl}" if bl else ""
        r["PER CTN"]       = row["pcs"]
        r["TOTAL PCS"]     = row["pcs"] * ctn
        r["GROSS WEIGHT"]  = round(row["gross"], 2)
        r["NET WEIGHT"]    = round(row["net"], 2)
        r["N.N.W"]         = round(row["n_net"], 2)
        r["TTL G.W.T"]     = round(row["gross"] * ctn, 2)
        r["TTL N.W.T"]     = round(row["net"] * ctn, 2)
        r["TTL N.N.W"]     = round(row["n_net"] * ctn, 2)
        r["L (in)"]        = row["L"]
        r["W (in)"]        = row["W"]
        r["H (in)"]        = row["H"]
        r["m3"]            = round(row["m3"] * ctn, 5)
        table_rows.append(r)
        ctn_no += ctn

    # Add a TOTAL row
    total_row = {
        "CARTON S/L":     "TOTAL",
        "CARTON NO FROM": "",
        "CARTON NO TO":   "",
        "TOTAL CTN":      sum(r["TOTAL CTN"] for r in table_rows),
    }
    for s in SIZES:
        total_row[s] = sum(int(r[s]) for r in table_rows if r[s] != "")
    if sh["mode"] == "assort_size":
        total_row["CTN PER BLSTR"] = ""
    total_row["PER CTN"]      = ""
    total_row["TOTAL PCS"]    = sum(r["TOTAL PCS"] for r in table_rows)
    total_row["GROSS WEIGHT"] = ""
    total_row["NET WEIGHT"]   = ""
    total_row["N.N.W"]        = ""
    total_row["TTL G.W.T"]    = round(sum(r["TTL G.W.T"] for r in table_rows), 2)
    total_row["TTL N.W.T"]    = round(sum(r["TTL N.W.T"] for r in table_rows), 2)
    total_row["TTL N.N.W"]    = round(sum(r["TTL N.N.W"] for r in table_rows), 2)
    total_row["L (in)"]       = ""
    total_row["W (in)"]       = ""
    total_row["H (in)"]       = ""
    total_row["m3"]           = round(sum(r["m3"] for r in table_rows), 5)
    table_rows.append(total_row)

    df = pd.DataFrame(table_rows)

    # Left-align descriptive columns, center everything else
    left_cols = ["CARTON S/L", "CARTON NO FROM", "CARTON NO TO", "CTN PER BLSTR"]
    left_cols = [c for c in left_cols if c in df.columns]
    styler = (
        df.style
        .set_properties(subset=left_cols, **{"text-align": "left", "font-weight": "bold"})
        .set_properties(
            subset=[c for c in df.columns if c not in left_cols],
            **{"text-align": "center"},
        )
        .apply(lambda r: [
            "font-weight: bold; background-color: #F0F0F0" if r["CARTON S/L"] == "TOTAL"
            else "" for _ in r
        ], axis=1)
    )
    st.dataframe(styler, use_container_width=True, hide_index=True)


# def _render_summary_table(sh, head):
#     """Render the summary breakdown table below the carton table."""
#     rows = sh["rows"]
#     if not rows:
#         return

#     # Compute per-size totals from the carton plan
#     packed_by_size = {s: 0 for s in SIZES}
#     for row in rows:
#         ctn = int(row["ctn"])
#         for s in SIZES:
#             packed_by_size[s] += row["sizes"].get(s, 0) * ctn

#     order_by_size = head.get("_order_qty", {s: 0 for s in SIZES})

#     summary_data = []
#     for label in ("ORDER QTY", "SHIP QTY", "SHORT / EXCESS", "PERCENTAGE"):
#         r = {"SIZE": label}
#         for s in SIZES:
#             if label == "ORDER QTY":
#                 r[s] = order_by_size.get(s, 0)
#             elif label == "SHIP QTY":
#                 r[s] = packed_by_size[s]
#             elif label == "SHORT / EXCESS":
#                 r[s] = packed_by_size[s] - order_by_size.get(s, 0)
#             else:  # PERCENTAGE
#                 o = order_by_size.get(s, 0)
#                 r[s] = 0.00 if o == 0 else round(packed_by_size[s] / o, 4)
#         total = sum(r[s] for s in SIZES if isinstance(r[s], (int, float)))
#         if label == "PERCENTAGE":
#             total_o = sum(order_by_size.values())
#             total_p = sum(packed_by_size.values())
#             total = 0.00 if total_o == 0 else round(total_p / total_o, 4)
#         r["TOTAL"] = total
#         summary_data.append(r)

#     df = pd.DataFrame(summary_data)
#     st.markdown("##### Summary Breakdown")

#     styler = (
#         df.style
#         .set_properties(subset=["SIZE"], **{"text-align": "left", "font-weight": "bold"})
#         .set_properties(
#             subset=[c for c in df.columns if c != "SIZE"],
#             **{"text-align": "center"},
#         )
#         .format({c: "{:.2f}" if c != "PERCENTAGE" and df.columns[0] != "PERCENTAGE" else "{}"
#                  for c in df.columns if c != "SIZE"})
#     )
#     # Format PERCENTAGE row and TOTAL column specifically
#     def _fmt(v):
#         if isinstance(v, float):
#             return f"{v:.4f}"
#         return v
#     st.dataframe(
#         df.style
#         .set_properties(subset=["SIZE"], **{"text-align": "left", "font-weight": "bold"})
#         .set_properties(
#             subset=[c for c in df.columns if c != "SIZE"],
#             **{"text-align": "center"},
#         )
#         .applymap(lambda v: "" if v == "" else v,
#                   subset=[c for c in df.columns if c != "SIZE"]),
#         use_container_width=True, hide_index=True,
#     )


def _render_summary_table(sh, head):
    """Render the summary breakdown table below the carton table."""
    rows = sh["rows"]
    if not rows:
        return

    # Compute per-size totals from the carton plan
    packed_by_size = {s: 0 for s in SIZES}
    for row in rows:
        ctn = int(row["ctn"])
        for s in SIZES:
            packed_by_size[s] += row["sizes"].get(s, 0) * ctn

    order_by_size = head.get("_order_qty", {s: 0 for s in SIZES})

    summary_data = []
    for label in ("ORDER QTY", "SHIP QTY", "SHORT / EXCESS", "PERCENTAGE"):
        r = {"SIZE": label}
        for s in SIZES:
            if label == "ORDER QTY":
                r[s] = int(order_by_size.get(s, 0))
            elif label == "SHIP QTY":
                r[s] = int(packed_by_size[s])
            elif label == "SHORT / EXCESS":
                r[s] = int(packed_by_size[s] - order_by_size.get(s, 0))
            else:  # PERCENTAGE
                o = order_by_size.get(s, 0)
                r[s] = 0.00 if o == 0 else round(packed_by_size[s] / o, 4)

        # TOTAL column
        if label == "PERCENTAGE":
            total_o = sum(order_by_size.values())
            total_p = sum(packed_by_size.values())
            r["TOTAL"] = 0.00 if total_o == 0 else round(total_p / total_o, 4)
        else:
            r["TOTAL"] = int(sum(r[s] for s in SIZES))
        summary_data.append(r)

    df = pd.DataFrame(summary_data)
    st.markdown("##### Summary Breakdown")

    # Build a plain-string display frame so the Styler only sees strings
    display_df = df.copy()
    for s in [c for c in df.columns if c != "SIZE"]:
        display_df[s] = df[s].apply(
            lambda v: f"{v:.4f}" if isinstance(v, float) else str(v)
        )

    styler = (
        display_df.style
        .set_properties(subset=["SIZE"],
                        **{"text-align": "left", "font-weight": "bold"})
        .set_properties(
            subset=[c for c in display_df.columns if c != "SIZE"],
            **{"text-align": "center"},
        )
    )
    st.dataframe(styler, use_container_width=True, hide_index=True)
def _line_ui(head, ln, key):
    total_q = sum(ln["qty"].values())
    kit.step(1, f"PO line {ln['line']}")
    kit.summary_cards([("Line", str(ln["line"])), ("PO qty", f"{total_q:,}"),
                       ("Style", str(head.get("style", ""))), ("Season", str(head.get("season", "")))])
    st.caption(f"Destination: {ln['dest'] or '(no destination found)'}  ·  " +
               "  ".join(f"{s}: {ln['qty'][s]}" for s in SIZES))

    if total_q == 0:
        st.error("No per-size quantity was found for this line. "
                 "Open the debug panel below and paste the text so we can adjust the parser.")
        with st.expander("Raw text of this block (debug)", expanded=True):
            st.code(ln["raw"][:3000])

    kit.step(2, "Packing settings", "Choose the packing mode, pieces per carton / ratio and weights.")
    mode = st.radio(
        "Packing mode",
        ["solid_size", "assort_size", "single_mixed"],
        index=1 if ln["has_ratio"] else 0,
        horizontal=True,
        format_func=lambda x: {"solid_size":  "Solid size (one size per ctn)",
                               "assort_size": "Assort size (mixed prepack)",
                               "single_mixed":"Single mixed carton"}[x],
        key=f"{key}_mode",
    )

    pack = {s: PACK_DEFAULT[s] for s in SIZES}
    ratio = list(ln["ratio"])
    blisters = 8
    dim = "MIX/RT981" if mode == "assort_size" else "DIM-EC002"

    if mode == "solid_size":
        cols = st.columns(6)
        for i, s in enumerate(SIZES):
            with cols[i]:
                pack[s] = st.number_input(f"{s} pcs/ctn", 1, 500, PACK_DEFAULT[s], 1,
                                          key=f"{key}_pk_{s}")
        dim = st.selectbox("Dim class", DIM_ORDER, index=1, key=f"{key}_dim1")
    elif mode == "assort_size":
        c1, c2 = st.columns(2)
        with c1:
            blisters = st.number_input("Blisters / carton", 1, 50, 8, 1, key=f"{key}_bl")
        with c2:
            dim = st.selectbox("Dim class", DIM_ORDER, index=3, key=f"{key}_dim2")
        st.markdown("**Ratio per blister**")
        rcols = st.columns(6)
        for i, s in enumerate(SIZES):
            with rcols[i]:
                ratio[i] = st.number_input(s, 0, 20, ln["ratio"][i], 1,
                                           key=f"{key}_rt_{s}")
    else:
        dim = st.selectbox("Dim class", DIM_ORDER, index=1, key=f"{key}_dim3")

    with st.expander("Per-piece weights (kg)", expanded=False):
        wcols = st.columns(6)
        netwt, n_netwt = {}, {}
        for i, s in enumerate(SIZES):
            with wcols[i]:
                netwt[s]   = st.number_input(f"{s} NET",   0.0, 1.0,
                                             NETWT_DEFAULT[s],   0.001, format="%.3f",
                                             key=f"{key}_nw_{s}")
                n_netwt[s] = st.number_input(f"{s} N.NET", 0.0, 1.0,
                                             N_NETWT_DEFAULT[s], 0.001, format="%.3f",
                                             key=f"{key}_nnw_{s}")

    Li, Wi, Hi, tare = DIM_DEFAULT.get(dim, (23.5, 15, 6.5, 0.98))
    m3 = Li * Wi * Hi * INCH2M3

    # Invalidate session rows if fundamental packing mode or pack/ratio changed
    pack_sig = (mode, tuple(sorted(pack.items())), tuple(ratio), blisters)
    if st.session_state.get(f"{key}_pack_sig") != pack_sig:
        st.session_state[f"{key}_pack_sig"] = pack_sig
        st.session_state.pop(f"{key}_carton_rows", None)
        st.session_state.pop(f"{key}_carton_defaults", None)

    # Initial default rows based on current settings
    base_rows = carton_rows(mode, ln["qty"], pack, ratio, blisters)
    initial_rows = []
    for r in base_rows:
        initial_rows.append({
            "kind": r.get("kind", "-"),
            "blisters": r.get("blisters", 0),
            "sizes": _pad_sizes(r.get("sizes", {})),
            "total_ctn": int(r.get("ctn", 1)),
        })

    kit.step(3, "Carton table",
             "Edit cells, add or delete carton rows in the grid. ORDER QTY / TOTAL / REMAIN stay pinned and update live; "
             "edited cells get a green border. Changing the packing settings above rebuilds the table.")
    # Header info block
    _render_header(head, {
        "line": ln["line"], "mode": mode, "dim": dim, "date": datetime.date.today(),
        "total_ctn": sum(r["ctn"] for r in base_rows),
        "gw": 0.0, "nw": 0.0, "cbm": 0.0,
    })
    st.markdown("---")

    # Render summary and measurement info below the grid inside the fragment
    def _draw_summary(grid_state):
        c_rows = []
        for r in grid_state["rows"]:
            sz = _pad_sizes(r["sizes"])
            ctn_count = int(r["total_ctn"])
            pcs = sum(sz.values())
            net = round(sum(sz[s] * netwt.get(s, 0.0) for s in SIZES), 4)
            n_net = round(sum(sz[s] * n_netwt.get(s, 0.0) for s in SIZES), 4)
            gross = round(net + tare, 2)
            c_rows.append({
                "kind": r.get("kind", "-"),
                "blisters": r.get("blisters", 0),
                "sizes": sz,
                "ctn": ctn_count,
                "pcs": pcs,
                "net": net,
                "n_net": n_net,
                "gross": gross,
                "m3": m3,
                "tare": tare, "netwt": dict(netwt), "n_netwt": dict(n_netwt),
                "L": Li, "W": Wi, "H": Hi,
            })
        tot_c = sum(r["ctn"] for r in c_rows)
        tot_cbm = sum(r["m3"] * r["ctn"] for r in c_rows)
        st.caption(f"**Measurement**: {Li:g}″ × {Wi:g}″ × {Hi:g}″ ({dim}) · **Total CBM**: `{tot_cbm:.5f} m³` ({tot_c} cartons)")
        temp_sh = {"rows": c_rows}
        head_with_qty = {**head, "_order_qty": ln["qty"]}
        _render_summary_table(temp_sh, head_with_qty)

    grid_res = grid_editor.render_carton_grid(
        key_prefix=key,
        sizes=SIZES,
        order_qty=ln["qty"],
        initial_rows=initial_rows,
        frozen_columns=[
            {"header": "S/L", "key": "kind", "width": 80},
        ],
        net_weight_by_size=netwt,
        empty_carton_weight=tare,
        show_weights=True,
        show_metrics=True,
        title=f"Main Carton Table (Editable) — Line {ln['line']}",
        render_below=_draw_summary,
    )

    # Convert edited grid rows to Carter's format
    edited_carton_rows = []
    for r in grid_res["rows"]:
        sz = _pad_sizes(r["sizes"])
        ctn_count = int(r["total_ctn"])
        pcs = sum(sz.values())
        net = round(sum(sz[s] * netwt.get(s, 0.0) for s in SIZES), 4)
        n_net = round(sum(sz[s] * n_netwt.get(s, 0.0) for s in SIZES), 4)
        gross = round(net + tare, 2)
        edited_carton_rows.append({
            "kind": r.get("kind", "-"),
            "blisters": r.get("blisters", 0),
            "sizes": sz,
            "ctn": ctn_count,
            "pcs": pcs,
            "net": net,
            "n_net": n_net,
            "gross": gross,
            "m3": m3,
            "tare": tare, "netwt": dict(netwt), "n_netwt": dict(n_netwt),
            "L": Li, "W": Wi, "H": Hi,
        })

    tot_ctn = sum(r["ctn"] for r in edited_carton_rows)
    tot_pcs = sum(r["pcs"] * r["ctn"] for r in edited_carton_rows)
    gw  = sum(r["gross"] * r["ctn"] for r in edited_carton_rows)
    nw  = sum(r["net"]   * r["ctn"] for r in edited_carton_rows)
    cbm = sum(r["m3"]    * r["ctn"] for r in edited_carton_rows)

    kit.step(4, "Check against the PO")
    packed = {sz: sum(r["sizes"].get(sz, 0) * r["ctn"] for r in edited_carton_rows) for sz in SIZES}
    kit.check_by_size(list(SIZES), {sz: ln["qty"].get(sz, 0) for sz in SIZES}, packed)
    kit.summary_cards([("Cartons", f"{tot_ctn:,}"), ("Pieces", f"{tot_pcs:,}"), ("CBM", f"{cbm:.3f}"),
                       ("Gross kg", f"{gw:,.1f}"), ("Net kg", f"{nw:,.1f}")])

    sh = {
        "line":      ln["line"],
        "mode":      mode,
        "dim":       dim,
        "ratio":     ratio,
        "date":      datetime.date.today(),
        "rows":      edited_carton_rows,
        "total_ctn": tot_ctn,
        "gw":  gw, "nw": nw, "cbm": cbm,
        "order_qty": dict(ln["qty"]),          # PO order qty per size (for the live summary block)
    }

    return sh


def run(pdf_file=None):
    st.title("Carter's / OshKosh — PO to Packing List")
    if pdf_file is None:
        st.info("Upload a Carter's e-PLM purchase order PDF.")
        return
    try:
        head, lines = _parse_cached(pdf_file.getvalue())
    except Exception as e:
        st.error(f"Could not read the PDF: {e}")
        return
    if not lines:
        st.warning("No line items found — the PDF layout may differ.")
        return

    st.success(f"PO **{head['po']}** · style **{head['style']}** · "
               f"{head['cdesc'] or '(colour not detected)'} · "
               f"{head['season']} {head['gender']} · {len(lines)} line(s)")

    sheets = []
    tabs = st.tabs([f"Line {ln['line']}" for ln in lines])
    for tab, ln in zip(tabs, lines):
        with tab:
            sheets.append(_line_ui(head, ln, key=f"ct_{head['po']}_{ln['line']}"))

    kit.step(5, "Download", "One Excel file with a sheet for every line. It always contains your latest edits.")
    wb = build_workbook(head, sheets)
    buf = io.BytesIO()
    wb.save(buf)
    kit.download_block("⬇ Download packing list (Excel)", buf.getvalue(), f"packing_list_{head['po']}.xlsx")


if __name__ == "__main__":
    st.set_page_config(page_title="Carter's PO -> Packing List", layout="wide")
    run(st.file_uploader("Upload Carter's PO PDF", type=["pdf"]))