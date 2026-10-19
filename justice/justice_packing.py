"""
Justice (Sports Products of America / Adjmi Apparel Group) buyer:
PO PDF -> factory packing list (Excel).

Pure logic, no Streamlit.  UI lives in justice_ui.py.

PO structure (one PO per PDF):
  * Each PO line = one STYLE/colour/label-code with
        total units, total CTN, pieces per carton, size ratio.
  * PREPACK line  (label W1..W5, WM): ratio over all sizes, e.g. 4 4 3 4 4 3 3 3
        -> every carton holds `pcs_per_ctn` pieces mixed by the ratio.
  * SOLID line    (label WL, 12J6 ...): ratio is one size only -> carton holds 12 pcs of that size.
  Cartons are exact (units = ctn x pcs/ctn), so there is no remainder carton.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import pdfplumber
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

# defaults (kg / cm)
DEF_NET = 0.35          # kg per garment (placeholder, user overrides)
DEF_TARE = 0.80         # kg per carton
DEF_DIM = (50.0, 40.0, 30.0)

THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"), bottom=Side(style="thin"))
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFFF00")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)

_SIZE_RE = re.compile(r"^J\d{1,2}$")
_NUM = r"[\d,]+"


# ------------------------------------------------------------------ data
@dataclass
class PoLine:
    label: str                 # WM / W1 / WL ...
    dim_pk: str                # 28JAA / 12J6
    kind: str                  # "PREPACK" | "SOLID"
    sizes: list[str]
    ratio: list[int]
    pcs_per_ctn: int
    ctns: int
    units: int
    unit_cost: float = 0.0
    delivery: str = ""

    @property
    def size_qty(self) -> dict[str, int]:
        return {s: r * self.ctns for s, r in zip(self.sizes, self.ratio)}


@dataclass
class PoData:
    po_no: str = ""
    ref_no: str = ""
    style: str = ""
    desc: str = ""
    color: str = ""
    wash: str = ""
    season: str = ""
    division: str = ""
    ship_date: str = ""
    ship_terms: str = ""
    ship_type: str = ""
    hs_code: str = ""
    buyer: str = "Sports Products of America"
    ship_to: str = "MIDAS IRWINDALE, 4889 4TH ST, IRWINDALE, CA 91706 UNITED STATES"
    agent: str = "CREATIVE COLLECTIONS LTD"
    total_ctn: int = 0
    total_units: int = 0
    lines: list[PoLine] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    order_qty: dict = field(default_factory=dict)   # size -> qty as printed on the PO (never edited)

    @property
    def sizes(self) -> list[str]:
        seen: list[str] = list(self.order_qty)
        for ln in self.lines:
            for s in ln.sizes:
                if s not in seen:
                    seen.append(s)
        return sorted(seen, key=lambda s: int(s[1:]))

    @property
    def order_by_size(self) -> dict[str, int]:
        return {s: int(self.order_qty.get(s, 0)) for s in self.sizes}

    def packed_by_size(self) -> dict[str, int]:
        out = {s: 0 for s in self.sizes}
        for ln in self.lines:
            for s, q in ln.size_qty.items():
                out[s] = out.get(s, 0) + q
        return out


def _i(s: str) -> int:
    return int(s.replace(",", ""))


# ------------------------------------------------------------------ parser
def parse_po(data: bytes) -> PoData:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [(pg.extract_text() or "") for pg in pdf.pages]
    full = "\n".join(pages)
    po = PoData()

    # ---- header (page 1 text, labels on one line / values on the next)
    m = re.search(r"P\.O\. NUMBER.*\n(\d{5,})\s+(\d\d/\d\d/\d\d)\s+(\w+)\s+(.+?)\s+(Net [\w ]+?)\s+(\w+)\s+(\w+)\s+(\d\d/\d\d/\d\d)", full)
    if m:
        po.po_no, po.ship_date, po.ship_terms = m.group(1), m.group(2), m.group(3)
        po.ship_type, po.division, po.season = m.group(4), m.group(6), m.group(7)
    else:
        m = re.search(r"\n(\d{7})\s+(\d\d/\d\d/\d\d)", full)
        if m:
            po.po_no, po.ship_date = m.group(1), m.group(2)
        else:
            po.warnings.append("PO number / ship date not found.")
    m = re.search(r"Issue/Make\s+(.+?)\s+HAMEEM", full)
    if m:
        po.ref_no = m.group(1)
    else:
        m = re.search(r"Issue/Make\s+(\S+)", full)
        po.ref_no = m.group(1) if m else ""
    m = re.search(r"HS CODE\s*\n?.*?(\d{4}\.\d\d\.\d{4})", full)
    po.hs_code = m.group(1) if m else ""

    # ---- blocks: every "STYLE: <no>" starts one PO line
    chunks = re.split(r"(?=STYLE:\s*\d{4,})", full)
    for ch in chunks:
        sm = re.match(r"STYLE:\s*(\d{4,})\s*(\d\d/\d\d/\d\d)?", ch)
        if not sm:
            continue
        # cut the block at page footer / totals
        ch = re.split(r"SPECIAL INSTRUCTIONS|PRINTED:", ch)[0]
        lines_ = [l.strip() for l in ch.splitlines() if l.strip()]

        hdr = re.search(r"Total((?:\s+J\d{1,2})+)", ch)
        sizes = hdr.group(1).split() if hdr else []

        nums = re.search(rf"^({_NUM})\s+({_NUM})\s+({_NUM})\s+\$([\d.,]+)\s+\$([\d.,]+)\s*$",
                         ch, re.M)
        ratio = None
        for i, l in enumerate(lines_):
            if l.startswith("COLOR CODE:") and i + 1 < len(lines_):
                nxt = lines_[i + 1]
                if re.fullmatch(r"\d+(?:\s+\d+)*", nxt):
                    ratio = [int(x) for x in nxt.split()]
        lab = re.search(r"LABEL CODE:\s*(\w+)", ch)
        pcs = re.search(r"Total\s+(\d+)\s+pieces in a carton", ch)
        dim = re.search(r"DIM/PK:\s*(\w+)", ch)
        col = re.search(r"COLOR CODE:\s*([A-Z0-9 ]+?)\s*(?:\n|$)", ch)
        desc = re.search(r"\n(JUSTICE[^\n]*?)\n", ch) or re.search(r"\n([A-Z][A-Z ]{6,})\n\d", ch)
        wash = re.search(r"\n([A-Z ]+ WASH)\b", ch)

        if not (nums and sizes and ratio and pcs and lab):
            po.warnings.append(f"Could not read a PO line near style {sm.group(1)}.")
            continue
        if len(ratio) != len(sizes):
            po.warnings.append(f"Ratio/size count mismatch in label {lab.group(1)}.")
            continue

        units, ctns = _i(nums.group(1)), _i(nums.group(2))
        pc = int(pcs.group(1))
        if units != ctns * pc:
            po.warnings.append(f"Label {lab.group(1)}: units {units} != ctn {ctns} x {pc} pcs.")
        if sum(ratio) != pc:
            po.warnings.append(f"Label {lab.group(1)}: ratio sums to {sum(ratio)}, not {pc} pcs/ctn.")

        po.style = po.style or sm.group(1)
        if desc and not po.desc:
            po.desc = desc.group(1).strip()
        if col and not po.color:
            po.color = col.group(1).strip()
        if wash and not po.wash:
            po.wash = wash.group(1).strip()
        po.lines.append(PoLine(
            label=lab.group(1), dim_pk=dim.group(1) if dim else "",
            kind="SOLID" if len(sizes) == 1 else "PREPACK",
            sizes=sizes, ratio=ratio, pcs_per_ctn=pc, ctns=ctns, units=units,
            unit_cost=float(nums.group(4).replace(",", "")),
            delivery=sm.group(2) or po.ship_date,
        ))

    for ln in po.lines:
        for sz, q in ln.size_qty.items():
            po.order_qty[sz] = po.order_qty.get(sz, 0) + q

    # ---- totals cross-check
    tm = re.search(r"TOTALS\s+([\d,]+)\s+([\d,]+)\s+\$([\d.,]+)", full)
    po.total_ctn = sum(l.ctns for l in po.lines)
    po.total_units = sum(l.units for l in po.lines)
    if tm:
        if _i(tm.group(1)) != po.total_ctn or _i(tm.group(2)) != po.total_units:
            po.warnings.append(
                f"PO totals say {tm.group(1)} ctn / {tm.group(2)} pcs but lines add up to "
                f"{po.total_ctn:,} ctn / {po.total_units:,} pcs - check that no line was missed.")
    return po


# ------------------------------------------------------------------ editable table <-> lines
COL_LABEL, COL_PACK, COL_CTNS = "Label", "Pack", "Ctns"
COL_L, COL_W, COL_H, COL_NET = "L (cm)", "W (cm)", "H (cm)", "Net kg / pc"


def _num(v, default=0.0) -> float:
    try:
        f = float(v)
        return default if f != f else f      # NaN -> default
    except (TypeError, ValueError):
        return default


def _txt(v) -> str:
    if v is None or (isinstance(v, float) and v != v):
        return ""
    return str(v).strip()


def rows_from_po(po: PoData) -> list[dict]:
    """The editable table the user starts from (one row per PO line)."""
    sizes = po.sizes
    out = []
    for ln in po.lines:
        r = {COL_LABEL: ln.label, COL_PACK: ln.dim_pk, COL_CTNS: ln.ctns}
        for s in sizes:
            r[s] = ln.ratio[ln.sizes.index(s)] if s in ln.sizes else 0
        r.update({COL_L: DEF_DIM[0], COL_W: DEF_DIM[1], COL_H: DEF_DIM[2], COL_NET: DEF_NET})
        out.append(r)
    return out


def lines_from_rows(rows: list[dict], sizes: list[str]):
    """Edited table -> (lines, dims, net_pc, problems). Blank / invalid rows are skipped and reported."""
    lines, dims, net, problems = [], {}, {}, []
    for n, r in enumerate(rows, 1):
        label = _txt(r.get(COL_LABEL))
        ratio_all = [int(max(0, _num(r.get(s)))) for s in sizes]
        ctns = int(_num(r.get(COL_CTNS)))
        if not label and not any(ratio_all) and ctns == 0:
            continue                                  # empty new row
        if not label:
            label = f"NEW{n}"
        used = [(s, q) for s, q in zip(sizes, ratio_all) if q > 0]
        if not used:
            problems.append(f"Row {n} ({label}): no size ratio entered - row ignored.")
            continue
        if ctns <= 0:
            problems.append(f"Row {n} ({label}): cartons must be above 0 - row ignored.")
            continue
        pcs = sum(q for _, q in used)
        lines.append(PoLine(
            label=label, dim_pk=_txt(r.get(COL_PACK)) or label,
            kind="SOLID" if len(used) == 1 else "PREPACK",
            sizes=[s for s, _ in used], ratio=[q for _, q in used],
            pcs_per_ctn=pcs, ctns=ctns, units=ctns * pcs))
        i = len(lines) - 1
        dims[i] = tuple(_num(r.get(c), d) if _num(r.get(c), d) > 0 else d
                        for c, d in zip((COL_L, COL_W, COL_H), DEF_DIM))
        net[i] = _num(r.get(COL_NET), DEF_NET)
    return lines, dims, net, problems


# ------------------------------------------------------------------ carton plan
@dataclass
class CartonRow:
    line: PoLine
    ctn_from: int
    ctn_to: int
    L: float
    W: float
    H: float
    net_pc: float
    tare: float

    @property
    def ctns(self): return self.ctn_to - self.ctn_from + 1
    @property
    def pcs_per_ctn(self): return self.line.pcs_per_ctn
    @property
    def total_pcs(self): return self.ctns * self.pcs_per_ctn
    @property
    def cbm(self): return round(self.L * self.W * self.H * self.ctns / 1e6, 3)
    @property
    def nw_ctn(self): return round(self.net_pc * self.pcs_per_ctn, 2)
    @property
    def gw_ctn(self): return round(self.nw_ctn + self.tare, 2)


def build_carton_plan(po: PoData, dims: dict, net_pc: dict, tare: float) -> list[CartonRow]:
    """dims: {line_index: (L,W,H)};  net_pc: {line_index: kg per piece}."""
    rows, start = [], 1
    for i, ln in enumerate(po.lines):
        L, W, H = dims.get(i, DEF_DIM)
        rows.append(CartonRow(ln, start, start + ln.ctns - 1, L, W, H,
                              net_pc.get(i, DEF_NET), tare))
        start += ln.ctns
    return rows


# ------------------------------------------------------------------ excel
def _label(ln: PoLine) -> str:
    return f"{ln.label}  ({ln.dim_pk})" if ln.dim_pk else ln.label


def make_packing_excel(po: PoData, plan: list[CartonRow],
                       company: str = "", address: str = "") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packing List"
    sizes = po.sizes
    ns = len(sizes)
    # columns: FROM TO CTNS LABEL | sizes... | PCS/CTN TOTAL PCS | L W H | CBM | NW/ctn GW/ctn | TOT NW TOT GW
    c_from, c_to, c_ct, c_lab = 1, 2, 3, 4
    c_s0 = 5
    c_pc = c_s0 + ns
    c_tp = c_pc + 1
    c_L, c_W, c_H = c_tp + 1, c_tp + 2, c_tp + 3
    c_cbm = c_H + 1
    c_nw, c_gw, c_tnw, c_tgw = c_cbm + 1, c_cbm + 2, c_cbm + 3, c_cbm + 4
    last = c_tgw

    def merged(r, text, size=10, bold=False, fill=None, c1=1, c2=last):
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        c = ws.cell(r, c1, text)
        c.font = F(size=size, bold=bold)
        c.alignment = CENTER
        if fill:
            c.fill = fill
        return c

    r = 1
    merged(r, company or "PACKING LIST", 16, True); ws.row_dimensions[r].height = 26; r += 1
    if address:
        merged(r, address, 10); r += 1
    merged(r, "PACKING LIST", 14, True); r += 2

    half = (last) // 2
    info = [
        ("PO NO", po.po_no, "BUYER", po.buyer),
        ("REF NO", po.ref_no, "SHIP TO", po.ship_to),
        ("STYLE", po.style, "SHIP DATE", po.ship_date),
        ("DESCRIPTION", po.desc, "SEASON / DIV", f"{po.season} / {po.division}"),
        ("COLOUR", f"{po.color} {('- ' + po.wash) if po.wash else ''}".strip(), "SHIP TERMS", f"{po.ship_terms} · {po.ship_type}"),
    ]
    for a, b, c, d in info:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=half)
        ws.merge_cells(start_row=r, start_column=half + 1, end_row=r, end_column=half + 2)
        ws.merge_cells(start_row=r, start_column=half + 3, end_row=r, end_column=last)
        for col, v, bold in ((1, a, True), (3, b, False), (half + 1, c, True), (half + 3, d, False)):
            cell = ws.cell(r, col, v)
            cell.font = F(bold=bold)
            cell.alignment = CENTER
        for col in range(1, last + 1):
            ws.cell(r, col).border = THIN
        r += 1
    r += 1

    # ---- header
    head = r
    ws.merge_cells(start_row=head, start_column=c_from, end_row=head, end_column=c_to)
    ws.cell(head, c_from, "CTN NO")
    ws.merge_cells(start_row=head, start_column=c_s0, end_row=head, end_column=c_s0 + ns - 1)
    ws.cell(head, c_s0, "SIZE RATIO  (pcs per carton)")
    ws.merge_cells(start_row=head, start_column=c_L, end_row=head, end_column=c_H)
    ws.cell(head, c_L, "CTN MEAS (cm)")
    ws.merge_cells(start_row=head, start_column=c_nw, end_row=head, end_column=c_gw)
    ws.cell(head, c_nw, "WEIGHT / CTN (kg)")
    ws.merge_cells(start_row=head, start_column=c_tnw, end_row=head, end_column=c_tgw)
    ws.cell(head, c_tnw, "TOTAL WEIGHT (kg)")
    sub = head + 1
    for col, t in ((c_from, "FROM"), (c_to, "TO"), (c_ct, "CTNS"), (c_lab, "LABEL / PACK"),
                   (c_pc, "PCS / CTN"), (c_tp, "TOTAL PCS"), (c_L, "L"), (c_W, "W"), (c_H, "H"),
                   (c_cbm, "CBM"), (c_nw, "NET"), (c_gw, "GROSS"), (c_tnw, "NET"), (c_tgw, "GROSS")):
        ws.cell(sub, col, t)
    for j, s in enumerate(sizes):
        ws.cell(sub, c_s0 + j, s)
    for rr in (head, sub):
        for col in range(1, last + 1):
            c = ws.cell(rr, col)
            c.font = F(bold=True); c.alignment = CENTER; c.fill = GREY; c.border = THIN
    for col in (c_ct, c_lab, c_pc, c_tp, c_cbm):
        txt = ws.cell(sub, col).value
        ws.merge_cells(start_row=head, start_column=col, end_row=sub, end_column=col)
        ws.cell(head, col, txt)
        ws.cell(head, col).font = F(bold=True); ws.cell(head, col).alignment = CENTER
    ws.row_dimensions[head].height = 30
    ws.row_dimensions[sub].height = 20

    # ---- body
    r = sub + 1
    first = r
    for row in plan:
        ln = row.line
        ws.cell(r, c_from, row.ctn_from)
        ws.cell(r, c_to, row.ctn_to)
        ws.cell(r, c_ct, row.ctns)
        ws.cell(r, c_lab, _label(ln))
        for j, s in enumerate(sizes):
            if s in ln.sizes:
                ws.cell(r, c_s0 + j, ln.ratio[ln.sizes.index(s)])
        ws.cell(r, c_pc, f"=SUM({CL(c_s0)}{r}:{CL(c_s0 + ns - 1)}{r})")
        ws.cell(r, c_tp, f"={CL(c_ct)}{r}*{CL(c_pc)}{r}")
        ws.cell(r, c_L, row.L); ws.cell(r, c_W, row.W); ws.cell(r, c_H, row.H)
        ws.cell(r, c_cbm, f"=ROUND({CL(c_L)}{r}*{CL(c_W)}{r}*{CL(c_H)}{r}*{CL(c_ct)}{r}/1000000,3)")
        ws.cell(r, c_nw, row.nw_ctn)
        ws.cell(r, c_gw, row.gw_ctn)
        ws.cell(r, c_tnw, f"=ROUND({CL(c_nw)}{r}*{CL(c_ct)}{r},2)")
        ws.cell(r, c_tgw, f"=ROUND({CL(c_gw)}{r}*{CL(c_ct)}{r},2)")
        for col in range(1, last + 1):
            c = ws.cell(r, col)
            c.font = F(); c.alignment = CENTER; c.border = THIN
        for col in (c_L, c_W, c_H, c_nw, c_gw):
            ws.cell(r, col).fill = YELLOW
        r += 1
    lastrow = r - 1

    # ---- totals
    ws.cell(r, c_from, "TOTAL")
    ws.merge_cells(start_row=r, start_column=c_from, end_row=r, end_column=c_to)
    ws.cell(r, c_ct, f"=SUM({CL(c_ct)}{first}:{CL(c_ct)}{lastrow})")
    for j in range(ns):
        col = c_s0 + j
        ws.cell(r, col, f"=SUMPRODUCT({CL(col)}{first}:{CL(col)}{lastrow},${CL(c_ct)}${first}:${CL(c_ct)}${lastrow})")
    for col in (c_tp, c_cbm, c_tnw, c_tgw):
        ws.cell(r, col, f"=SUM({CL(col)}{first}:{CL(col)}{lastrow})")
    for col in range(1, last + 1):
        c = ws.cell(r, col)
        c.font = F(bold=True); c.alignment = CENTER; c.border = THIN; c.fill = GREY
    tot = r
    r += 2

    # ---- summary by size
    ws.cell(r, c_lab, "SIZE SUMMARY"); ws.cell(r, c_lab).font = F(bold=True)
    r += 1
    labels = ["ORDER QTY", "PACKED QTY", "EXCESS / SHORT"]
    order = po.order_by_size
    for k, lab in enumerate(labels):
        rr = r + k
        ws.merge_cells(start_row=rr, start_column=c_from, end_row=rr, end_column=c_lab)
        ws.cell(rr, c_from, lab)
        for j, s in enumerate(sizes):
            col = c_s0 + j
            if k == 0:
                ws.cell(rr, col, order[s])
            elif k == 1:
                ws.cell(rr, col, f"={CL(col)}{tot}")
            else:
                ws.cell(rr, col, f"={CL(col)}{r + 1}-{CL(col)}{r}")
        ws.cell(rr, c_tp, f"=SUM({CL(c_s0)}{rr}:{CL(c_s0 + ns - 1)}{rr})")
        for col in range(c_from, c_tp + 1):
            if col in range(c_to + 1, c_lab) or (col == c_pc):
                continue
            c = ws.cell(rr, col)
            c.font = F(bold=(col in (c_from, c_tp))); c.alignment = CENTER
            c.border = THIN
            if col == c_from:
                c.fill = GREY
    r += 4

    ws.cell(r, c_from, "Packed by: ______________").font = F()
    ws.cell(r, c_s0 + 2, "Checked by: ______________").font = F()

    widths = {c_from: 7, c_to: 7, c_ct: 8, c_lab: 20, c_pc: 9, c_tp: 11, c_L: 6, c_W: 6, c_H: 6,
              c_cbm: 9, c_nw: 8, c_gw: 8, c_tnw: 10, c_tgw: 10}
    for col in range(1, last + 1):
        ws.column_dimensions[CL(col)].width = widths.get(col, 8)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "FFC000"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.freeze_panes = ws.cell(sub + 1, c_s0)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
